"""Phase 0 proof: one-way cash flow with scenario events and a native Monte Carlo.

Builds a self-checking workbook that needs no add-in, no macros and no Python in
Excel to calculate:

- a scenario events register (bullet, repeat, spread, start and end, percentage)
  driving a monthly cash engine with a facility that draws to hold a cash floor;
- a risk register (PERT, triangular, normal, discrete, Bernoulli);
- seeded draws from a counter-based generator (HDR, 2019 form), Latin hypercube
  sampling and a Gaussian copula built from a plain-formula Cholesky grid;
- one native data table that reruns the whole model once per trial;
- statistics, distribution report, fan chart, convergence checkpoints, rank
  sensitivity, an inspect-trial replay and a check against this file's own
  Python reference calculation.

    python prototypes/montecarlo/simulation.py [--trials 2000] [out.xlsx]

The reference calculation below is the specification; the workbook must match it.
"""

from __future__ import annotations

import argparse
from dataclasses import dataclass, field
from pathlib import Path

import numpy as np
from scipy import stats

MONTHS = 24
FIRST_MONTH_COL = 10  # J
HDR_SOURCE = "Hubbard Decision Research, HDR PRNG (2019)"


# --------------------------------------------------------------------------
# Inputs (fictional demo numbers)
# --------------------------------------------------------------------------

@dataclass
class Driver:
    name: str
    dist: str            # PERT, Triangular, Normal, Discrete, Bernoulli
    var_id: int
    base: float
    p1: float | None = None   # min, mean or probability
    p2: float | None = None   # mode or SD
    p3: float | None = None   # max
    units: str = ""


@dataclass
class Event:
    id: str
    label: str
    kind: str            # Bullet, Repeat, Spread, Start and end, Percentage
    amount: float | None = None
    start: int | None = None
    end: int | None = None
    every: int | None = None
    pct: float | None = None
    of: str | None = None
    lag: int | None = None
    scale_by: str | None = None
    shift_by: str | None = None
    only_if: str | None = None


@dataclass
class Inputs:
    seed: int = 1
    opening_cash: float = 200.0
    cash_floor: float = 50.0
    facility_limit: float = 800.0
    base_receipts: float = 100.0
    base_payments: float = 60.0
    drivers: list[Driver] = field(default_factory=lambda: [
        Driver("Receipts growth", "PERT", 101, 0.01, -0.01, 0.01, 0.03, "% a month"),
        Driver("Construction overrun", "Triangular", 102, 0.05, 0.0, 0.05, 0.25, "%"),
        Driver("Settlement delay", "Discrete", 103, 0, units="months"),
        Driver("Interest rate", "Normal", 104, 0.065, 0.065, 0.01, units="% a year"),
        Driver("Consent delay", "Bernoulli", 105, 0, 0.2, units="flag"),
    ])
    delay_table: list[tuple[int, float]] = field(default_factory=lambda: [(0, 0.5), (1, 0.3), (2, 0.15), (3, 0.05)])
    correlation: dict[tuple[int, int], float] = field(default_factory=lambda: {(0, 3): -0.4, (1, 3): 0.3})
    events: list[Event] = field(default_factory=lambda: [
        Event("E1", "Land settlement", "Bullet", -500, 2),
        Event("E2", "Construction", "Spread", -1200, 3, 14, scale_by="Construction overrun"),
        Event("E3", "Insurance", "Repeat", -30, 1, 24, every=6),
        Event("E4", "GST refund on construction", "Percentage", pct=-0.15, of="E2", lag=2),
        Event("E5", "Operating grant", "Start and end", 40, 6, 17),
        Event("E6", "Sale of units", "Bullet", 1800, 16, shift_by="Settlement delay"),
        Event("E7", "Consent delay cost", "Bullet", -150, 4, only_if="Consent delay"),
    ])

    def corr_matrix(self) -> np.ndarray:
        k = len(self.drivers)
        a = np.eye(k)
        for (i, j), r in self.correlation.items():
            a[i, j] = a[j, i] = r
        return a


# --------------------------------------------------------------------------
# Reference calculation
# --------------------------------------------------------------------------

def hdr(trial: int, var: int, ent: int) -> float:
    """Counter-based uniform on (0, 1): a pure function of trial, variable and entity."""
    a = (999999999999989 % ((trial * 2499997 + var * 1800451 + ent * 2000371) % 7450589 * 4658 + 7450581)) * 383 % 99991
    b = (999999999999989 % ((trial * 2246527 + var * 2399993 + ent * 2100869) % 7450987 * 7580 + 7560584)) * 17669 % 7440893
    return (((a * 7440893 + b) * 1343) % 4294967296 + 0.5) / 4294967296


def draws(inp: Inputs, n: int) -> np.ndarray:
    """Sampled driver values, trials x drivers."""
    k = len(inp.drivers)
    trials = np.arange(1, n + 1)
    z = np.zeros((n, k))
    for j, d in enumerate(inp.drivers):
        u1 = np.array([hdr(int(t), d.var_id, inp.seed) for t in trials])
        u2 = np.array([hdr(int(t), d.var_id + 1000, inp.seed) for t in trials])
        rank = stats.rankdata(u2, method="min")
        z[:, j] = stats.norm.ppf((rank - 1 + u1) / n)
    chol = np.linalg.cholesky(inp.corr_matrix())
    uc = stats.norm.cdf(z @ chol.T)
    out = np.zeros((n, k))
    for j, d in enumerate(inp.drivers):
        u = uc[:, j]
        if d.dist == "PERT":
            lo, mode, hi = d.p1, d.p2, d.p3
            a = 1 + 4 * (mode - lo) / (hi - lo)
            b = 1 + 4 * (hi - mode) / (hi - lo)
            out[:, j] = stats.beta.ppf(u, a, b, loc=lo, scale=hi - lo)
        elif d.dist == "Triangular":
            lo, mode, hi = d.p1, d.p2, d.p3
            c = (mode - lo) / (hi - lo)
            out[:, j] = np.where(u < c, lo + np.sqrt(u * (hi - lo) * (mode - lo)),
                                 hi - np.sqrt((1 - u) * (hi - lo) * (hi - mode)))
        elif d.dist == "Normal":
            out[:, j] = stats.norm.ppf(u, d.p1, d.p2)
        elif d.dist == "Discrete":
            lower = np.cumsum([0] + [p for _, p in inp.delay_table[:-1]])
            values = np.array([v for v, _ in inp.delay_table])
            out[:, j] = values[np.searchsorted(lower, u, side="right") - 1]
        elif d.dist == "Bernoulli":
            out[:, j] = (u < d.p1).astype(float)
    return out


def cash_model(inp: Inputs, live: dict[str, float]) -> dict:
    """The monthly one-way cash engine for one set of driver values."""
    t = np.arange(1, MONTHS + 1)
    receipts = inp.base_receipts * (1 + live["Receipts growth"]) ** (t - 1)
    payments = -inp.base_payments * np.ones(MONTHS)
    rows: dict[str, np.ndarray] = {}
    for ev in inp.events:
        shift = live[ev.shift_by] if ev.shift_by else 0
        scale = 1 + (live[ev.scale_by] if ev.scale_by else 0)
        include = live[ev.only_if] if ev.only_if else 1
        s = (ev.start or 0) + shift
        e = (ev.end + shift) if ev.end else s
        if ev.kind == "Bullet":
            v = (t == s) * ev.amount
        elif ev.kind == "Repeat":
            v = (t >= s) * (t <= e) * ((t - s) % ev.every == 0) * ev.amount
        elif ev.kind == "Spread":
            v = (t >= s) * (t <= e) * ev.amount / (e - s + 1)
        elif ev.kind == "Start and end":
            v = (t >= s) * (t <= e) * ev.amount
        elif ev.kind == "Percentage":
            src = rows[ev.of]
            v = np.array([ev.pct * src[m - ev.lag - 1] if m - ev.lag >= 1 else 0.0 for m in t])
        rows[ev.id] = v * scale * include
    events = sum(rows.values())
    cash, debt = inp.opening_cash, 0.0
    before, closing_debt, interest = [], [], []
    for m in range(MONTHS):
        i = -debt * live["Interest rate"] / 12
        cb = cash + receipts[m] + payments[m] + events[m] + i
        draw = max(0.0, inp.cash_floor - cb)
        repay = min(debt, max(0.0, cb - inp.cash_floor))
        debt = debt + draw - repay
        cash = cb + draw - repay
        before.append(cb)
        closing_debt.append(debt)
        interest.append(i)
    before, closing_debt = np.array(before), np.array(closing_debt)
    above = closing_debt > inp.facility_limit
    return {
        "peak_debt": closing_debt.max(), "peak_month": int(closing_debt.argmax()) + 1,
        "low_cash": before.min(), "low_month": int(before.argmin()) + 1,
        "interest": -sum(interest), "net_cash": cash - debt,
        "breach": float(closing_debt.max() > inp.facility_limit),
        "first_above": int(above.argmax()) + 1 if above.any() else 0,
        "debt": closing_debt,
    }


OUTPUTS = [("peak_debt", "Peak debt", "$"), ("peak_month", "Month of peak debt", "month"),
           ("low_cash", "Lowest cash before funding", "$"), ("low_month", "Month of lowest cash", "month"),
           ("interest", "Total interest", "$"), ("net_cash", "Closing net cash", "$"),
           ("breach", "Limit breached", "flag"), ("first_above", "First month above limit", "month")]


def simulate(inp: Inputs, n: int) -> dict:
    d = draws(inp, n)
    names = [x.name for x in inp.drivers]
    trials = [cash_model(inp, dict(zip(names, row))) for row in d]
    table = np.array([[r[k] for k, _, _ in OUTPUTS] + list(r["debt"]) for r in trials])
    base = cash_model(inp, {x.name: x.base for x in inp.drivers})
    rho = []
    peak_rank = stats.rankdata(table[:, 0])
    for j in range(len(names)):
        rho.append(np.corrcoef(stats.rankdata(d[:, j]), peak_rank)[0, 1])
    rho = np.array(rho)
    return {"draws": d, "table": table, "base": base, "rho": rho, "contribution": rho ** 2 / (rho ** 2).sum()}


# --------------------------------------------------------------------------
# Workbook
# --------------------------------------------------------------------------

def _xl_hdr(trial: str, var: str, ent: str) -> str:
    return ("(MOD((MOD(MOD(999999999999989,MOD({t}*2499997+{v}*1800451+{e}*2000371,7450589)*4658+7450581)*383,99991)"
            "*7440893+MOD(MOD(999999999999989,MOD({t}*2246527+{v}*2399993+{e}*2100869,7450987)*7580+7560584)"
            "*17669,7440893))*1343,4294967296)+0.5)/4294967296").format(t=f"({trial})", v=f"({var})", e=f"({ent})")


def build(inp: Inputs, n: int, path: Path) -> dict:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.styles import Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter as L
    from openpyxl.workbook.defined_name import DefinedName
    from openpyxl.worksheet.formula import DataTableFormula
    from openpyxl.worksheet.hyperlink import Hyperlink

    ref = simulate(inp, n)
    k = len(inp.drivers)
    last = 6 + n
    TEXT = "404040"
    f_body = Font(name="Segoe UI", size=9, color=TEXT)
    f_bold = Font(name="Segoe UI", size=9, color=TEXT, bold=True)
    f_head = Font(name="Segoe UI", size=10, color="FFFFFF", bold=True)
    f_title = Font(name="Segoe UI", size=10, color=TEXT, bold=True)
    f_check = Font(name="Segoe UI", size=9, color="9C0006")
    fill_in = PatternFill("solid", fgColor="FFF2CC")
    fill_head = PatternFill("solid", fgColor="44546A")
    top = Border(top=Side(style="thin", color=TEXT))
    NUM = '#,##0.00;(#,##0.00);"-"'
    PCT = "0.0%"
    wb = Workbook()
    wb.remove(wb.active)

    def sheet(name: str, purpose: str, width_g: int = 30):
        ws = wb.create_sheet(name)
        ws.sheet_view.showGridLines = False
        ws["B1"], ws["B2"] = name, purpose
        ws["B1"].font, ws["B2"].font = f_title, f_body
        ws.column_dimensions["A"].width = 2.5
        for c in "BCDEF":
            ws.column_dimensions[c].width = 2.5
        ws.column_dimensions["G"].width = width_g
        ws.column_dimensions["H"].width = 9
        ws.column_dimensions["I"].width = 12
        return ws

    def put(ws, cell, value, font=None, fmt=None, fill=None):
        c = ws[cell]
        c.value = value
        c.font = font or f_body
        if fmt:
            c.number_format = fmt
        if fill:
            c.fill = fill
        return c

    def heading(ws, row, text, last_col=FIRST_MONTH_COL + MONTHS - 1):
        for c in range(2, last_col + 1):
            ws.cell(row, c).fill = fill_head
        put(ws, f"B{row}", text, f_head)

    def name(nm, ws, cell):
        col = "".join(ch for ch in cell if ch.isalpha())
        row = "".join(ch for ch in cell if ch.isdigit())
        wb.defined_names[nm] = DefinedName(nm, attr_text=f"'{ws.title}'!${col}${row}")

    contents = sheet("Contents", "One-way cash flow with scenario events and a native Monte Carlo (demo data)", 40)

    # Control -----------------------------------------------------------
    ctl = sheet("Control", "Settings for the cash model and the simulation")
    heading(ctl, 5, "Simulation", 9)
    rows = [("MC_Seed", "Seed (feeds the generator)", inp.seed, "", True),
            ("MC_Trials", "Trials (fixed when the model is built)", n, "", False),
            ("MC_Inspect", "Inspect trial (0 shows the base case)", 0, "", True),
            ("MC_Trial", "Trial in use", "=IF(Run!$C$3>0,Run!$C$3,MC_Inspect)", "", False)]
    heading(ctl, 11, "Cash model", 9)
    rows2 = [("Cash_Opening", "Opening cash", inp.opening_cash, "$", True),
             ("Cash_Floor", "Minimum cash", inp.cash_floor, "$", True),
             ("Fac_Limit", "Facility limit", inp.facility_limit, "$", True),
             ("Base_Receipts", "Base receipts in month 1", inp.base_receipts, "$", True),
             ("Base_Payments", "Base payments a month", inp.base_payments, "$", True)]
    for start, block in ((6, rows), (12, rows2)):
        for i, (nm, label, v, unit, is_input) in enumerate(block):
            r = start + i
            put(ctl, f"C{r}", label)
            put(ctl, f"H{r}", unit)
            put(ctl, f"I{r}", v, fmt=NUM if unit == "$" else "0", fill=fill_in if is_input else None)
            name(nm, ctl, f"I{r}")

    # Risk register -------------------------------------------------------
    risk = sheet("Risk", "Risk register: each driver's distribution, base value and live value", 4)
    hdr_cols = ["ID", "Driver", "Distribution", "Var ID", "Base", "Min / mean / p", "Mode / SD", "Max", "Units",
                "Live value", "PERT alpha", "PERT beta"]
    for i, h in enumerate(hdr_cols):
        put(risk, f"{L(2 + i)}6", h, f_bold)
    for c, w in zip("BCDEFGHIJKLM", [4, 22, 11, 7, 9, 12, 10, 9, 10, 10, 9, 9]):
        risk.column_dimensions[c].width = w
    for j, d in enumerate(inp.drivers):
        r = 7 + j
        vals = [j + 1, d.name, d.dist, d.var_id, d.base, d.p1, d.p2, d.p3, d.units]
        for i, v in enumerate(vals):
            put(risk, f"{L(2 + i)}{r}", v, fill=fill_in if 4 <= i <= 7 and v is not None else None)
        vcol = L(3 + j * 8 + 7)
        put(risk, f"K{r}", f"=IF(MC_Trial=0,F{r},INDEX(Draws!${vcol}$7:${vcol}${last},MC_Trial))", f_bold)
        if d.dist == "PERT":
            put(risk, f"L{r}", f"=1+4*(H{r}-G{r})/(I{r}-G{r})")
            put(risk, f"M{r}", f"=1+4*(I{r}-H{r})/(I{r}-G{r})")
        name("Drv_" + d.name.replace(" ", "_"), risk, f"K{r}")
    put(risk, "B14", "Settlement delay table", f_bold)
    for i, h in enumerate(["Months", "Probability", "Cumulative from"]):
        put(risk, f"{L(2 + i)}15", h, f_bold)
    dt0 = 16
    for i, (v, p) in enumerate(inp.delay_table):
        r = dt0 + i
        put(risk, f"B{r}", v, fill=fill_in)
        put(risk, f"C{r}", p, fmt=PCT, fill=fill_in)
        put(risk, f"D{r}", 0 if i == 0 else f"=D{r - 1}+C{r - 1}", fmt=PCT)
    dt1 = dt0 + len(inp.delay_table) - 1
    put(risk, f"B{dt1 + 1}", "Total")
    put(risk, f"C{dt1 + 1}", f"=SUM(C{dt0}:C{dt1})", f_bold, PCT)

    # Correlation -----------------------------------------------------------
    cor = sheet("Correlation", "Target correlation between drivers and its Cholesky factor (plain formulas)", 22)
    put(cor, "B5", "Correlation (input)", f_bold)
    put(cor, "B13", "Cholesky factor", f_bold)
    a = inp.corr_matrix()
    for i, d in enumerate(inp.drivers):
        put(cor, f"G{6 + i}", d.name)
        put(cor, f"{L(8 + i)}5", d.name[:12], f_bold)
        put(cor, f"G{14 + i}", d.name)
        cor.column_dimensions[L(8 + i)].width = 12
        for j in range(k):
            put(cor, f"{L(8 + j)}{6 + i}", float(a[i, j]), fmt="0.00", fill=fill_in)
    for i in range(k):
        for j in range(k):
            cell = f"{L(8 + j)}{14 + i}"
            A = f"{L(8 + j)}{6 + i}"
            if j > i:
                put(cor, cell, 0, fmt="0.0000")
            elif i == j:
                prev = f"-SUMSQ({L(8)}{14 + i}:{L(8 + j - 1)}{14 + i})" if j > 0 else ""
                put(cor, cell, f"=SQRT({A}{prev})", fmt="0.0000")
            else:
                prev = (f"-SUMPRODUCT({L(8)}{14 + i}:{L(8 + j - 1)}{14 + i},{L(8)}{14 + j}:{L(8 + j - 1)}{14 + j})"
                        if j > 0 else "")
                put(cor, cell, f"=({A}{prev})/{L(8 + j)}{14 + j}", fmt="0.0000")
    lrange = f"$H$14:${L(8 + k - 1)}${14 + k - 1}"
    put(cor, "G21", "Factor failed (matrix not valid)")
    put(cor, "I21", f"=IF(SUMPRODUCT(--ISERROR({lrange}))>0,1,0)", f_check)
    name("Corr_Error", cor, "I21")
    sym = "+".join(f"ABS({L(8 + j)}{6 + i}-{L(8 + i)}{6 + j})" for i in range(k) for j in range(i + 1, k))
    put(cor, "G22", "Not symmetric")
    put(cor, "I22", f"=IF({sym}>0.000001,1,0)", f_check)
    name("Corr_Asym", cor, "I22")

    # Draws -------------------------------------------------------------------
    drw = sheet("Draws", "Seeded uniforms, Latin hypercube, copula and sampled values; calculated once, outside the trial loop", 4)
    put(drw, "B6", "Trial", f_bold)
    parts = ["U", "U2", "Rank", "U LHS", "Z", "Z corr", "U corr", "Value"]
    for j, d in enumerate(inp.drivers):
        c0 = 3 + j * 8
        put(drw, f"{L(c0)}5", d.name, f_bold)
        for p, lab in enumerate(parts):
            put(drw, f"{L(c0 + p)}6", lab, f_bold)
            drw.column_dimensions[L(c0 + p)].width = 9
    zcols = [L(3 + j * 8 + 4) for j in range(k)]
    for i in range(n):
        r = 7 + i
        put(drw, f"B{r}", i + 1 if i == 0 else f"=B{r - 1}+1")
        for j, d in enumerate(inp.drivers):
            c = [L(3 + j * 8 + p) for p in range(8)]
            rr = 7 + j
            put(drw, f"{c[0]}{r}", "=" + _xl_hdr(f"$B{r}", f"Risk!$E${rr}", "MC_Seed"))
            put(drw, f"{c[1]}{r}", "=" + _xl_hdr(f"$B{r}", f"Risk!$E${rr}+1000", "MC_Seed"))
            put(drw, f"{c[2]}{r}", f"=RANK({c[1]}{r},{c[1]}$7:{c[1]}${last},1)")
            put(drw, f"{c[3]}{r}", f"=({c[2]}{r}-1+{c[0]}{r})/MC_Trials")
            put(drw, f"{c[4]}{r}", f"=_xlfn.NORM.S.INV({c[3]}{r})")
            terms = "+".join(f"Correlation!${L(8 + m)}${14 + j}*{zcols[m]}{r}" for m in range(j + 1))
            put(drw, f"{c[5]}{r}", f"=IF(Corr_Error=0,{terms},{c[4]}{r})")
            put(drw, f"{c[6]}{r}", f"=_xlfn.NORM.S.DIST({c[5]}{r},TRUE)")
            u = f"{c[6]}{r}"
            G, H, I = f"Risk!$G${rr}", f"Risk!$H${rr}", f"Risk!$I${rr}"
            if d.dist == "PERT":
                v = f"=_xlfn.BETA.INV({u},Risk!$L${rr},Risk!$M${rr},{G},{I})"
            elif d.dist == "Triangular":
                v = f"=IF({u}<({H}-{G})/({I}-{G}),{G}+SQRT({u}*({I}-{G})*({H}-{G})),{I}-SQRT((1-{u})*({I}-{G})*({I}-{H})))"
            elif d.dist == "Normal":
                v = f"=_xlfn.NORM.INV({u},{G},{H})"
            elif d.dist == "Discrete":
                v = f"=INDEX(Risk!$B${dt0}:$B${dt1},MATCH({u},Risk!$D${dt0}:$D${dt1},1))"
            else:
                v = f"=IF({u}<{G},1,0)"
            put(drw, f"{c[7]}{r}", v)

    # Events register -----------------------------------------------------------
    evs = sheet("Events", "Scenario events register; each event becomes one row of the cash engine", 4)
    ecols = ["ID", "Event", "Type", "Amount", "Start", "End", "Every", "Percent", "Of event", "Lag",
             "Scale by", "Shift by", "Only if", "Start used", "End used", "Scale", "Include", "Source row", "Refers ahead"]
    widths = [5, 24, 12, 9, 6, 6, 6, 8, 8, 5, 18, 15, 13, 8, 8, 7, 7, 8, 8]
    for i, (h, w) in enumerate(zip(ecols, widths)):
        put(evs, f"{L(2 + i)}6", h, f_bold)
        evs.column_dimensions[L(2 + i)].width = w
    ne = len(inp.events)
    e0, e1 = 7, 7 + ne - 1
    drv_names = f"Risk!$C$7:$C${6 + k}"
    drv_live = f"Risk!$K$7:$K${6 + k}"
    for i, ev in enumerate(inp.events):
        r = e0 + i
        vals = [ev.id, ev.label, ev.kind, ev.amount, ev.start, ev.end, ev.every, ev.pct, ev.of, ev.lag,
                ev.scale_by, ev.shift_by, ev.only_if]
        for j, v in enumerate(vals):
            put(evs, f"{L(2 + j)}{r}", v, fill=fill_in, fmt=PCT if j == 7 and v is not None else None)
        put(evs, f"O{r}", f'=F{r}+IF(M{r}="",0,INDEX({drv_live},MATCH(M{r},{drv_names},0)))')
        put(evs, f"P{r}", f'=IF(G{r}="",O{r},G{r}+O{r}-F{r})')
        put(evs, f"Q{r}", f'=1+IF(L{r}="",0,INDEX({drv_live},MATCH(L{r},{drv_names},0)))', fmt="0.000")
        put(evs, f"R{r}", f'=IF(N{r}="",1,INDEX({drv_live},MATCH(N{r},{drv_names},0)))')
        put(evs, f"S{r}", f'=IF(J{r}="",0,MATCH(J{r},$B${e0}:$B${e1},0))')
        put(evs, f"T{r}", f'=IF(D{r}="Percentage",IF(AND(S{r}>0,S{r}<{i + 1}),0,1),0)', f_check)

    # Cash engine ---------------------------------------------------------------
    cash = sheet("Cash", "Monthly one-way cash engine: base cash, events, funding to hold the floor, results")
    lastm = FIRST_MONTH_COL + MONTHS - 1
    J, AG = L(FIRST_MONTH_COL), L(lastm)
    put(cash, "B5", "Month", f_bold)
    put(cash, "I5", "Total", f_bold)
    for m in range(MONTHS):
        put(cash, f"{L(FIRST_MONTH_COL + m)}5", m + 1, f_bold)
        cash.column_dimensions[L(FIRST_MONTH_COL + m)].width = 9
    cash.freeze_panes = "J6"

    def series(row, label, first, rest=None, unit="$", total="sum", style=None):
        put(cash, f"C{row}", label, f_bold if style == "total" else None)
        put(cash, f"H{row}", unit)
        for m in range(MONTHS):
            col = L(FIRST_MONTH_COL + m)
            tpl = first if (m == 0 or rest is None) else rest
            c = put(cash, f"{col}{row}", tpl.replace("@", col).replace("«p»", L(FIRST_MONTH_COL + m - 1) if m else col),
                    fmt=NUM if unit == "$" else "0")
            if style == "total":
                c.font, c.border = f_bold, top
        if total == "sum":
            put(cash, f"I{row}", f"=SUM({J}{row}:{AG}{row})", f_bold if style == "total" else None, NUM)

    heading(cash, 7, "Base cash")
    series(8, "Base receipts", "=Base_Receipts*(1+Drv_Receipts_growth)^(@$5-1)")
    series(9, "Base payments", "=-Base_Payments")
    heading(cash, 10, "Scenario events")
    ev0 = 11
    for i, ev in enumerate(inp.events):
        r, er = ev0 + i, e0 + i
        E = lambda c: f"Events!${c}{er}"  # noqa: E731
        pct_branch = "0"
        if i > 0:
            pct_branch = (f"IF(@$5-{E('K')}>=1,{E('I')}*INDEX(${J}${ev0}:${AG}${r - 1},{E('S')},@$5-{E('K')}),0)")
        f = (f"=IF({E('R')}=0,0,{E('Q')}*IF({E('D')}=\"Bullet\",(@$5={E('O')})*{E('E')},"
             f"IF({E('D')}=\"Repeat\",(@$5>={E('O')})*(@$5<={E('P')})*(MOD(@$5-{E('O')},{E('H')})=0)*{E('E')},"
             f"IF({E('D')}=\"Spread\",(@$5>={E('O')})*(@$5<={E('P')})*{E('E')}/({E('P')}-{E('O')}+1),"
             f"IF({E('D')}=\"Start and end\",(@$5>={E('O')})*(@$5<={E('P')})*{E('E')},"
             f"IF({E('D')}=\"Percentage\",{pct_branch},0))))))")
        series(r, f"={E('C')}", f)
    tr = ev0 + ne
    series(tr, "Total events", f"=SUM(@{ev0}:@{tr - 1})", style="total")
    heading(cash, tr + 1, "Funding")
    R = {key: tr + 2 + i for i, key in enumerate(
        ["open_cash", "open_debt", "interest", "before", "draw", "repay", "above", "close_debt", "close_cash"])}
    series(R["open_cash"], "Opening cash", "=Cash_Opening", f"=«p»{R['close_cash']}", total=None)
    series(R["open_debt"], "Opening debt", "=0", f"=«p»{R['close_debt']}", total=None)
    series(R["interest"], "Interest", f"=-@{R['open_debt']}*Drv_Interest_rate/12")
    series(R["before"], "Cash before funding", f"=@{R['open_cash']}+@8+@9+@{tr}+@{R['interest']}", total=None, style="total")
    series(R["draw"], "Facility draw", f"=MAX(0,Cash_Floor-@{R['before']})")
    series(R["repay"], "Facility repayment", f"=MIN(@{R['open_debt']},MAX(0,@{R['before']}-Cash_Floor))")
    series(R["above"], "Debt above the limit", f"=IF(@{R['close_debt']}>Fac_Limit,1,0)", unit="flag", total=None)
    series(R["close_debt"], "Closing debt", f"=@{R['open_debt']}+@{R['draw']}-@{R['repay']}", total=None, style="total")
    series(R["close_cash"], "Closing cash", f"=@{R['before']}+@{R['draw']}-@{R['repay']}", total=None, style="total")
    res0 = R["close_cash"] + 2
    heading(cash, res0 - 1, "Results")
    bd, cd = f"{J}{R['before']}:{AG}{R['before']}", f"{J}{R['close_debt']}:{AG}{R['close_debt']}"
    results = [("low_cash", "Lowest cash before funding", f"=MIN({bd})", "$"),
               ("low_month", "Month of lowest cash", f"=MATCH(I{res0},{bd},0)", "month"),
               ("peak_debt", "Peak debt", f"=MAX({cd})", "$"),
               ("peak_month", "Month of peak debt", f"=MATCH(I{res0 + 2},{cd},0)", "month"),
               ("interest", "Total interest", f"=-I{R['interest']}", "$"),
               ("net_cash", "Closing net cash", f"={AG}{R['close_cash']}-{AG}{R['close_debt']}", "$"),
               ("breach", "Limit breached", f"=IF(I{res0 + 2}>Fac_Limit,1,0)", "flag"),
               ("first_above", "First month above limit", f"=IFERROR(MATCH(1,{J}{R['above']}:{AG}{R['above']},0),0)", "month")]
    res_row = {}
    for i, (key, label, f, unit) in enumerate(results):
        r = res0 + i
        res_row[key] = r
        put(cash, f"C{r}", label)
        put(cash, f"H{r}", unit)
        put(cash, f"I{r}", f, f_bold, NUM if unit == "$" else "0")

    # Run: the one data table ---------------------------------------------------
    run = sheet("Run", "One data table: each row reruns the whole model with that trial's draws", 4)
    put(run, "B3", "Trial (set by the data table; leave at 0)")
    put(run, "C3", 0, fill=fill_in)
    nout = len(OUTPUTS) + MONTHS
    lastc = L(3 + nout - 1)
    for i, (key, label, unit) in enumerate(OUTPUTS):
        put(run, f"{L(3 + i)}5", label, f_bold)
        put(run, f"{L(3 + i)}6", f"=Cash!$I${res_row[key]}", fmt=NUM if unit == "$" else "0")
    for m in range(MONTHS):
        c = L(3 + len(OUTPUTS) + m)
        put(run, f"{c}5", f"Debt month {m + 1}", f_bold)
        put(run, f"{c}6", f"=Cash!{L(FIRST_MONTH_COL + m)}${R['close_debt']}", fmt=NUM)
    for i in range(nout):
        run.column_dimensions[L(3 + i)].width = 11
    for i in range(n):
        put(run, f"B{7 + i}", i + 1)
    run["C7"] = DataTableFormula(ref=f"C7:{lastc}{last}", r1="C3")

    # Stats ---------------------------------------------------------------------
    st = sheet("Stats", "Distribution report: percentiles, breach probability, fan chart and convergence", 4)
    labels = [("Mean", "AVERAGE({r})"), ("SD", "_xlfn.STDEV.S({r})"), ("Min", "MIN({r})")]
    labels += [(f"P{p}", f"_xlfn.PERCENTILE.INC({{r}},{p / 100})") for p in (5, 10, 25, 50, 75, 90, 95)]
    labels += [("Max", "MAX({r})")]
    for i, (key, label, unit) in enumerate(OUTPUTS):
        col = L(3 + i)
        put(st, f"{col}6", label, f_bold)
        st.column_dimensions[col].width = 12
        rng = f"Run!{col}$7:{col}${last}"
        for j, (lab, f) in enumerate(labels):
            put(st, f"B{7 + j}", lab, f_bold)
            put(st, f"{col}{7 + j}", "=" + f.format(r=rng), fmt=NUM if unit == "$" else "0.00")
    br = 7 + len(labels) + 1
    put(st, f"B{br}", "Probability the facility limit is breached", f_bold)
    put(st, f"I{br}", f"=AVERAGE(Run!I$7:I${last})", f_bold, PCT)
    name("MC_Breach_Prob", st, f"I{br}")
    # Fan chart of closing debt
    fan = br + 3
    put(st, f"B{fan - 1}", "Closing debt by month (fan chart)", f_bold)
    for j, p in enumerate((10, 50, 90)):
        put(st, f"B{fan + j}", f"P{p}", f_bold)
        for m in range(MONTHS):
            rc = L(3 + len(OUTPUTS) + m)
            put(st, f"{L(3 + m)}{fan + j}", f"=_xlfn.PERCENTILE.INC(Run!{rc}$7:{rc}${last},{p / 100})", fmt=NUM)
    for m in range(MONTHS):
        put(st, f"{L(3 + m)}{fan - 1}", m + 1, f_bold)
    # Histogram of peak debt
    hist = fan + 5
    put(st, f"B{hist - 1}", "Peak debt: histogram", f_bold)
    put(st, f"C{hist}", "From", f_bold)
    put(st, f"D{hist}", "To", f_bold)
    put(st, f"E{hist}", "Trials", f_bold)
    bins = 20
    pd_rng = f"Run!$C$7:$C${last}"
    for b in range(bins):
        r = hist + 1 + b
        put(st, f"C{r}", f"=$C$9+({b}/{bins})*($C$17-$C$9)", fmt=NUM)
        put(st, f"D{r}", f"=$C$9+({b + 1}/{bins})*($C$17-$C$9)", fmt=NUM)
        op = "<=" if b == bins - 1 else "<"
        put(st, f"E{r}", f'=COUNTIFS({pd_rng},">="&C{r},{pd_rng},"{op}"&D{r})')
    # Convergence checkpoints
    conv = hist + bins + 3
    put(st, f"B{conv - 1}", "Convergence (peak debt and breach probability)", f_bold)
    for i, h in enumerate(["Trials", "Mean", "Half-width (95%)", "P90", "Breach prob.", "Breach ± (95%)"]):
        put(st, f"{L(3 + i)}{conv}", h, f_bold)
    checkpoints = sorted({max(2, n // 8), max(2, n // 4), max(2, n // 2), n})
    for i, cp in enumerate(checkpoints):
        r = conv + 1 + i
        pr = f"Run!$C$7:INDEX(Run!$C$7:$C${last},C{r})"
        brr = f"Run!$I$7:INDEX(Run!$I$7:$I${last},C{r})"
        put(st, f"C{r}", cp)
        put(st, f"D{r}", f"=AVERAGE({pr})", fmt=NUM)
        put(st, f"E{r}", f"=1.96*_xlfn.STDEV.S({pr})/SQRT(C{r})", fmt=NUM)
        put(st, f"F{r}", f"=_xlfn.PERCENTILE.INC({pr},0.9)", fmt=NUM)
        put(st, f"G{r}", f"=AVERAGE({brr})", fmt=PCT)
        put(st, f"H{r}", f"=1.96*SQRT(G{r}*(1-G{r})/C{r})", fmt=PCT)
    conv_last = conv + len(checkpoints)
    put(st, f"B{conv_last + 1}", "Mean settled (half-width under 2% of the mean)", f_bold)
    put(st, f"I{conv_last + 1}", f"=IF(E{conv_last}<0.02*ABS(D{conv_last}),1,0)", f_bold)
    name("MC_Converged", st, f"I{conv_last + 1}")
    ch = BarChart()
    ch.title, ch.y_axis.title, ch.x_axis.title = "Peak debt: trials by range", "Trials", "Peak debt (upper bound of range)"
    ch.add_data(Reference(st, min_col=5, min_row=hist, max_row=hist + bins), titles_from_data=True)
    ch.set_categories(Reference(st, min_col=4, min_row=hist + 1, max_row=hist + bins))
    ch.height, ch.width, ch.legend = 7, 16, None
    st.add_chart(ch, f"L{hist}")
    lc = LineChart()
    lc.title, lc.y_axis.title, lc.x_axis.title = "Closing debt: P10, P50, P90", "$", "Month"
    lc.add_data(Reference(st, min_col=2, max_col=3 + MONTHS - 1, min_row=fan, max_row=fan + 2), from_rows=True,
                titles_from_data=True)
    lc.set_categories(Reference(st, min_col=3, max_col=3 + MONTHS - 1, min_row=fan - 1))
    lc.height, lc.width = 7, 16
    st.add_chart(lc, f"L{conv - 2}")

    # Sensitivity ---------------------------------------------------------------
    sen = sheet("Sensitivity", "Rank correlation of each driver with peak debt, and its approximate share of the variance", 4)
    for i, h in enumerate(["Driver", "Rank correlation", "Contribution (approx.)"]):
        put(sen, f"{L(2 + i)}6", h, f_bold)
    sen.column_dimensions["B"].width = 22
    sen.column_dimensions["C"].width = 16
    sen.column_dimensions["D"].width = 20
    t0, t1 = 16, 15 + n
    put(sen, "B15", "Trial", f_bold)
    for j, d in enumerate(inp.drivers):
        put(sen, f"B{7 + j}", d.name)
        rc = L(3 + j)
        put(sen, f"C{7 + j}", f"=CORREL({rc}${t0}:{rc}${t1},$H${t0}:$H${t1})", fmt="0.000")
        put(sen, f"D{7 + j}", f"=C{7 + j}^2/SUMSQ($C$7:$C${6 + k})", fmt=PCT)
        put(sen, f"{rc}15", d.name[:12], f_bold)
    put(sen, "H15", "Peak debt", f_bold)
    for i in range(n):
        r = t0 + i
        put(sen, f"B{r}", i + 1)
        for j in range(k):
            vc = L(3 + j * 8 + 7)
            put(sen, f"{L(3 + j)}{r}", f"=_xlfn.RANK.AVG(Draws!{vc}{7 + i},Draws!{vc}$7:{vc}${last},1)")
        put(sen, f"H{r}", f"=_xlfn.RANK.AVG(Run!C{7 + i},Run!C$7:C${last},1)")

    # Checks ----------------------------------------------------------------------
    chk = sheet("Checks", "Error checks must be nil; alerts are for review", 44)
    put(ctl, "C10", "Simulation has run")
    put(ctl, "I10", f"=IF(COUNT(Run!$C$7:$C${last})={n},1,0)")
    name("MC_Ran", ctl, "I10")
    expected = [("Peak debt, trial 1", "=Run!C7", ref["table"][0, 0]),
                ("Peak debt, trial 2", "=Run!C8", ref["table"][1, 0]),
                (f"Peak debt, trial {n}", f"=Run!C{last}", ref["table"][-1, 0]),
                ("Mean peak debt", "=Stats!C7", ref["table"][:, 0].mean()),
                ("P90 peak debt", "=Stats!C15", np.percentile(ref["table"][:, 0], 90)),
                ("Breach probability", "=MC_Breach_Prob", ref["table"][:, 6].mean()),
                ("Base case peak debt", f"=IF(MC_Trial=0,Cash!I{res_row['peak_debt']},H{{r}})", ref["base"]["peak_debt"])]
    errors = [("Correlation factor failed (matrix not valid)", "=Corr_Error"),
              ("Correlation matrix not symmetric", "=Corr_Asym"),
              ("Delay probabilities do not add to 100%", f"=IF(ABS(Risk!C{dt1 + 1}-1)>0.000001,1,0)"),
              ("A percentage event refers to a later event", f"=IF(SUM(Events!T{e0}:T{e1})>0,1,0)"),
              ("Results differ from the build's reference", "=IF(MC_Ran=1,IF(MAX({refs})>0.00001,1,0),0)")]
    alerts = [("Simulation not run (press F9)", "=1-MC_Ran"),
              ("Facility limit breached in the case shown", f"=Cash!I{res_row['breach']}"),
              ("Breach probability above 10%", "=IF(MC_Ran=1,IF(MC_Breach_Prob>0.1,1,0),0)"),
              ("Mean peak debt not settled at this many trials", "=IF(MC_Ran=1,1-MC_Converged,0)")]
    r = 5
    put(chk, f"C{r}", "Error checks", f_bold)
    err0 = r + 1
    err_tot = err0 + len(errors)
    al0 = err_tot + 2
    al_tot = al0 + len(alerts)
    ref0 = al_tot + 3
    refs = f"J{ref0}:J{ref0 + len(expected) - 1}"
    for i, (label, f) in enumerate(errors):
        put(chk, f"C{err0 + i}", label)
        put(chk, f"I{err0 + i}", f.replace("{refs}", refs), f_check)
    put(chk, f"C{err_tot}", "Error checks failing", f_bold)
    put(chk, f"I{err_tot}", f"=SUM(I{err0}:I{err_tot - 1})", f_bold)
    put(chk, f"C{al0 - 1}", "Alerts", f_bold)
    for i, (label, f) in enumerate(alerts):
        put(chk, f"C{al0 + i}", label)
        put(chk, f"I{al0 + i}", f, f_check)
    put(chk, f"C{al_tot}", "Alerts raised", f_bold)
    put(chk, f"I{al_tot}", f"=SUM(I{al0}:I{al_tot - 1})", f_bold)
    name("Chk_Errors", chk, f"I{err_tot}")
    name("Chk_Alerts", chk, f"I{al_tot}")
    put(chk, f"C{ref0 - 1}", "Reference values from the build (Python), compared once the simulation has run", f_bold)
    for i, h in enumerate(["Expected", "Workbook", "Difference"]):
        put(chk, f"{L(8 + i)}{ref0 - 1}", h, f_bold)
    for i, (label, f, v) in enumerate(expected):
        rr = ref0 + i
        put(chk, f"C{rr}", label)
        put(chk, f"H{rr}", float(v), fmt=NUM)
        put(chk, f"I{rr}", f.replace("{r}", str(rr)), fmt=NUM)
        put(chk, f"J{rr}", f"=ABS(I{rr}-H{rr})/MAX(1,ABS(H{rr}))", fmt="0.0E+00")
    chk.column_dimensions["J"].width = 11

    # Contents --------------------------------------------------------------------
    put(contents, "B4", "Press F9 to run the simulation. Calculation is set to automatic except tables, so the data table runs only when asked.", f_bold)
    put(contents, "B5", f"Random draws use the {HDR_SOURCE} counter-based generator, so the same seed gives the same results in any Excel.")
    toc = [("Control", "Seed, trials, inspect trial, cash model settings"),
           ("Risk", "Risk register and live driver values"),
           ("Correlation", "Correlation matrix and Cholesky factor"),
           ("Draws", f"{n} trials x {k} drivers: uniforms, Latin hypercube, copula, values"),
           ("Events", "Scenario events register (bullet, repeat, spread, start and end, percentage)"),
           ("Cash", "Monthly cash engine, funding and results"),
           ("Run", "The data table"),
           ("Stats", "Percentiles, breach probability, histogram, fan chart, convergence"),
           ("Sensitivity", "Rank correlation and contribution to variance"),
           ("Checks", "Error checks, alerts and the reference comparison")]
    for i, (s, d) in enumerate(toc):
        r = 7 + i
        c = put(contents, f"C{r}", s)
        c.hyperlink = Hyperlink(ref=f"C{r}", location=f"'{s}'!A1")
        put(contents, f"H{r}", d)
    put(contents, "C18", "Error checks failing", f_bold)
    put(contents, "I18", "=Chk_Errors", f_check)
    put(contents, "C19", "Alerts raised", f_bold)
    put(contents, "I19", "=Chk_Alerts", f_bold)

    for ws in wb.worksheets:
        for r in range(1, ws.max_row + 1):
            ws.row_dimensions[r].height = 15
    wb.calculation.calcMode = "autoNoTable"
    wb.calculation.fullCalcOnLoad = True
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return {"reference": ref, "rows": {"cash_results": res_row, "close_debt": R["close_debt"], "last": last,
                                       "stats_breach": br, "conv": conv}}


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("out", nargs="?", default=str(Path(__file__).resolve().parents[2] / "build" / "montecarlo" / "mc_demo.xlsx"))
    ap.add_argument("--trials", type=int, default=2000)
    ap.add_argument("--seed", type=int, default=1)
    a = ap.parse_args()
    info = build(Inputs(seed=a.seed), a.trials, Path(a.out))
    ref = info["reference"]
    t = ref["table"]
    print(f"Built {a.out} with {a.trials} trials")
    print(f"Base case peak debt {ref['base']['peak_debt']:,.2f}; simulated mean {t[:, 0].mean():,.2f}, "
          f"P90 {np.percentile(t[:, 0], 90):,.2f}; breach probability {t[:, 6].mean():.1%}")
    print("Contribution to variance of peak debt: " + ", ".join(
        f"{d.name} {c:.0%}" for d, c in zip(Inputs().drivers, ref["contribution"])))


if __name__ == "__main__":
    main()
