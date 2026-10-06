"""Contents, section covers and links, as the package writer writes them (openpyxl).

The same design as the assembly proof's live writer, for whole workbooks:

- a Contents sheet first: model name and status, notes, then a numbered table of contents
  (sections 1, 2, 3 linking to their covers; sheets a., b., c. in each section; and the
  headings on each sheet, marked "-"), then the error and alert totals;
- a cover sheet before each section: its title, "Section N.", the model name, a link to the
  contents and links to the sheets before and after it, and notes;
- on every other sheet, a link to the contents in A1 and to the checks in A2;
- every link is a HYPERLINK formula to a defined name (HL_Home, HL_Err_Chk, HL_Sheet_..,
  HL_Toc_..), and every title in the contents reads the sheet's own title cell, so a rename
  or a move keeps the links and the wording right.

Sections are not stored anywhere but the workbook's order: a cover starts a section and the
sheets after it belong to it until the next cover. So `refresh` after adding, moving or
removing a sheet files the sheet where it sits and rebuilds the contents, the covers'
section numbers and their previous and next links.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.defined_name import DefinedName

TEXT = "404040"
ACCENT = "679DB5"
F_BODY = Font(name="Segoe UI", size=9, color=TEXT)
F_BOLD = Font(name="Segoe UI", size=9, color=TEXT, bold=True)
F_HEAD = Font(name="Segoe UI", size=10, color=TEXT, bold=True)
F_NOTE = Font(name="Segoe UI", size=9, color="808080", italic=True)
F_LINK = Font(name="Segoe UI", size=9, color=TEXT)
F_LINK_BOLD = Font(name="Segoe UI", size=9, color=TEXT, bold=True)
F_NAV = Font(name="Segoe UI Symbol", size=9, color=ACCENT, bold=True)
F_COVER = Font(name="Segoe UI", size=20, color=TEXT, bold=True)
F_CHECK = Font(name="Segoe UI", size=9, color="9C0006", bold=True)
BAR = PatternFill("solid", fgColor="44546A")
F_BAR = Font(name="Segoe UI", size=10, color="FFFFFF", bold=True)
RULE = Border(bottom=Side(style="medium", color=ACCENT))
CONTENTS = "Contents"
HOME, NEXT = "⌂", "›"          # house and single right angle quote, from Segoe UI Symbol


@dataclass
class Cover:
    title: str
    note: str


@dataclass
class Navigation:
    model_name: str
    model_kind: str                                  # shown under the name, e.g. "Budget and actuals model"
    covers: dict[str, Cover]                         # cover sheet name -> section title and note
    notes: list[str] = field(default_factory=list)
    headings: dict[str, list[tuple[int, str]]] = field(default_factory=dict)   # sheet -> (row, text) shown at level 3
    entity: str = "Fictional numbers for illustration"   # shown in B3 of the contents
    checks: str | None = "Checks"                    # sheet the A2 links go to
    errors: str | None = "Chk_Errors"                # names holding the totals, if the model has them
    alerts: str | None = "Chk_Alerts"


def code(sheet: str) -> str:
    return re.sub(r"\W", "_", sheet)


def sheet_name(sheet: str) -> str:
    return f"HL_Sheet_{code(sheet)}"


def ref(sheet: str, cell: str) -> str:
    plain = re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]*", sheet) is not None
    return (sheet if plain else "'" + sheet.replace("'", "''") + "'") + "!" + cell


def link(name: str, text: str) -> str:
    """A HYPERLINK formula to a defined name; text is a formula fragment (a reference or a quoted string)."""
    return f'=HYPERLINK("#{name}",{text})'


def sections(wb, nav: Navigation) -> list[tuple[str | None, list[str]]]:
    """Sections as they stand in the workbook: each cover with the sheets after it."""
    out: list[tuple[str | None, list[str]]] = []
    for ws in wb.worksheets:
        if ws.title == CONTENTS:
            continue
        if ws.title in nav.covers:
            out.append((ws.title, []))
        elif out:
            out[-1][1].append(ws.title)
        else:
            out.append((None, [ws.title]))           # sheets before the first cover sit in an unnamed section
    return out


def arrange(wb, order: list[tuple[str, list[str]]], nav: Navigation) -> None:
    """Create the covers and put the sheets in section order: Contents, then each cover and its sheets."""
    if CONTENTS not in wb.sheetnames:
        wb.create_sheet(CONTENTS, 0)
    for cover, _ in order:
        if cover not in wb.sheetnames:
            wb.create_sheet(cover)
    wanted = [CONTENTS] + [s for cover, sheets in order for s in [cover] + sheets]
    rest = [ws.title for ws in wb.worksheets if ws.title not in wanted]
    wb._sheets = [wb[s] for s in wanted + rest]
    wb.active = 0


def refresh(wb, nav: Navigation) -> None:
    """Rebuild the contents, the covers and every link from the workbook's current order."""
    secs = sections(wb, nav)
    order = [CONTENTS] + [s for cover, sheets in secs for s in ([cover] if cover else []) + sheets]
    for nm in [n for n in list(wb.defined_names) if n.startswith("HL_") or n == "Model_Name"]:
        del wb.defined_names[nm]

    def define(nm, sheet, cell):
        wb.defined_names[nm] = DefinedName(nm, attr_text=_abs(sheet, cell))

    define("HL_Home", CONTENTS, "B1")
    define("Model_Name", CONTENTS, "B1")
    if nav.checks and nav.checks in wb.sheetnames:
        define("HL_Err_Chk", nav.checks, "B1")
    for s in order:
        define(sheet_name(s), s, "A1")

    # Links on every sheet but the contents.
    for s in order[1:]:
        ws = wb[s]
        ws["A1"].hyperlink = ws["A2"].hyperlink = None     # links written as cell hyperlinks give way to the formulas
        ws["A1"] = link("HL_Home", f'"{HOME}"')
        ws["A1"].font = F_NAV
        if nav.checks and nav.checks in wb.sheetnames:
            flag = f'IF({nav.errors}=0,"✓","!")' if nav.errors else '"✓"'
            ws["A2"] = link("HL_Err_Chk", flag)
            ws["A2"].font = F_NAV
        ws.column_dimensions["A"].width = max(ws.column_dimensions["A"].width or 0, 2.5)

    numbered = [(cover, sheets) for cover, sheets in secs if cover]
    for n, (cover, sheets) in enumerate(numbered, start=1):
        i = order.index(cover)
        _cover(wb[cover], nav, n, order[i - 1], order[i + 1] if i + 1 < len(order) else None)
    _contents(wb, nav, secs, define)


def _abs(sheet: str, cell: str) -> str:
    col = "".join(ch for ch in cell if ch.isalpha())
    row = "".join(ch for ch in cell if ch.isdigit())
    return ref(sheet, f"${col}${row}")


def _clear(ws, first_row: int = 1) -> None:
    for row in ws.iter_rows(min_row=first_row, max_row=max(ws.max_row, first_row)):
        for c in row:
            c.value = None
            c.font = F_BODY
            c.fill = PatternFill()
            c.border = Border()


def _cover(ws, nav: Navigation, n: int, prev: str, nxt: str | None) -> None:
    c = nav.covers[ws.title]
    _clear(ws, 3)
    ws.sheet_view.showGridLines = False
    for col, w in (("A", 2.5), ("B", 3), ("C", 3), ("D", 3), ("E", 40), ("F", 12), ("G", 12)):
        ws.column_dimensions[col].width = w
    put = lambda r, v, f=F_BODY: _put(ws, r, 2, v, f)   # noqa: E731
    put(1, c.title, F_HEAD)
    ws["B2"] = "=Model_Name"
    ws["B2"].font = F_BODY
    put(9, c.title, F_COVER)
    ws.row_dimensions[9].height = 30
    for col in range(2, 8):
        ws.cell(9, col).border = RULE
    put(10, f"Section {n}.", F_BOLD)
    ws["B11"] = "=Model_Name"
    ws["B11"].font = F_BODY
    put(12, link("HL_Home", '"Go to contents"'), F_LINK_BOLD)
    shown = '"< Contents"' if prev == CONTENTS else f'"< "&{ref(prev, "$B$1")}'   # the contents' B1 is the model name
    put(13, link(sheet_name(prev), shown), F_LINK)
    if nxt:
        put(14, link(sheet_name(nxt), f'{ref(nxt, "$B$1")}&" >"'), F_LINK)
    put(17, "Section notes", F_BOLD)
    put(18, c.note, F_BODY)
    for r in range(1, 19):
        if r != 9:
            ws.row_dimensions[r].height = 15


def _put(ws, r, col, v, font):
    cell = ws.cell(r, col)
    cell.value = v
    cell.font = font
    return cell


def _contents(wb, nav: Navigation, secs, define) -> None:
    ws = wb[CONTENTS]
    _clear(ws)
    ws.sheet_view.showGridLines = False
    for col, w in (("A", 2.5), ("B", 4), ("C", 4), ("D", 4), ("E", 46), ("F", 10), ("G", 10)):
        ws.column_dimensions[col].width = w
    _put(ws, 1, 2, nav.model_name, F_HEAD)
    if nav.errors:
        _put(ws, 2, 2, f'="{nav.model_kind}"&IF({nav.errors}=0,""," ("&{nav.errors}&" error checks failing)")', F_BODY)
    else:
        _put(ws, 2, 2, nav.model_kind, F_BODY)
    _put(ws, 3, 2, nav.entity, F_NOTE)
    r = 5
    if nav.notes:
        _put(ws, r, 2, "Notes", F_HEAD)
        r += 2
        for k, text in enumerate(nav.notes, start=1):
            _put(ws, r, 2, k, F_BODY).alignment = Alignment(horizontal="left")
            _put(ws, r, 3, text, F_BODY)
            r += 1
        r += 1
    _bar(ws, r, "Table of contents")
    r += 2
    number = 0
    for cover, sheets in secs:
        if cover:
            number += 1
            _put(ws, r, 2, number, F_BOLD).alignment = Alignment(horizontal="left")
            _put(ws, r, 3, link(sheet_name(cover), ref(cover, "$B$1")), F_LINK_BOLD)
            r += 1
        for k, s in enumerate(sheets):
            _put(ws, r, 3, f"{chr(97 + k)}.", F_BODY)
            _put(ws, r, 4, link(sheet_name(s), ref(s, "$B$1")), F_LINK)
            r += 1
            for j, (row, text) in enumerate(nav.headings.get(s, []), start=1):
                nm = f"HL_Toc_{code(s)}_{j}"
                define(nm, s, f"B{row}")
                _put(ws, r, 4, "-", F_BODY).alignment = Alignment(horizontal="right")
                shown = text[1:] if text.startswith("=") else '"' + text.replace('"', '""') + '"'
                _put(ws, r, 5, link(nm, shown), F_LINK)
                r += 1
        r += 1
    if nav.errors:
        _bar(ws, r, "Checks")
        r += 2
        _put(ws, r, 3, link("HL_Err_Chk", '"Error checks failing"'), F_LINK)
        _put(ws, r, 6, f"={nav.errors}", F_CHECK).number_format = "0"
        r += 1
        if nav.alerts:
            _put(ws, r, 3, link("HL_Err_Chk", '"Alerts raised"'), F_LINK)
            _put(ws, r, 6, f"={nav.alerts}", F_BOLD).number_format = "0"
            r += 1
    for rr in range(1, r + 1):
        ws.row_dimensions[rr].height = 15
    ws.freeze_panes = None


def _bar(ws, r: int, text: str) -> None:
    for col in range(2, 8):
        ws.cell(r, col).fill = BAR
    _put(ws, r, 2, text, F_BAR)


def apply(wb, nav: Navigation, order: list[tuple[str, list[str]]]) -> None:
    """Put a built workbook into sections with covers, then write the contents and every link."""
    arrange(wb, order, nav)
    refresh(wb, nav)
