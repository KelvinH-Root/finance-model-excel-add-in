"""Development example: the development test example finished as a model the add-in would hand over.

    HFG_TEMPLATE=/path/to/Budget_Template.xlsx python prototypes/models/development.py [out.xlsx]

Builds the development test example into Kelvin's template (examples/development), then
finishes it the way New model and Remove unused sheets will in Phase 1: keeps the frame sheets
it uses (Time, Lookups) and the development sheets, drops the template's other sheets, adds a
Development summary dashboard with native charts, and writes the contents, section covers
and links (navigation.py). Formulas on kept template sheets that read a removed sheet are
replaced by their values, and names that pointed at removed sheets are dropped.

The template is not in git and neither is the output; the numbers are the test example's
fictional sites.
"""

from __future__ import annotations

import importlib.util
import os
import re
import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.chart import BarChart, LineChart
from openpyxl.styles import Alignment

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(HERE))

import navigation as N  # noqa: E402

RECIPE = ROOT / "examples" / "development" / "recipe.yaml"
COVERS = {
    "Dashboards": N.Cover("Dashboards", "The development summary: key figures for the portfolio and its charts."),
    "Model": N.Cover("Development Model", "Sites, costs, sales and exit, GST, funding, financial statements and returns."),
    "Appendices": N.Cover("Appendices", "The timeline, lookups and the checks."),
}
DASHBOARD = "Development summary"
KEEP_FRAME = ["Time", "Lookups"]


def _reports_module():
    """Chart helpers shared with the report charts proof (styling, finish)."""
    spec = importlib.util.spec_from_file_location("reports_build_dev", ROOT / "prototypes" / "reports" / "build.py")
    mod = importlib.util.module_from_spec(spec)
    sys.modules["reports_build_dev"] = mod
    spec.loader.exec_module(mod)
    return mod


def _sheets_in(text: str) -> set[str]:
    return {a or b for a, b in re.findall(r"'((?:[^']|'')+)'!|([A-Za-z_][\w.]*)!", text or "")}


REBOUND = {"Model_Name", "HL_Home", "HL_Err_Chk"}   # names navigation points at the new contents and checks


def finish_template_sheets(wb, cached, keep: list[str]) -> list[str]:
    """Drop every sheet not kept, with the names that pointed at them; freeze formulas that read them."""
    drop = [s for s in wb.sheetnames if s not in keep]
    bad = set()
    for name in list(wb.defined_names):
        if _sheets_in(wb.defined_names[name].attr_text) & set(drop):
            if name not in REBOUND:
                bad.add(name)
            del wb.defined_names[name]
    token = re.compile(r"[A-Za-z_][A-Za-z0-9_.]*")
    frozen = []
    for s in keep:
        ws = wb[s]
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if not (isinstance(v, str) and v.startswith("=")):
                    continue
                v = re.sub(r"\bCover!\$?B\$?1\b", "Contents!$B$3", v)
                if (_sheets_in(v) & set(drop)) or (set(token.findall(v)) & bad):
                    c.value = cached[s][c.coordinate].value
                    frozen.append(f"{s}!{c.coordinate}")
                else:
                    c.value = v
    for s in drop:
        del wb[s]
    if "Lookups" in wb.sheetnames:
        wb["Lookups"]["B1"] = "Lookups"        # the template's Lookups sheet carries another sheet's title
    return frozen


def dashboard(wb, model, recipe, R):
    """Key figures for the portfolio and four native charts, every number a formula on the development sheets."""
    ws = wb.create_sheet(DASHBOARD)
    ws.sheet_view.showGridLines = False
    put = R.put
    put(ws, 1, 2, DASHBOARD, R.F_HEAD)
    put(ws, 2, 2, "=Model_Name")
    put(ws, 3, 2, "=Contents!$B$3", R.F_NOTE)
    for col, w in (("A", 2.5), ("B", 2.5), ("C", 2.5), ("D", 2.5), ("E", 2.5), ("F", 2.5), ("G", 40), ("H", 14), ("I", 4)):
        ws.column_dimensions[col].width = w
    headings = []

    def section(r, text):
        headings.append((r, text))
        put(ws, r, 2, text, R.F_HEAD, border=R.UNDER)
        for c in range(3, 9):
            ws.cell(r, c).border = R.UNDER

    section(5, "Key figures for the portfolio")
    money, pct = '#,##0.0,,"m";(#,##0.0,,"m");"-"', "0.0%"
    figures = [("Units", 11, "0"), ("Sale revenue (excluding GST)", 12, money), ("Development cost", 13, money),
               ("Finance costs (forecast months)", 14, money), ("Development surplus", 15, money), ("Margin on revenue", 16, pct),
               ("Facility limit", 21, money), ("Peak facility balance", 22, money), ("Month of the peak", 23, "@"),
               ("Peak net equity invested", 24, money), ("Peak loan to cost", 25, pct), ("Equity IRR (per year)", 29, pct),
               ("Project IRR before finance (per year)", 30, pct), ("Equity multiple", 31, '0.00"x"')]
    for k, (label, row, fmt) in enumerate(figures):
        r = 6 + k
        put(ws, r, 3, label)
        put(ws, r, 8, f"=Dev_Returns!$K${row}", R.F_BOLD if k in (1, 4, 11) else R.F_BODY, fmt, align=Alignment(horizontal="right"))

    periods = recipe.periods
    first = 10
    grid_top = 22
    section(grid_top - 1, "Charts")
    data_top = grid_top + 2 * 17 + 1
    section(data_top, "Chart data")
    pos = model.index
    month_row = data_top + 1
    put(ws, month_row, 7, "Month", R.F_BOLD)
    for t in range(periods):
        put(ws, month_row, first + t, f"=Dev_Costs!{R.L(first + t)}$5", R.F_BOLD, align=Alignment(horizontal="right"))
        ws.column_dimensions[R.L(first + t)].width = 9
    lines = [("cost", "Development cost by month", "port.cost.total"), ("cum", "Cumulative development cost", None),
             ("debt", "Facility balance", "port.fund.close"), ("equity", "Equity invested", "port.fund.eq_bal"),
             ("rev", "Settlement revenue", "port.sale.revenue"), ("wip", "Work in progress", "port.fs.wip")]
    rows = {}
    for k, (key, label, src) in enumerate(lines):
        r = month_row + 1 + k
        rows[key] = r
        put(ws, r, 7, label)
        put(ws, r, 8, "$")
        for t in range(periods):
            c = R.L(first + t)
            if src:
                sheet, srow = pos[src]
                f = f"={sheet}!{c}{srow}"
            else:
                f = f"=SUM(${R.L(first)}{rows['cost']}:{c}{rows['cost']})"
            put(ws, r, first + t, f, fmt=money)
    last = first + periods - 1
    cats = R.absref(ws, first, month_row, last, month_row)

    def ser(key):
        r = rows[key]
        return R.series(ws, r, first, last, R.absref(ws, 7, r), cats)

    charts = []

    title_row = month_row + len(lines) + 2
    put(ws, title_row, 7, "Chart titles", R.F_BOLD)

    def place(ch, k, title):
        R.finish_axes(ch, money)
        ch.x_axis.tickLblSkip = 6
        r = title_row + 1 + k
        put(ws, r, 7, title)
        R.finish_chart(ch, R.absref(ws, 7, r), title, legend="b")
        ch.width, ch.height = 15.5, 8
        ch.anchor = f"{'B' if k % 2 == 0 else 'L'}{grid_top + 1 + (k // 2) * 17}"
        ws.add_chart(ch)
        charts.append(title)

    c1 = BarChart()
    c1.type, c1.grouping, c1.gapWidth = "col", "clustered", 40
    s = ser("cost")
    R.style_fill(s, R.PALETTE[0])
    c1.series.append(s)
    place(c1, 0, "Development cost by month")
    c2 = LineChart()
    s = ser("cum")
    R.style_line(s, R.INK, 2.25)
    c2.series.append(s)
    place(c2, 1, "Cumulative development cost (S-curve)")
    c3 = LineChart()
    for key, colour, dash in (("debt", R.INK, None), ("equity", R.PALETTE[0], "dash")):
        s = ser(key)
        R.style_line(s, colour, 2, dash)
        c3.series.append(s)
    place(c3, 2, "Facility balance and equity invested")
    c4 = BarChart()
    c4.type, c4.grouping, c4.gapWidth = "col", "clustered", 40
    s = ser("rev")
    R.style_fill(s, R.PALETTE[2])
    c4.series.append(s)
    ln = LineChart()
    s = ser("wip")
    R.style_line(s, R.INK, 2)
    ln.series.append(s)
    c4 += ln
    place(c4, 3, "Work in progress and settlement revenue")
    for r in range(1, title_row + 6):
        ws.row_dimensions[r].height = 15
    return charts, headings


def build(out: Path, template: str | None = None) -> Path:
    from examples.development.assemble import assemble
    template = template or os.environ.get("HFG_TEMPLATE")
    if not template:
        raise SystemExit("Set HFG_TEMPLATE to the path of Budget_Template.xlsx (it is not kept in git).")
    out = Path(out)
    out.parent.mkdir(parents=True, exist_ok=True)
    base = out.with_name(out.stem + "_template_build.xlsx")
    r, model, _ = assemble(RECIPE, base, template)
    R = _reports_module()
    wb = load_workbook(base)
    cached = load_workbook(base, data_only=True)
    dev = [sl.name for sl in model.sheets]
    finish_template_sheets(wb, cached, KEEP_FRAME + dev)
    charts, dash_headings = dashboard(wb, model, r, R)
    headings = {sl.name: [(row.row, ("=" + row.label[1]) if isinstance(row.label, tuple) else row.label)
                          for row in sl.rows if row.kind == "section"] for sl in model.sheets}
    headings[DASHBOARD] = dash_headings
    nav = N.Navigation(model_name="Demo Development LP", model_kind="Development feasibility and funding model",
                       covers=COVERS, entity="Fictional sites for illustration (the development test example)",
                       notes=["Three fictional sites, monthly over five years, actuals to the last actual month then forecast.",
                              "Costs phased by S-curve, sales through an SPV to the Fund, GST, facility and equity funding.",
                              "Every number on the summary is a formula on the development sheets."],
                       headings=headings, checks="Dev_Checks", errors="Dev_Err_Chk", alerts="Dev_Alt_Chk")
    N.apply(wb, nav, [("Dashboards", [DASHBOARD]), ("Model", [d for d in dev if d != "Dev_Checks"]),
                      ("Appendices", KEEP_FRAME + ["Dev_Checks"])])
    wb.save(out)
    R.finish(out, {DASHBOARD: [f"D{k + 1} {t}" for k, t in enumerate(charts)]})
    base.unlink()
    return out


if __name__ == "__main__":
    dst = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "build" / "models" / "development_example.xlsx"
    print(build(dst))
