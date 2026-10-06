"""Phase 0 proof: native, live report charts and controls written by the package writer.

Builds a workbook with fictional numbers that shows what the Build tab's Charts and
Controls commands will produce, as ordinary Excel objects driven by formulas:

- a Z chart (monthly actual and forecast bars, budget bars behind, cumulative actual then
  forecast, cumulative budget, moving annual total), with the line chosen by a classic
  form control drop-down (combo box) linked to a named cell;
- two waterfalls (P&L walk, budget to actual) built as stacked columns, so totals,
  increases and decreases are coloured by formula and stay right when numbers change;
- IBCS column and variance charts (prior year grey, plan outlined, actual solid,
  forecast hatched; good variances green, bad red), with the comparison chosen by an
  in-cell drop-down;
- a contents sheet, a Lookups sheet with the lists, and a Checks sheet.

    python prototypes/charts/build.py [out.xlsx]

Nothing here needs the add-in to calculate; the charts update when the numbers or the
selections change. Form controls cannot be created through Office.js, so the package
writer writes them as file parts (a control properties part, a VML shape and the sheet's
controls element), the same parts Excel writes.
"""

from __future__ import annotations

import os
import re
import shutil
import sys
import tempfile
import zipfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference, ScatterChart, Series
from openpyxl.chart.data_source import StrData, StrRef, StrVal
from openpyxl.chart.label import DataLabel, DataLabelList
from openpyxl.chart.marker import Marker
from openpyxl.chart.series import SeriesLabel
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText, Text
from openpyxl.chart.title import Title
from openpyxl.chart.legend import LegendEntry
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.drawing.fill import ColorChoice, PatternFillProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.text import CharacterProperties, Paragraph, ParagraphProperties
from openpyxl.drawing.text import Font as DrawingFont
from openpyxl.styles import Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.hyperlink import Hyperlink

INK, GREY, LIGHT, GOOD, BAD = "262626", "A6A6A6", "D9D9D9", "8CB400", "FF0000"
TEXT = "404040"
F_BODY = Font(name="Segoe UI", size=9, color=TEXT)
F_BOLD = Font(name="Segoe UI", size=9, color=TEXT, bold=True)
F_HEAD = Font(name="Segoe UI", size=10, color=TEXT, bold=True)
F_CHECK = Font(name="Segoe UI", size=9, color="9C0006")
FILL_IN = PatternFill("solid", fgColor="FFF2CC")
TOP = Border(top=Side(style="thin", color=TEXT))
NUM = '#,##0;(#,##0);"-"'

MONTHS = ["Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec", "Jan", "Feb", "Mar"]
LAST_ACTUAL = 6  # September
LINES = {
    # line: (prior year monthly, actual or forecast monthly, budget monthly)
    # Revenue reproduces the template's Z-Chart sheet: the same actuals and forecast, and a
    # budget that passes through every labelled point of its budget line (6, 27, 38, 46, 66, 76).
    "Revenue": ([6, 6, 7, 6, 6, 4, 6, 6, 5, 5, 3, 2],
                [8, 8, 13, 7, 6, 2, 6, 9, 7, 7, 4, 5],
                [6, 6, 7, 8, 6, 5, 8, 7, 6, 7, 5, 5]),
    "Cost of sales": ([3, 4, 4, 4, 4, 4, 4, 3, 3, 4, 4, 4],
                      [5, 5, 7, 4, 4, 2, 4, 5, 4, 4, 3, 3],
                      [4, 4, 3, 6, 5, 2, 5, 4, 4, 4, 3, 4]),
}
COST_LINES = {"Cost of sales"}   # a fall against budget is good


def style_sheet(ws, title: str, purpose: str):
    ws.sheet_view.showGridLines = False
    ws["B1"], ws["B2"] = title, purpose
    ws["B1"].font, ws["B2"].font = F_HEAD, F_BODY
    ws.column_dimensions["A"].width = 2.5


def put(ws, ref, value, font=None, fmt=None, fill=None, border=None):
    c = ws[ref]
    c.value = value
    c.font = font or F_BODY
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if border:
        c.border = border
    return c


def name(wb, nm, ws, ref):
    col = "".join(ch for ch in ref if ch.isalpha())
    row = "".join(ch for ch in ref if ch.isdigit())
    wb.defined_names[nm] = DefinedName(nm, attr_text=f"'{ws.title}'!${col}${row}")


def no_line():
    return LineProperties(noFill=True)


# Z chart --------------------------------------------------------------------------------
# Look taken from the template's Z-Chart picture: black actual bars, hatched forecast bars,
# grey budget bars set back and to the left, labelled lines instead of a legend, a divider
# between actual and forecast, AC and FC under the months, and year to date and full year
# variance boxes on the right.
Z_INK, Z_GREY, Z_BAR, Z_GOOD, Z_BAD = "000000", "808080", "D0D0D0", "008F00", "C00000"
Z_FONT = "Segoe UI"
Z_W, Z_H = 30.0, 15.0            # chart size, cm
Z_PLOT = (0.045, 0.075, 0.95, 0.785)   # plot area x, y, width, height as shares of the chart
Z_CATS = 15                      # twelve months and three spare places for the variance boxes
Z_YTD_X, Z_FY_X = 13.8, 14.5     # box positions on the category scale (1 is the first month)
Z_BOX = 1 / 3                    # box width as a share of one month
Z_QUARTERS = (0, 3, 6, 9)        # points labelled on the lines: first month of each quarter
Z_HELPER_NAMES = {
    "divider": "Divider, actual to forecast", "ac": "AC label", "fc": "FC label",
    "ytd_level": "YTD actual reference line", "ytd_bud_level": "YTD budget reference line",
    "fy_level": "FY reference line", "fy_bud_level": "FY budget reference line",
    "ytd_good": "YTD box, good", "ytd_bad": "YTD box, bad", "ytd_edge": "YTD box outline",
    "fy_good": "FY box, good", "fy_bad": "FY box, bad", "fy_edge": "FY box outline",
    "ytd_name": "YTD Δ label", "ytd_var_good": "YTD variance label, good", "ytd_var_bad": "YTD variance label, bad",
    "fy_name": "FY Δ label", "fy_var_good": "FY variance label, good", "fy_var_bad": "FY variance label, bad",
    "mat_name": "MAT label", "acfc_name": "AC+FC label", "bud_name": "BUD label", "scale": "Scale (keeps both axes alike)",
}


def z_text(size=8, bold=False, colour=Z_INK):
    props = CharacterProperties(sz=int(size * 100), b=bold, solidFill=colour, latin=DrawingFont(typeface=Z_FONT))
    return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=props), r=[], endParaRPr=CharacterProperties())])


def z_labels(points=None, pos="t", size=8, bold=False, colour=Z_INK, name=False, fmt="0"):
    """Data labels on every point, or only on the listed points; value or series name."""
    dl = DataLabelList()
    dl.numFmt = fmt
    dl.txPr = z_text(size, bold, colour)
    flags = dict(showLegendKey=False, showCatName=False, showPercent=False, showBubbleSize=False)
    if points is None:
        dl.showVal, dl.showSerName, dl.position = not name, name, pos
    else:
        dl.showVal = dl.showSerName = False
        dl.dLbl = [DataLabel(idx=i, numFmt=fmt, txPr=z_text(size, bold, colour), dLblPos=pos,
                             showVal=not name, showSerName=name, **flags) for i in points]
    for k, v in flags.items():
        setattr(dl, k, v)
    return dl


def z_line(ser, colour=Z_INK, width=1.25, dash=None, marker=True, cap=None):
    if ser.graphicalProperties is None:
        ser.graphicalProperties = GraphicalProperties()
    ser.graphicalProperties.line = LineProperties(solidFill=colour, w=int(width * 12700), prstDash=dash, cap=cap)
    ser.marker = Marker(symbol="circle" if marker else "none")
    if marker:
        ser.marker.size = 4
        ser.marker.graphicalProperties = GraphicalProperties(solidFill=colour, ln=LineProperties(solidFill=colour))
    ser.smooth = False
    return ser


def z_hidden(ser):
    if ser.graphicalProperties is None:
        ser.graphicalProperties = GraphicalProperties()
    ser.graphicalProperties.line = no_line()
    ser.marker = Marker(symbol="none")
    ser.smooth = False
    return ser


def z_chart_sheet(wb):
    zs = wb.create_sheet("Z chart")
    style_sheet(zs, "Z chart", "Monthly bars, year to date and moving annual total, actual then forecast against budget")
    put(zs, "B4", "Line", F_BOLD)
    put(zs, "D4", "=INDEX(LU_Chart_Lines,DD_Chart_Line)", F_HEAD)
    put(zs, "H4", "Cost line", F_BOLD)
    put(zs, "J4", "=INDEX(LU_Chart_Cost,DD_Chart_Line)", fmt="0")
    put(zs, "B5", "Chart title", F_BOLD)
    put(zs, "D5", '=D4&" Bridge (Z-Chart)"')

    # Monthly table: rows 8 to 19 are the months, 20 to 22 the spare places on the right.
    top, first, last = 7, 8, 7 + Z_CATS
    heads = ["Month", "AC", "FC", "Budget", "Offset", "YTD AC", "YTD FC", "YTD budget", "MAT", "AC+FC", "Scale"]
    for i, h in enumerate(heads):
        put(zs, f"{L(2 + i)}{top}", h, F_BOLD)
    sel = "(DD_Chart_Line-1)*5"
    row_of = lambda k: f"INDEX(Data!$J$1:$U$40,{k}+{sel},0)"          # noqa: E731  whole row of the selected line
    to_date = lambda k, m: f"SUM(INDEX(Data!$J$1:$U$40,{k}+{sel},1):INDEX(Data!$J$1:$U$40,{k}+{sel},{m}))"  # noqa: E731
    for m in range(12):
        rr, col = first + m, L(10 + m)
        acfc = f"INDEX(Data!${col}$1:${col}$40,12+{sel})"
        put(zs, f"B{rr}", MONTHS[m])
        put(zs, f"C{rr}", f"=IF({m + 1}<=Ts_Last_Actual,{acfc},NA())", fmt=NUM)
        put(zs, f"D{rr}", f"=IF({m + 1}>Ts_Last_Actual,{acfc},NA())", fmt=NUM)
        put(zs, f"E{rr}", f"=INDEX(Data!${col}$1:${col}$40,13+{sel})", fmt=NUM)
        put(zs, f"G{rr}", f"=IF({m + 1}<=Ts_Last_Actual,K{rr},NA())", fmt=NUM)
        put(zs, f"H{rr}", f"=IF({m + 1}>=Ts_Last_Actual,K{rr},NA())", fmt=NUM)
        put(zs, f"I{rr}", f"={to_date(13, m + 1)}", fmt=NUM)
        # MAT: this year's months to date plus the prior year's remaining months
        py_rest = f"+SUM(INDEX(Data!$J$1:$U$40,11+{sel},{m + 2}):INDEX(Data!$J$1:$U$40,11+{sel},12))" if m < 11 else ""
        put(zs, f"J{rr}", f"={to_date(12, m + 1)}{py_rest}", fmt=NUM)
        put(zs, f"K{rr}", f"={to_date(12, m + 1)}", fmt=NUM)
    m12 = first + 11

    # Totals, variances and the values that place the helpers.
    s = 24
    summary = [
        ("Year to date actual", f"=IF(Ts_Last_Actual>=1,INDEX(K{first}:K{m12},Ts_Last_Actual),0)", NUM),
        ("Year to date budget", f"=IF(Ts_Last_Actual>=1,INDEX(I{first}:I{m12},Ts_Last_Actual),0)", NUM),
        ("YTD variance to budget", f"=F{s}-F{s + 1}", '+#,##0;-#,##0;"-"'),
        ("Full year actual then forecast", f"=K{m12}", NUM),
        ("Full year budget", f"=I{m12}", NUM),
        ("Full year variance to budget", f"=F{s + 3}-F{s + 4}", '+#,##0;-#,##0;"-"'),
        ("Good direction (1 up, -1 down)", "=1-2*J4", "0"),
        ("Highest value plotted", f"=MAX(K{first}:K{m12},I{first}:I{m12},J{first}:J{m12},E{first}:E{m12},{row_of(12)})", NUM),
        ("Top of the divider", f"=F{s + 7}", "0.0"),
        ("YTD variance label", f'=TEXT(F{s + 2},"+0;-0;0")', None),
        ("Full year variance label", f'=TEXT(F{s + 5},"+0;-0;0")', None),
    ]
    for i, (label, f, fmt) in enumerate(summary):
        put(zs, f"B{s + i}", label, F_BOLD if "variance to" in label else F_BODY)
        put(zs, f"F{s + i}", f, F_BOLD if "variance to" in label else F_BODY, fmt)
    ytd_ac, ytd_bud, ytd_var, fy, fy_bud, fy_var, good_dir, high, top_v, ytd_txt, fy_txt = (f"$F${s + i}" for i in range(len(summary)))
    put(zs, f"L{first}", f"={top_v}", fmt="0.0")

    # Helper points for the scatter series: up to five points each (x in C:G, y in H:L, name in M).
    h0 = s + len(summary) + 2
    for i, head in enumerate(["Helper"] + ["x"] * 5 + ["y"] * 5 + ["Name"]):
        put(zs, f"{L(2 + i)}{h0}", head, F_BOLD)
    has_ac, has_fc = "Ts_Last_Actual>=1", "Ts_Last_Actual<=11"
    ytd_good, fy_good = f"{good_dir}*{ytd_var}>=0", f"{good_dir}*{fy_var}>=0"
    lo = lambda a, b: f"MIN({a},{b})"   # noqa: E731
    hi = lambda a, b: f"MAX({a},{b})"   # noqa: E731
    mid = lambda a, b: f"({a}+{b})/2"   # noqa: E731
    na_unless = lambda cond, v: f"=IF({cond},{v},NA())"   # noqa: E731

    def outline(x, a, b):
        """Five points round a box centred on x, from a to b."""
        x0, x1 = round(x - Z_BOX / 2, 4), round(x + Z_BOX / 2, 4)
        return [x0, x1, x1, x0, x0], [lo(a, b), lo(a, b), hi(a, b), hi(a, b), lo(a, b)]

    def box(x, a, b):
        return [x, x], [lo(a, b), hi(a, b)]

    gap = round(Z_BOX / 2 + 0.04, 4)
    show_ytd = f"AND({has_ac},{has_fc})"          # with a full year of actuals the YTD box would repeat the FY box
    helpers = [
        # key, shown when, (x values, y values), name; every value turns to #N/A when hidden
        ("divider", show_ytd, (["Ts_Last_Actual+0.5"] * 2, [0, top_v]), None),
        ("ac", has_ac, (["(1+Ts_Last_Actual)/2"], [0]), '=CHAR(10)&"AC"'),
        ("fc", has_fc, (["(Ts_Last_Actual+13)/2"], [0]), '=CHAR(10)&"FC"'),
        ("ytd_level", show_ytd, (["Ts_Last_Actual", Z_YTD_X], [ytd_ac] * 2), None),
        ("ytd_bud_level", show_ytd, (["Ts_Last_Actual", Z_YTD_X], [ytd_bud] * 2), None),
        ("fy_level", "TRUE", ([12, Z_FY_X], [fy] * 2), None),
        ("fy_bud_level", "TRUE", ([12, Z_FY_X], [fy_bud] * 2), None),
        ("ytd_good", f"AND({show_ytd},{ytd_good})", box(Z_YTD_X, ytd_ac, ytd_bud), None),
        ("ytd_bad", f"AND({show_ytd},NOT({ytd_good}))", box(Z_YTD_X, ytd_ac, ytd_bud), None),
        ("ytd_edge", show_ytd, outline(Z_YTD_X, ytd_ac, ytd_bud), None),
        ("fy_good", fy_good, box(Z_FY_X, fy, fy_bud), None),
        ("fy_bad", f"NOT({fy_good})", box(Z_FY_X, fy, fy_bud), None),
        ("fy_edge", "TRUE", outline(Z_FY_X, fy, fy_bud), None),
        ("ytd_name", show_ytd, ([Z_YTD_X - gap], [mid(ytd_ac, ytd_bud)]), "YTD Δ"),
        ("ytd_var_good", f"AND({show_ytd},{ytd_good})", ([Z_YTD_X + gap], [mid(ytd_ac, ytd_bud)]), f"={ytd_txt}"),
        ("ytd_var_bad", f"AND({show_ytd},NOT({ytd_good}))", ([Z_YTD_X + gap], [mid(ytd_ac, ytd_bud)]), f"={ytd_txt}"),
        ("fy_name", "TRUE", ([Z_FY_X - gap], [mid(fy, fy_bud)]), "FY Δ"),
        ("fy_var_good", fy_good, ([Z_FY_X + gap], [mid(fy, fy_bud)]), f"={fy_txt}"),
        ("fy_var_bad", f"NOT({fy_good})", ([Z_FY_X + gap], [mid(fy, fy_bud)]), f"={fy_txt}"),
        ("mat_name", "TRUE", ([1], [f"J{first}"]), "MAT"),
        ("acfc_name", "TRUE", ([12], [f"K{m12}"]), "AC+FC"),
        ("bud_name", "TRUE", ([12], [f"I{m12}"]), "BUD"),
        ("scale", "TRUE", ([1], [top_v]), None),
    ]
    cell = lambda cond, v: f"={v}" if cond == "TRUE" else na_unless(cond, v)   # noqa: E731
    helpers = [(key, ([cell(c, x) for x in xs], [cell(c, y) for y in ys]), nm) for key, c, (xs, ys), nm in helpers]
    rows, sizes = {}, {}
    for i, (key, (xs, ys), nm) in enumerate(helpers):
        rr = h0 + 1 + i
        rows[key], sizes[key] = rr, len(xs)
        put(zs, f"B{rr}", Z_HELPER_NAMES[key])
        for j, v in enumerate(xs):
            put(zs, f"{L(3 + j)}{rr}", v, fmt="0.00")
        for j, v in enumerate(ys):
            put(zs, f"{L(8 + j)}{rr}", v, fmt="0.0")
        if nm is not None:
            put(zs, f"M{rr}", nm)
    for c, w in zip("BCDEFGHIJKLM", [26, 7, 7, 8, 8, 8, 8, 10, 7, 7, 7, 8]):
        zs.column_dimensions[c].width = w

    cats = Reference(zs, min_col=2, min_row=first, max_row=last)
    plot_w_pt = Z_W / 2.54 * 72 * Z_PLOT[2]
    box_pt = plot_w_pt / Z_CATS * Z_BOX

    # Primary axes: budget bars set back and to the left, then the helpers.
    bud = BarChart()
    bud.type, bud.grouping, bud.overlap, bud.gapWidth = "col", "clustered", 50, 17
    bud.add_data(Reference(zs, min_col=5, max_col=6, min_row=top, max_row=last), titles_from_data=True)
    bud.set_categories(cats)
    b_ser, offset = bud.series
    b_ser.graphicalProperties.solidFill = Z_BAR
    b_ser.graphicalProperties.line = no_line()
    offset.graphicalProperties.noFill = True
    offset.graphicalProperties.line = no_line()

    sc = ScatterChart()
    sc.scatterStyle = "lineMarker"
    sc.x_axis.axId, sc.y_axis.axId = bud.x_axis.axId, bud.y_axis.axId    # share the category axis: x 1 is the first month

    def helper(key, title=None):
        rr, n = rows[key], sizes[key]
        ser = Series(Reference(zs, min_col=8, max_col=7 + n, min_row=rr), Reference(zs, min_col=3, max_col=2 + n, min_row=rr))
        if title == "cell":
            ser.tx = SeriesLabel(strRef=StrRef(f"'Z chart'!$M${rr}"))
        elif title:
            ser.tx = SeriesLabel(v=title)
        sc.series.append(ser)
        return ser

    # Excel draws scatter series in order, so each box's outline comes after its fill.
    z_line(helper("divider"), width=1, marker=False)
    for key, colour in (("ytd_level", Z_INK), ("ytd_bud_level", Z_GREY), ("fy_level", Z_INK), ("fy_bud_level", Z_GREY)):
        z_line(helper(key), colour, width=1, dash="sysDot", marker=False)
    for key, colour, width, cap in (("ytd_good", Z_GOOD, box_pt, "flat"), ("ytd_bad", Z_BAD, box_pt, "flat"), ("ytd_edge", Z_INK, 0.75, None),
                                    ("fy_good", Z_GOOD, box_pt, "flat"), ("fy_bad", Z_BAD, box_pt, "flat"), ("fy_edge", Z_INK, 0.75, None)):
        z_line(helper(key), colour, width=width, marker=False, cap=cap)
    for key, pos, size, bold, colour in (
            ("ac", "b", 10, True, Z_INK), ("fc", "b", 10, True, Z_INK),
            ("ytd_name", "l", 8, False, Z_GREY), ("fy_name", "l", 8, False, Z_GREY),
            ("ytd_var_good", "r", 9, True, Z_GOOD), ("ytd_var_bad", "r", 9, True, Z_BAD),
            ("fy_var_good", "r", 9, True, Z_GOOD), ("fy_var_bad", "r", 9, True, Z_BAD),
            ("mat_name", "l", 8, True, Z_INK), ("acfc_name", "l", 8, True, Z_INK), ("bud_name", "b", 8, True, Z_GREY)):
        ser = z_hidden(helper(key, "cell"))
        ser.dLbls = z_labels(pos=pos, size=size, bold=bold, colour=colour, name=True)
    z_hidden(helper("scale", "Scale"))

    # Secondary axes, drawn on top: actual and forecast bars, then the lines.
    acfc = BarChart()
    acfc.type, acfc.grouping, acfc.overlap, acfc.gapWidth = "col", "clustered", 100, 100
    acfc.add_data(Reference(zs, min_col=3, max_col=4, min_row=top, max_row=last), titles_from_data=True)
    acfc.set_categories(cats)
    ac, fc = acfc.series
    ac.graphicalProperties.solidFill = Z_INK
    ac.graphicalProperties.line = no_line()
    fc.graphicalProperties.pattFill = PatternFillProperties(prst="upDiag", fgClr=ColorChoice(srgbClr=Z_INK), bgClr=ColorChoice(srgbClr="FFFFFF"))
    fc.graphicalProperties.line = LineProperties(solidFill=Z_INK, w=9525)
    for ser in (ac, fc):
        ser.dLbls = z_labels(pos="outEnd", bold=True)

    lines = LineChart()
    lines.add_data(Reference(zs, min_col=7, max_col=12, min_row=top, max_row=last), titles_from_data=True)
    lines.set_categories(cats)
    l_ac, l_fc, l_bud, l_mat, l_cum, l_scale = lines.series
    z_line(l_ac)
    z_line(l_fc, dash="dash")
    z_line(l_bud, Z_GREY)
    z_line(l_mat)
    z_hidden(l_cum)
    z_hidden(l_scale)
    l_cum.dLbls = z_labels(Z_QUARTERS, "t")
    l_bud.dLbls = z_labels(Z_QUARTERS, "b", colour=Z_GREY)
    l_mat.dLbls = z_labels(Z_QUARTERS, "t", bold=True)
    for ch in (acfc, lines):
        ch.x_axis.axId, ch.y_axis.axId = 300, 200
    acfc.x_axis.crossAx, acfc.y_axis.crossAx = 200, 300
    acfc.x_axis.delete = True
    acfc.y_axis.delete = True
    acfc.y_axis.crosses = "max"
    acfc.y_axis.scaling.min = 0

    z = bud
    z += sc
    z += acfc
    z += lines
    z.legend = None
    title_ref = StrRef("'Z chart'!$D$5")
    title_ref.strCache = StrData(ptCount=1, pt=[StrVal(idx=0, v=f"{list(LINES)[0]} Bridge (Z-Chart)")])
    z.title = Title(tx=Text(strRef=title_ref), overlay=False, txPr=z_text(11, True),
                    layout=Layout(manualLayout=ManualLayout(xMode="edge", yMode="edge", x=0.005, y=0.01)))
    z.y_axis.delete = True
    z.y_axis.majorGridlines = None
    z.y_axis.scaling.min = 0
    z.x_axis.delete = False
    z.x_axis.majorTickMark = "none"
    z.x_axis.minorTickMark = "none"
    z.x_axis.tickLblPos = "nextTo"
    z.x_axis.txPr = z_text(9)
    z.x_axis.spPr = GraphicalProperties(ln=LineProperties(solidFill=Z_INK, w=9525))
    z.plot_area.layout = Layout(manualLayout=ManualLayout(layoutTarget="inner", xMode="edge", yMode="edge",
                                                          x=Z_PLOT[0], y=Z_PLOT[1], w=Z_PLOT[2], h=Z_PLOT[3]))
    z.graphical_properties = GraphicalProperties(ln=LineProperties(noFill=True))
    z.width, z.height = Z_W, Z_H
    zs.add_chart(z, "N4")
    return zs


def build(path: Path) -> Path:
    wb = Workbook()
    contents = wb.active
    contents.title = "Contents"
    style_sheet(contents, "Contents", "Charts and controls the Build tab will insert (fictional numbers)")

    # Lookups --------------------------------------------------------------------
    lk = wb.create_sheet("Lookups")
    style_sheet(lk, "Lookups", "Lists behind the drop-downs; a list made from a category set grows with it")
    put(lk, "B4", "Chart lines", F_BOLD)
    names = list(LINES) + ["Gross profit"]
    for i, n in enumerate(names):
        put(lk, f"B{5 + i}", n)
    wb.defined_names["LU_Chart_Lines"] = DefinedName("LU_Chart_Lines", attr_text=f"Lookups!$B$5:$B${4 + len(names)}")
    put(lk, "C4", "Cost line", F_BOLD)
    for i, n in enumerate(names):
        put(lk, f"C{5 + i}", 1 if n in COST_LINES else 0, fmt="0")
    wb.defined_names["LU_Chart_Cost"] = DefinedName("LU_Chart_Cost", attr_text=f"Lookups!$C$5:$C${4 + len(names)}")
    put(lk, "D4", "Compare with", F_BOLD)
    put(lk, "D5", "Plan")
    put(lk, "D6", "Prior year")
    wb.defined_names["LU_Compare"] = DefinedName("LU_Compare", attr_text="Lookups!$D$5:$D$6")
    lk.column_dimensions["B"].width = 18
    lk.column_dimensions["C"].width = 10
    lk.column_dimensions["D"].width = 14

    # Data ------------------------------------------------------------------------
    data = wb.create_sheet("Data")
    style_sheet(data, "Data", "Monthly prior year, actual then forecast, and budget for each line")
    put(data, "B4", "Line shown in the Z chart", F_BOLD)
    put(data, "H4", 1, fmt="0", fill=FILL_IN)
    name(wb, "DD_Chart_Line", data, "H4")
    put(data, "I4", "=INDEX(LU_Chart_Lines,DD_Chart_Line)", F_BOLD)
    put(data, "B5", "Last actual month", F_BOLD)
    put(data, "H5", LAST_ACTUAL, fmt="0", fill=FILL_IN)
    name(wb, "Ts_Last_Actual", data, "H5")
    put(data, "I5", "=INDEX($J$8:$U$8,Ts_Last_Actual)", F_BODY)
    put(data, "B8", "Month", F_BOLD)
    for m, mon in enumerate(MONTHS):
        put(data, f"{L(10 + m)}8", mon, F_BOLD)
    put(data, "H8", "Total", F_BOLD)
    r = 10
    blocks = {}
    for line in names:
        put(data, f"B{r}", line, F_HEAD)
        rows = {}
        for k, label in (("py", "Prior year"), ("acfc", "Actual then forecast"), ("bud", "Budget")):
            rr = r + 1 + len(rows)
            rows[k] = rr
            put(data, f"C{rr}", label)
            for m in range(12):
                cell = f"{L(10 + m)}{rr}"
                if line == "Gross profit":
                    a, b = blocks["Revenue"][k], blocks["Cost of sales"][k]
                    put(data, cell, f"={L(10 + m)}{a}-{L(10 + m)}{b}", fmt=NUM)
                else:
                    put(data, cell, LINES[line][("py", "acfc", "bud").index(k)][m], fmt=NUM, fill=FILL_IN)
            put(data, f"H{rr}", f"=SUM(J{rr}:U{rr})", F_BOLD, NUM)
        blocks[line] = rows
        r += 5
    data.column_dimensions["B"].width = 2.5
    data.column_dimensions["C"].width = 22
    for m in range(12):
        data.column_dimensions[L(10 + m)].width = 7

    # Classic form control: a combo box linked to DD_Chart_Line, list LU_Chart_Lines.
    # Written into the file after openpyxl saves (see add_form_control).

    zs = z_chart_sheet(wb)

    # Waterfalls --------------------------------------------------------------------
    wf = wb.create_sheet("Waterfall")
    style_sheet(wf, "Waterfall", "Bridges built as stacked columns: totals grey, increases green, decreases red")
    walk = [("Revenue", 1000, "total"), ("Materials", -350, "step"), ("Labour", -200, "step"), ("Other direct", -50, "step"),
            ("Gross profit", None, "total"), ("Staff", -120, "step"), ("Marketing", -45, "step"), ("Admin", -35, "step"),
            ("Other opex", -20, "step"), ("EBITDA", None, "total"), ("D&A", -30, "step"), ("Interest", -15, "step"),
            ("Tax", -40, "step"), ("Net profit", None, "total")]
    bva = [("Budget", 500, "total"), ("Revenue Δ", 85, "step"), ("COGS Δ", -30, "step"), ("Opex Δ", -15, "step"),
           ("Other Δ", 10, "step"), ("Actual", None, "total")]

    def waterfall(top_row, title, items, anchor):
        heads = ["Item", "Value", "Running", "Base", "Total", "Increase", "Decrease", "Label"]
        put(wf, f"B{top_row}", title, F_HEAD)
        for i, h in enumerate(heads):
            put(wf, f"{L(2 + i)}{top_row + 1}", h, F_BOLD)
        first = top_row + 2
        for k, (label, value, kind) in enumerate(items):
            rr = first + k
            put(wf, f"B{rr}", label)
            if kind == "total" and value is None:
                put(wf, f"C{rr}", f"=D{rr - 1}", fmt=NUM)          # subtotal = running total so far
                put(wf, f"D{rr}", f"=C{rr}", fmt=NUM)
            elif kind == "total":
                put(wf, f"C{rr}", value, fmt=NUM, fill=FILL_IN)
                put(wf, f"D{rr}", f"=C{rr}", fmt=NUM)
            else:
                put(wf, f"C{rr}", value, fmt=NUM, fill=FILL_IN)
                put(wf, f"D{rr}", f"=D{rr - 1}+C{rr}", fmt=NUM)
            is_total = kind == "total"
            put(wf, f"E{rr}", "=0" if is_total else f"=MIN(D{rr - 1},D{rr})", fmt=NUM)
            put(wf, f"F{rr}", f"=C{rr}" if is_total else "=0", fmt=NUM)
            put(wf, f"G{rr}", "=0" if is_total else f"=MAX(C{rr},0)", fmt=NUM)
            put(wf, f"H{rr}", "=0" if is_total else f"=MAX(-C{rr},0)", fmt=NUM)
            put(wf, f"I{rr}", f"=C{rr}", fmt='+#,##0;-#,##0;"-"' if not is_total else NUM)
        last = first + len(items) - 1
        ch = BarChart()
        ch.type = "col"
        ch.grouping = "stacked"
        ch.overlap = 100
        ch.gapWidth = 30
        ch.add_data(Reference(wf, min_col=5, max_col=8, min_row=top_row + 1, max_row=last), titles_from_data=True)
        ch.set_categories(Reference(wf, min_col=2, min_row=first, max_row=last))
        base, tot, inc, dec = ch.series
        base.graphicalProperties.noFill = True
        base.graphicalProperties.line = no_line()
        for ser, colour, fmt in ((tot, "808080", '#,##0;;'), (inc, GOOD, '"+"#,##0;;'), (dec, BAD, '"-"#,##0;;')):
            ser.graphicalProperties.solidFill = colour
            ser.graphicalProperties.line = LineProperties(solidFill="404040", w=6350)
            ser.dLbls = DataLabelList()
            ser.dLbls.showVal = True
            ser.dLbls.showSerName = ser.dLbls.showCatName = ser.dLbls.showLegendKey = False
            ser.dLbls.numFmt = fmt
            ser.dLbls.position = "inEnd" if ser is tot else "ctr"
        ch.legend = None
        ch.y_axis.delete = True
        ch.y_axis.majorGridlines = None
        ch.title = title
        ch.height, ch.width = 8.5, 20
        wf.add_chart(ch, anchor)
        return first, last

    w1 = waterfall(4, "P&L walk", walk, "L4")
    w2 = waterfall(22, "Budget to actual", bva, "L22")
    for c, w in zip("BCDEFGHI", [14, 8, 9, 8, 8, 9, 9, 8]):
        wf.column_dimensions[c].width = w

    # IBCS ---------------------------------------------------------------------------
    ib = wb.create_sheet("IBCS")
    style_sheet(ib, "IBCS", "Prior year grey, plan outlined, actual solid, forecast hatched; variances good green, bad red")
    put(ib, "B4", "Compare with", F_BOLD)
    put(ib, "E4", "Plan", fill=FILL_IN)
    dv = DataValidation(type="list", formula1="LU_Compare", allow_blank=False, showDropDown=False)
    dv.error, dv.errorTitle = "Choose Plan or Prior year", "Compare with"
    ib.add_data_validation(dv)
    dv.add("E4")
    name(wb, "DD_Compare", ib, "E4")
    heads = ["Line", "Cost line", "PY", "PL", "AC", "FC", "Compare", "Δ", "Good", "Bad"]
    for i, h in enumerate(heads):
        put(ib, f"{L(2 + i)}6", h, F_BOLD)
    lines_ib = [("Revenue", 0, 920, 1000, 1085, None), ("Materials", 1, 330, 350, 380, None), ("Labour", 1, 190, 200, 215, None),
                ("Overheads", 1, 140, 150, 140, None), ("EBITDA", 0, 260, 300, None, 350)]
    for k, (label, cost, py, pl, acv, fcv) in enumerate(lines_ib):
        rr = 7 + k
        put(ib, f"B{rr}", label)
        put(ib, f"C{rr}", cost, fmt="0", fill=FILL_IN)
        put(ib, f"D{rr}", py, fmt=NUM, fill=FILL_IN)
        put(ib, f"E{rr}", pl, fmt=NUM, fill=FILL_IN)
        put(ib, f"F{rr}", acv if acv is not None else "=NA()", fmt=NUM, fill=FILL_IN if acv is not None else None)
        put(ib, f"G{rr}", fcv if fcv is not None else "=NA()", fmt=NUM, fill=FILL_IN if fcv is not None else None)
        put(ib, f"H{rr}", f'=IF(DD_Compare="Prior year",D{rr},E{rr})', fmt=NUM)
        put(ib, f"I{rr}", f"=IFERROR(F{rr},G{rr})-H{rr}", fmt='+#,##0;-#,##0;"-"')
        put(ib, f"J{rr}", f"=IF((1-2*C{rr})*I{rr}>=0,I{rr},NA())", fmt=NUM)
        put(ib, f"K{rr}", f"=IF((1-2*C{rr})*I{rr}<0,I{rr},NA())", fmt=NUM)
    last_ib = 6 + len(lines_ib)
    col = BarChart()
    col.type = "col"
    col.grouping = "clustered"
    col.overlap = 30
    col.gapWidth = 70
    col.add_data(Reference(ib, min_col=4, max_col=7, min_row=6, max_row=last_ib), titles_from_data=True)
    col.set_categories(Reference(ib, min_col=2, min_row=7, max_row=last_ib))
    py_s, pl_s, ac_s, fc_s = col.series
    py_s.graphicalProperties.solidFill = GREY
    py_s.graphicalProperties.line = no_line()
    pl_s.graphicalProperties.noFill = True
    pl_s.graphicalProperties.line = LineProperties(solidFill=INK, w=12700)
    ac_s.graphicalProperties.solidFill = INK
    ac_s.graphicalProperties.line = no_line()
    fc_s.graphicalProperties.pattFill = PatternFillProperties(prst="upDiag", fgClr=ColorChoice(srgbClr=INK), bgClr=ColorChoice(srgbClr="FFFFFF"))
    fc_s.graphicalProperties.line = LineProperties(solidFill=INK, w=9525)
    col.y_axis.delete = True
    col.y_axis.majorGridlines = None
    col.legend.position = "b"
    col.title = "PY, PL, AC and FC"
    col.height, col.width = 8, 16
    ib.add_chart(col, "M4")
    var = BarChart()
    var.type = "bar"
    var.grouping = "clustered"
    var.overlap = 100
    var.gapWidth = 60
    var.add_data(Reference(ib, min_col=10, max_col=11, min_row=6, max_row=last_ib), titles_from_data=True)
    var.set_categories(Reference(ib, min_col=2, min_row=7, max_row=last_ib))
    good, bad = var.series
    for ser, colour in ((good, GOOD), (bad, BAD)):
        ser.graphicalProperties.solidFill = colour
        ser.graphicalProperties.line = no_line()
        ser.dLbls = DataLabelList()
        ser.dLbls.showVal = True
        ser.dLbls.showSerName = ser.dLbls.showCatName = ser.dLbls.showLegendKey = False
        ser.dLbls.numFmt = '+#,##0;-#,##0;"0"'
    var.x_axis.scaling.orientation = "maxMin"
    var.x_axis.tickLblPos = "low"
    var.y_axis.delete = True
    var.y_axis.majorGridlines = None
    var.legend = None
    var.title = "Variance (green good, red bad)"
    var.height, var.width = 8, 12
    ib.add_chart(var, "M21")
    for c, w in zip("BCDEFGHIJK", [12, 8, 8, 8, 8, 8, 9, 8, 8, 8]):
        ib.column_dimensions[c].width = w

    # Checks -----------------------------------------------------------------------
    ck = wb.create_sheet("Checks")
    style_sheet(ck, "Checks", "Error checks must be nil")
    checks = [
        ("Line selection outside the list", "=IF(OR(DD_Chart_Line<1,DD_Chart_Line>ROWS(LU_Chart_Lines)),1,0)"),
        ("Last actual month outside the year", "=IF(OR(Ts_Last_Actual<1,Ts_Last_Actual>12),1,0)"),
        ("Z chart full year differs from the data", f"=IF(ABS('Z chart'!K19-INDEX(Data!$H$1:$H$40,12+(DD_Chart_Line-1)*5))>0.000001,1,0)"),
        ("Z chart MAT does not end at the full year", "=IF(ABS('Z chart'!J19-'Z chart'!K19)>0.000001,1,0)"),
        ("P&L walk does not reach net profit", f"=IF(ABS(Waterfall!C{w1[1]}-SUM(Waterfall!C{w1[0]}:C{w1[0] + 3})-SUM(Waterfall!C{w1[0] + 5}:C{w1[0] + 8})-SUM(Waterfall!C{w1[0] + 10}:C{w1[0] + 12}))>0.000001,1,0)"),
        ("Budget to actual bridge does not close", f"=IF(ABS(Waterfall!C{w2[1]}-SUM(Waterfall!C{w2[0]}:C{w2[1] - 1}))>0.000001,1,0)"),
        ("Compare with is not Plan or Prior year", '=IF(COUNTIF(LU_Compare,DD_Compare)=0,1,0)'),
    ]
    for i, (label, f) in enumerate(checks):
        put(ck, f"C{4 + i}", label)
        put(ck, f"H{4 + i}", f, F_CHECK, "0")
    tot = 4 + len(checks)
    put(ck, f"C{tot}", "Error checks failing", F_BOLD)
    put(ck, f"H{tot}", f"=SUM(H4:H{tot - 1})", F_BOLD, "0", border=TOP)
    name(wb, "Chk_Errors", ck, f"H{tot}")
    ck.column_dimensions["C"].width = 40

    # Contents ---------------------------------------------------------------------
    toc = [("Data", "Monthly numbers, the line selector (classic drop-down) and the last actual month"),
           ("Z chart", "Z chart for the selected line"), ("Waterfall", "P&L walk and budget to actual bridges"),
           ("IBCS", "IBCS columns and variances, with the comparison chosen in-cell"),
           ("Lookups", "Lists behind the drop-downs"), ("Checks", "Error checks")]
    for i, (s, d) in enumerate(toc):
        c = put(contents, f"C{4 + i}", s)
        c.hyperlink = Hyperlink(ref=f"C{4 + i}", location=f"'{s}'!A1")
        put(contents, f"F{4 + i}", d)
    put(contents, "C11", "Error checks failing", F_BOLD)
    put(contents, "F11", "=Chk_Errors", F_CHECK, "0")
    contents.column_dimensions["C"].width = 14

    for ws in wb.worksheets:
        for rr in range(1, ws.max_row + 1):
            ws.row_dimensions[rr].height = 15
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 1
        ws.sheet_properties.pageSetUpPr.fitToPage = True
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    add_form_control(path, "Data", link="DD_Chart_Line", lst="LU_Chart_Lines", col=(3, 7), row=3, lines=len(names))
    finish_charts(path)
    return path


NA_AS_BLANK = ('<extLst><ext uri="{56B9EC1D-385E-4148-901F-78D8002777C0}" '
               'xmlns:c16r3="http://schemas.microsoft.com/office/drawing/2017/03/chart">'
               '<c16r3:dataDisplayOptions16><c16r3:dispNaAsBlank val="1"/></c16r3:dataDisplayOptions16></ext></extLst>')


def finish_charts(path: Path) -> None:
    """Settings openpyxl does not write: Excel's "Show #N/A as an empty cell" on every chart, so #N/A
    draws nothing and shows no label, and label number formats that do not follow the cells."""
    with zipfile.ZipFile(path) as z:
        parts = {n: z.read(n) for n in z.namelist()}
    for n in parts:
        if n.startswith("xl/charts/chart") and n.endswith(".xml"):
            x = parts[n].decode()
            if "dispNaAsBlank" not in x:
                x = x.replace("</chart>", NA_AS_BLANK + "</chart>", 1)
            x = re.sub(r'<numFmt formatCode="([^"]*)"/>', r'<numFmt formatCode="\1" sourceLinked="0"/>', x)
            parts[n] = x.encode()
    fd, tmp = tempfile.mkstemp(suffix=".xlsx")
    os.close(fd)
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as z:
        for n, b in parts.items():
            z.writestr(n, b)
    shutil.move(tmp, path)


def add_form_control(path: Path, sheet: str, link: str, lst: str, col: tuple[int, int], row: int, lines: int) -> None:
    """Write a classic combo box (form control) into a saved workbook, as Excel stores one."""
    with zipfile.ZipFile(path) as z:
        parts = {n: z.read(n) for n in z.namelist()}
    wbx = parts["xl/workbook.xml"].decode()
    rels = parts["xl/_rels/workbook.xml.rels"].decode()
    rid = re.search(rf'<sheet [^>]*name="{re.escape(sheet)}"[^>]*r:id="(rId\d+)"', wbx).group(1)
    target = re.search(rf'<Relationship [^>]*Id="{rid}"[^>]*Target="([^"]+)"', rels) or \
        re.search(rf'<Relationship [^>]*Target="([^"]+)"[^>]*Id="{rid}"', rels)
    sheet_part = "xl/" + target.group(1).lstrip("/").replace("xl/", "")
    sheet_rels = sheet_part.replace("worksheets/", "worksheets/_rels/") + ".rels"
    c0, c1 = col
    shape_id = 1025
    parts["xl/ctrlProps/ctrlProp1.xml"] = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<formControlPr xmlns="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main" objectType="Drop" '
        f'dropLines="{lines}" dropStyle="combo" dx="26" fmlaLink="{link}" fmlaRange="{lst}" noThreeD="1" sel="1" val="0"/>').encode()
    parts["xl/drawings/vmlDrawing1.vml"] = (
        '<xml xmlns:v="urn:schemas-microsoft-com:vml" xmlns:o="urn:schemas-microsoft-com:office:office" '
        'xmlns:x="urn:schemas-microsoft-com:office:excel">\n'
        ' <o:shapelayout v:ext="edit"><o:idmap v:ext="edit" data="1"/></o:shapelayout>'
        '<v:shapetype id="_x0000_t201" coordsize="21600,21600" o:spt="201" path="m,l,21600r21600,l21600,xe">'
        '<v:stroke joinstyle="miter"/><v:path shadowok="f" o:extrusionok="f" strokeok="f" fillok="f" o:connecttype="rect"/>'
        '<o:lock v:ext="edit" shapetype="t"/></v:shapetype>'
        f'<v:shape id="_x0000_s{shape_id}" type="#_x0000_t201" style="position:absolute;margin-left:0;margin-top:0;width:150pt;height:15pt;z-index:1;mso-wrap-style:tight" '
        'stroked="f" strokecolor="windowText [64]" o:insetmode="auto"><o:lock v:ext="edit" rotation="t" text="t"/>'
        '<x:ClientData ObjectType="Drop"><x:SizeWithCells/>'
        f'<x:Anchor>{c0}, 0, {row}, 0, {c1}, 0, {row + 1}, 0</x:Anchor><x:AutoLine>False</x:AutoLine>'
        f'<x:FmlaLink>{link}</x:FmlaLink><x:Val>0</x:Val><x:Min>0</x:Min><x:Max>{lines}</x:Max><x:Inc>1</x:Inc><x:Page>{lines}</x:Page>'
        f'<x:Dx>26</x:Dx><x:FmlaRange>{lst}</x:FmlaRange><x:Sel>1</x:Sel><x:NoThreeD2/><x:SelType>Single</x:SelType>'
        f'<x:LCT>Normal</x:LCT><x:DropStyle>Combo</x:DropStyle><x:DropLines>{lines}</x:DropLines></x:ClientData></v:shape></xml>').encode()
    anchor = (f'<xdr:from><xdr:col>{c0}</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>{row}</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:from>'
              f'<xdr:to><xdr:col>{c1}</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>{row + 1}</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:to>')
    parts["xl/drawings/drawingControls1.xml"] = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<xdr:wsDr xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        '<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006">'
        '<mc:Choice xmlns:a14="http://schemas.microsoft.com/office/drawing/2010/main" Requires="a14">'
        f'<xdr:twoCellAnchor editAs="oneCell">{anchor}<xdr:sp macro="" textlink=""><xdr:nvSpPr>'
        f'<xdr:cNvPr id="{shape_id}" name="Drop Down 1" hidden="1"><a:extLst><a:ext uri="{{63B3BB69-23CF-44E3-9099-C40C66FF867C}}">'
        f'<a14:compatExt spid="_x0000_s{shape_id}"/></a:ext></a:extLst></xdr:cNvPr><xdr:cNvSpPr/></xdr:nvSpPr>'
        '<xdr:spPr bwMode="auto"><a:xfrm><a:off x="0" y="0"/><a:ext cx="0" cy="0"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom>'
        '<a:noFill/><a:ln><a:noFill/></a:ln></xdr:spPr></xdr:sp><xdr:clientData/></xdr:twoCellAnchor></mc:Choice><mc:Fallback/></mc:AlternateContent>'
        '</xdr:wsDr>').encode()
    srels = parts.get(sheet_rels, b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"></Relationships>').decode()
    if "/drawing\"" in srels or 'relationships/drawing"' in srels:
        raise RuntimeError(f"{sheet} already has a drawing; put the control on a sheet without charts")
    srels = srels.replace("</Relationships>",
        '<Relationship Id="rIdCtlD" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/drawing" Target="../drawings/drawingControls1.xml"/>'
        '<Relationship Id="rIdCtlV" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/vmlDrawing" Target="../drawings/vmlDrawing1.vml"/>'
        '<Relationship Id="rIdCtlP" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/ctrlProp" Target="../ctrlProps/ctrlProp1.xml"/>'
        "</Relationships>")
    parts[sheet_rels] = srels.encode()
    sx = parts[sheet_part].decode()
    if 'xmlns:r="' not in sx[:600]:
        sx = sx.replace("<worksheet ", '<worksheet xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" ', 1)
    if 'xmlns:xdr="' not in sx[:800]:
        sx = sx.replace("<worksheet ", '<worksheet xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" ', 1)
    controls = ('<drawing r:id="rIdCtlD"/><legacyDrawing r:id="rIdCtlV"/>'
                '<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"><mc:Choice Requires="x14"><controls>'
                '<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"><mc:Choice Requires="x14">'
                f'<control shapeId="{shape_id}" r:id="rIdCtlP" name="Drop Down 1"><controlPr defaultSize="0" autoLine="0" autoPict="0">'
                f'<anchor moveWithCells="1">{anchor}</anchor></controlPr></control></mc:Choice></mc:AlternateContent>'
                '</controls></mc:Choice></mc:AlternateContent>')
    tail = re.search(r"(<tableParts|<extLst|</worksheet>)", sx)
    sx = sx[:tail.start()] + controls + sx[tail.start():]
    parts[sheet_part] = sx.encode()
    ct = parts["[Content_Types].xml"].decode()
    if 'Extension="vml"' not in ct:
        ct = ct.replace("</Types>", '<Default Extension="vml" ContentType="application/vnd.openxmlformats-officedocument.vmlDrawing"/></Types>')
    ct = ct.replace("</Types>", '<Override PartName="/xl/ctrlProps/ctrlProp1.xml" ContentType="application/vnd.ms-excel.controlproperties+xml"/>'
                    '<Override PartName="/xl/drawings/drawingControls1.xml" ContentType="application/vnd.openxmlformats-officedocument.drawing+xml"/></Types>')
    parts["[Content_Types].xml"] = ct.encode()
    fd, tmp = tempfile.mkstemp(suffix=".xlsx")
    os.close(fd)
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as z:
        for n, b in parts.items():
            z.writestr(n, b)
    shutil.move(tmp, path)


if __name__ == "__main__":
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).resolve().parents[2] / "build" / "charts" / "charts_demo.xlsx"
    print(f"Built {build(out)}")
