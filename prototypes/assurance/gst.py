"""GST return periods and due dates, for the GST module's cash timing.

New Zealand GST returns and payments are due on the 28th of the month after the period ends, with
two exceptions: a period ending in November is due on 15 January, and a period ending in March on
7 May. An entity files monthly, two-monthly or six-monthly, and a two-monthly or six-monthly filer
has its own cycle (periods ending in odd or even months; six-monthly ending in, say, March and
September). Each entity in a model carries its frequency and cycle as inputs.

This proof writes the timing as formulas (the shape the GST module's lookup will take), with an
independent Python calculation to check it: for each month of the model, the return period it
falls in, the period's due date and the month the GST is paid, and the cash flow when net GST for
a period is paid on the due date or refunded a set number of months later.
"""

from __future__ import annotations

import calendar
from datetime import date
from pathlib import Path

FONT, TEXT = "Segoe UI", "404040"
FIRST = 10          # first month row
START_CELL = "F5"   # model start (first month end)
FREQ_CELL = "F6"
CYCLE_CELL = "F7"
LAG_CELL = "F8"


def month_end(y: int, m: int) -> date:
    y, m = y + (m - 1) // 12, (m - 1) % 12 + 1
    return date(y, m, calendar.monthrange(y, m)[1])


def add_months(d: date, n: int) -> date:
    return month_end(d.year, d.month + n)


def period_end(d: date, months: int, cycle: int) -> date:
    """End of the return period month d falls in: periods end in months where month mod frequency = cycle."""
    return add_months(d, (cycle - d.month) % months)


def due_date(end: date) -> date:
    if end.month == 11:
        return date(end.year + 1, 1, 15)
    if end.month == 3:
        return date(end.year, 5, 7)
    nxt = add_months(end, 1)
    return date(nxt.year, nxt.month, 28)


def months_between(a: date, b: date) -> int:
    return (b.year - a.year) * 12 + b.month - a.month


def reference(start: date, n: int, months: int, cycle: int, lag: int, net: list[float]) -> list[dict]:
    """Month by month: period end, due date, month paid (1 = first month), and the GST cash flow."""
    rows = []
    for t in range(n):
        m = add_months(start, t)
        pe = period_end(m, months, cycle)
        rows.append({"month": t + 1, "end": m, "period_end": pe, "due": due_date(pe),
                     "due_month": months_between(start, due_date(pe)) + 1, "net": net[t]})
    cash = [0.0] * n
    for r in rows:
        if r["end"] == r["period_end"]:
            total = sum(x["net"] for x in rows if x["period_end"] == r["period_end"])
            r["period_net"] = total
            r["paid_month"] = r["due_month"] + (lag if total < 0 else 0)
            if 1 <= r["paid_month"] <= n:
                cash[r["paid_month"] - 1] -= total
        else:
            r["period_net"], r["paid_month"] = 0.0, None
    for r, c in zip(rows, cash):
        r["cash"] = c
    return rows


def write_workbook(path: Path, start: date, n: int, months: int, cycle: int, lag: int, net: list[float]) -> Path:
    """The GST timing block as formulas. Inputs at the top; one row per month from row 10."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    body, bold = Font(name=FONT, size=9, color=TEXT), Font(name=FONT, size=9, color=TEXT, bold=True)
    title = Font(name=FONT, size=10, color=TEXT, bold=True)
    inp = PatternFill("solid", fgColor="FFF2CC")
    wb = Workbook()
    ws = wb.active
    ws.title = "GST timing"
    ws.sheet_view.showGridLines = False
    ws["B1"], ws["B2"] = "GST return periods and due dates", "Assurance proof (demo data)"
    ws["B1"].font, ws["B2"].font = title, body
    inputs = [("Model start (first month end)", start, "d mmmm yyyy"), ("Filing frequency (months)", months, "0"),
              ("Periods end where month mod frequency is", cycle, "0"), ("Refunds received, months after the due date", lag, "0")]
    for k, (label, v, fmt) in enumerate(inputs):
        r = 5 + k
        ws.cell(r, 2, label).font = body
        c = ws.cell(r, 6, v)
        c.font, c.fill, c.number_format = body, inp, fmt
    heads = ["Month", "Month end", "Period ends", "Return due", "Due in month", "Net GST for the month",
             "Net GST for the period", "Paid or received in month", "GST cash flow"]
    for k, h in enumerate(heads):
        ws.cell(FIRST - 1, 2 + k, h).font = bold
    last = FIRST + n - 1
    for t in range(n):
        r = FIRST + t
        ws.cell(r, 2, t + 1)
        ws.cell(r, 3, f"=EOMONTH({START_CELL},B{r}-1)")
        ws.cell(r, 4, f"=EOMONTH(C{r},MOD({CYCLE_CELL}-MONTH(C{r}),{FREQ_CELL}))")
        ws.cell(r, 5, f"=IF(MONTH(D{r})=11,DATE(YEAR(D{r})+1,1,15),IF(MONTH(D{r})=3,DATE(YEAR(D{r}),5,7),"
                      f"DATE(YEAR(D{r}),MONTH(D{r})+1,28)))")
        ws.cell(r, 6, f"=(YEAR(E{r})-YEAR({START_CELL}))*12+MONTH(E{r})-MONTH({START_CELL})+1")
        ws.cell(r, 7, net[t]).fill = inp
        ws.cell(r, 8, f"=IF(C{r}=D{r},SUMIF($D${FIRST}:$D${last},D{r},$G${FIRST}:$G${last}),0)")
        ws.cell(r, 9, f'=IF(C{r}=D{r},F{r}+IF(H{r}<0,{LAG_CELL},0),"")')
        ws.cell(r, 10, f"=-SUMIF($I${FIRST}:$I${last},B{r},$H${FIRST}:$H${last})")
        for c in range(2, 11):
            ws.cell(r, c).font = body
        for c in (3, 4, 5):
            ws.cell(r, c).number_format = "d mmm yyyy"
        for c in (7, 8, 10):
            ws.cell(r, c).number_format = '#,##0.00;(#,##0.00);"-"'
        ws.row_dimensions[r].height = 15
    ws.column_dimensions["A"].width = 2.5
    ws.column_dimensions["B"].width = 8
    for col, w in zip("CDEFGHIJ", (13, 13, 13, 14, 22, 22, 24, 15)):
        ws.column_dimensions[col].width = w
    wb.save(path)
    return Path(path)
