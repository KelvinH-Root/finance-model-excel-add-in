"""Phase 0 proof: insert a module into a built model and have it link itself in.

The engine is model-agnostic. It knows nothing about revenue, debtors or debt:
it reads module definitions (library/*.yaml), resolves the links between module
instances to a fixed point, lays the model out on sheets, and computes the change
plan between two layouts. The same plan builds a new workbook (package writer)
or is applied to an open one (live writer), so both end up identical.

Formulas are held as structural references (markers naming a row id, not a cell
address) until they are rendered, which is what lets the engine tell whether an
existing formula still points at the same rows after an insert.
"""

from __future__ import annotations

import copy
import json
import re
from dataclasses import dataclass, field
from pathlib import Path

import yaml

HERE = Path(__file__).parent
LIBRARY = HERE / "library"

# Frame standard columns: labels B to G, units H, totals I, timeline from J.
LABEL_COLS = {0: 2, 1: 3, 2: 4}
UNIT_COL = 8
TOTAL_COL = 9
FIRST_PERIOD_COL = 10
PERIOD_ROW = 5
FIRST_ROW = 7
CONTENTS = "Contents"
META_NS = "urn:hfg:model-metadata:v0"


class AssemblyError(Exception):
    """A module cannot be inserted or the links cannot be resolved."""


def col_letter(c: int) -> str:
    s = ""
    while c:
        c, r = divmod(c - 1, 26)
        s = chr(65 + r) + s
    return s


# --------------------------------------------------------------------------
# Library and model
# --------------------------------------------------------------------------

@dataclass
class Library:
    areas: list[str]
    modules: dict[str, dict]

    @classmethod
    def load(cls, path: Path = LIBRARY) -> "Library":
        path = Path(path)
        areas = yaml.safe_load((path / "areas.yaml").read_text())["areas"]
        modules = {}
        for f in sorted(path.glob("*.yaml")):
            if f.name == "areas.yaml":
                continue
            d = yaml.safe_load(f.read_text())
            if d["area"] not in areas:
                raise AssemblyError(f"{f.name}: area {d['area']!r} is not in areas.yaml")
            modules[d["id"]] = d
        return cls(areas, modules)

    def kind(self, module_id: str) -> str:
        m = self.modules[module_id]
        if m.get("mirror"):
            return "mirror"
        return "category" if m.get("as_category") else "single"


@dataclass
class Instance:
    module: str
    number: int
    settings: dict

    @property
    def uid(self) -> str:
        return f"{self.module}#{self.number}"


class Model:
    """The list of module instances a workbook is built from."""

    def __init__(self, library: Library, periods: int = 12):
        self.lib = library
        self.periods = periods
        self.instances: list[Instance] = []
        self.counters: dict[str, int] = {}

    # Step 1 (compatibility) and step 2 (instance naming) of the insert.
    def insert(self, module_id: str, **settings) -> Instance:
        if module_id not in self.lib.modules:
            raise AssemblyError(f"no module {module_id!r} in the library")
        mod = self.lib.modules[module_id]
        if self.lib.kind(module_id) != "category" and any(i.module == module_id for i in self.instances):
            raise AssemblyError(f"{mod['title']} is already in the model and can only be inserted once")
        known = {s["key"]: s for s in mod.get("settings", [])}
        unknown = set(settings) - set(known)
        if unknown:
            raise AssemblyError(f"{mod['title']} has no setting {', '.join(sorted(unknown))}")
        values = {k: s.get("default") for k, s in known.items()}
        values.update(settings)
        n = self.counters.get(module_id, 0) + 1  # numbers are never reused
        self.counters[module_id] = n
        inst = Instance(module_id, n, values)
        self.instances.append(inst)
        return inst

    def remove(self, uid: str) -> None:
        before = len(self.instances)
        self.instances = [i for i in self.instances if i.uid != uid]
        if len(self.instances) == before:
            raise AssemblyError(f"no instance {uid} in the model")

    def copy(self) -> "Model":
        m = Model(self.lib, self.periods)
        m.instances = copy.deepcopy(self.instances)
        m.counters = dict(self.counters)
        return m

    def title(self, inst: Instance) -> str:
        mod = self.lib.modules[inst.module]
        return f"{mod['title']} {inst.number}" if self.lib.kind(inst.module) == "category" else mod["title"]

    def to_dict(self) -> dict:
        return {"periods": self.periods, "counters": self.counters,
                "instances": [{"module": i.module, "number": i.number, "settings": i.settings} for i in self.instances]}

    @classmethod
    def from_dict(cls, library: Library, d: dict) -> "Model":
        m = cls(library, d["periods"])
        m.counters = dict(d["counters"])
        m.instances = [Instance(i["module"], i["number"], dict(i["settings"])) for i in d["instances"]]
        return m


# --------------------------------------------------------------------------
# Link resolution (steps 3, 5 and 6)
# --------------------------------------------------------------------------

@dataclass
class Block:
    id: str
    inst: Instance
    mod: dict
    title: str
    src: tuple[str, str] | None = None  # (producer block id, row key) for a mirror block


def _producers(blocks: list[Block]) -> dict[str, list[tuple[Block, str]]]:
    out: dict[str, list[tuple[Block, str]]] = {}
    for b in blocks:
        for o in b.mod.get("outputs", []):
            out.setdefault(o["link"], []).append((b, o["row"]))
        for r in b.mod.get("rows", []):
            if r.get("check"):
                out.setdefault(f"check.{r['check']}", []).append((b, r["key"]))
    return out


def resolve(model: Model) -> tuple[list[Block], dict]:
    lib = model.lib
    order = {i.uid: k for k, i in enumerate(model.instances)}
    ordered = sorted(model.instances, key=lambda i: (lib.areas.index(lib.modules[i.module]["area"]), order[i.uid]))
    blocks: list[Block] = []
    for _ in range(25):  # mirror blocks can produce links other mirrors take, so repeat until nothing new appears
        producers = _producers(blocks)
        new: list[Block] = []
        for inst in ordered:
            mod = lib.modules[inst.module]
            title = model.title(inst)
            if mod.get("mirror"):
                for pb, row in producers.get(mod["mirror"], []):
                    if pb.inst.uid == inst.uid:
                        continue
                    new.append(Block(f"{inst.uid}/{pb.id}", inst, mod, f"{title}: {pb.title}", (pb.id, row)))
            else:
                new.append(Block(inst.uid, inst, mod, title))
        if [b.id for b in new] == [b.id for b in blocks]:
            return blocks, producers
        blocks = new
    raise AssemblyError("links did not settle after 25 passes (a mirror feeds itself)")


# --------------------------------------------------------------------------
# Layout (step 4 and step 7)
# --------------------------------------------------------------------------

@dataclass
class LRow:
    id: str
    kind: str                 # heading, subheading, section, setting, series, text
    label: str
    indent: int = 0
    unit: str = ""
    style: str = ""           # total, check
    value: object = None      # setting value
    name: str | None = None   # defined name on the value or total cell
    first: str | None = None  # marker formula for the first period
    formula: str | None = None  # marker formula for later periods
    total: str = "sum"        # sum, last, none
    cells: dict = field(default_factory=dict)  # col -> marker formula or text (Contents rows)

    def signature(self) -> tuple:
        """What the engine owns in this row. Input values are left out: a structural change never overwrites what someone typed."""
        return (self.kind, self.label, self.indent, self.unit, self.style, self.name, self.first,
                self.formula, self.total, tuple(sorted(self.cells.items())))


@dataclass
class Layout:
    periods: int
    sheets: list[tuple[str, list[LRow]]]
    names: dict[str, str]                   # name -> row id
    records: list[dict]                     # link resolution records
    warnings: list[str]
    blocks: dict[str, str]                  # block id -> title

    def positions(self) -> dict[str, tuple[str, int]]:
        pos = {}
        for sheet, rows in self.sheets:
            for k, r in enumerate(rows):
                pos[r.id] = (sheet, FIRST_ROW + k)
        return pos

    def sheet_rows(self) -> dict[str, list[LRow]]:
        return dict(self.sheets)


_SUM = re.compile(r"\[sum:([\w.]+)\]")
_PREV = re.compile(r"\[(\w+)@prev\]")
_REF = re.compile(r"\[(\w+)\]")
_SET = re.compile(r"\$([a-z]\w*)")


def _compile(tpl: str | None, b: Block, keys: dict[str, str], names: dict[str, str],
             collects: dict[str, list[str]], src: str | None) -> str | None:
    if tpl is None:
        return None
    where = f"{b.mod['id']} ({b.title})"

    def sum_(m):
        ids = collects.get(m.group(1))
        if ids is None:
            raise AssemblyError(f"{where}: [sum:{m.group(1)}] but the module does not take {m.group(1)}")
        return f"SUM(«R|{ids[0]}»:«R|{ids[-1]}»)" if ids else "0"

    def key(m, kind):
        k = m.group(1)
        if k not in keys:
            raise AssemblyError(f"{where}: no row {k!r}")
        return f"«{kind}|{keys[k]}»"

    def setting(m):
        k = m.group(1)
        if k not in names:
            raise AssemblyError(f"{where}: no setting {k!r}")
        return names[k]

    out = _SUM.sub(sum_, tpl)
    out = _PREV.sub(lambda m: key(m, "P"), out)
    out = _REF.sub(lambda m: key(m, "R"), out)
    out = _SET.sub(setting, out)
    if "{src}" in out:
        if src is None:
            raise AssemblyError(f"{where}: {{src}} is only allowed in a mirror module")
        out = out.replace("{src}", f"«R|{src}»")
    return out


def _name_part(key: str) -> str:
    return "".join(p.capitalize() for p in key.split("_"))


def assemble(model: Model) -> Layout:
    lib = model.lib
    blocks, producers = resolve(model)
    by_id = {b.id: b for b in blocks}
    row_id = lambda b, key: f"{b.id}/{key}"  # noqa: E731
    records: list[dict] = []
    warnings: list[str] = []
    names: dict[str, str] = {}
    sheets: dict[str, list[LRow]] = {}
    headed: set[str] = set()

    for b in blocks:
        mod = b.mod
        rows = sheets.setdefault(mod["area"], [])
        inst_title = model.title(b.inst)
        mirror = b.src is not None
        if b.inst.uid not in headed:  # one section bar per instance; a mirror has one sub-block per sender
            headed.add(b.inst.uid)
            rows.append(LRow(f"{b.inst.uid}/heading", "heading", inst_title))
        if mirror:
            rows.append(LRow(f"{b.id}/subheading", "subheading", b.title.split(": ", 1)[1], indent=1))

        # Settings become named input cells (step 7 registers the names).
        set_names = {}
        code = mod.get("code", "M")
        for s in mod.get("settings", []):
            nm = f"{code}{b.inst.number}_{_name_part(s['key'])}"
            set_names[s["key"]] = nm
            rid = row_id(b, f"set.{s['key']}")
            rows.append(LRow(rid, "setting", s["label"], indent=1, unit=s.get("unit", ""),
                             value=b.inst.settings.get(s["key"]), name=nm))
            names[nm] = rid

        # Row ids for keys first, so formulas can point forwards as well as back.
        keys = {r["key"]: row_id(b, r["key"]) for r in mod.get("rows", []) if "key" in r}
        collects: dict[str, list[str]] = {}
        planned: list[tuple[dict, list]] = []
        for r in mod.get("rows", []):
            if "collect" in r:
                link = r["collect"]
                spec = next((i for i in mod.get("inputs", []) if i["link"] == link), None)
                if spec is None:
                    raise AssemblyError(f"{mod['id']}: collects {link} but does not list it in inputs")
                found = producers.get(link, [])
                if spec.get("mode", "each") == "each":
                    collect_rows = []
                    for pb, prow in found:
                        rid = f"{b.id}/in/{link}/{pb.id}"
                        collect_rows.append((rid, pb, prow))
                    collects[link] = [c[0] for c in collect_rows]
                    planned.append((r, collect_rows))
                else:  # total: one row adding every sender
                    rid = f"{b.id}/in/{link}"
                    collects[link] = [rid]
                    planned.append((r, [(rid, None, found)]))
            else:
                planned.append((r, []))
        for spec in mod.get("inputs", []):
            if spec.get("required") and not producers.get(spec["link"]):
                warnings.append(f"{b.title}: required link {spec['link']} has no sender")

        src = row_id(by_id[b.src[0]], b.src[1]) if mirror else None
        if mirror:
            records.append({"link": mod["mirror"], "mode": "mirror", "from": src, "to": b.id})
        for r, collect_rows in planned:
            if "section" in r:
                rows.append(LRow(f"{b.id}/section/{_name_part(r['section'])}", "section", r["section"], indent=1))
                continue
            if "collect" in r:
                link = r["collect"]
                for rid, pb, prow in collect_rows:
                    if pb is None:  # total mode
                        senders = prow
                        f = "=" + ("+".join(f"«R|{row_id(x, k)}»" for x, k in senders) if senders else "0")
                        rows.append(LRow(rid, "series", f"{link} (all senders)", indent=2, unit=r.get("unit", ""),
                                         formula=f, total=r.get("total", "sum")))
                        for x, k in senders:
                            records.append({"link": link, "mode": "total", "from": row_id(x, k), "to": rid})
                    else:
                        prow_def = next(x for x in pb.mod["rows"] if x.get("key") == prow)
                        label = pb.title if not link.startswith("check.") else f"{pb.title}: {prow_def['label']}"
                        rows.append(LRow(rid, "series", label, indent=2, unit=r.get("unit", prow_def.get("unit", "")),
                                         formula=f"=«R|{row_id(pb, prow)}»", total=r.get("total", prow_def.get("total", "sum"))))
                        records.append({"link": link, "mode": "each", "from": row_id(pb, prow), "to": rid})
                continue
            rid = keys[r["key"]]
            nm = r.get("name")
            if nm:
                names[nm] = rid
            rows.append(LRow(rid, "series", r["label"], indent=2 if mirror else 1, unit=r.get("unit", ""),
                             style="check" if r.get("check") else r.get("style", ""), name=nm,
                             first=_compile(r.get("first"), b, keys, set_names, collects, src),
                             formula=_compile(r.get("formula"), b, keys, set_names, collects, src),
                             total=r.get("total", "sum")))
        rows.append(LRow(f"{b.id}/end", "blank", ""))

    linked_in = {rec["from"] for rec in records}
    for link, plist in producers.items():
        for pb, prow in plist:
            if row_id(pb, prow) not in linked_in and not link.startswith("check."):
                warnings.append(f"{pb.title}: {link} is not taken by any module")

    ordered = [(a, sheets[a]) for a in lib.areas if a in sheets]
    ordered.insert(0, (CONTENTS, _contents(model, blocks, names)))
    return Layout(model.periods, ordered, names, records, warnings, {b.id: b.title for b in blocks})


def _contents(model: Model, blocks: list[Block], names: dict[str, str]) -> list[LRow]:
    rows = [LRow("contents/heading", "heading", "Modules")]
    seen = set()
    for b in blocks:
        if b.inst.uid in seen:
            continue
        seen.add(b.inst.uid)
        rows.append(LRow(f"contents/{b.inst.uid}", "text", model.title(b.inst), indent=1,
                         cells={UNIT_COL: b.mod["area"]}))
    rows.append(LRow("contents/end", "blank", ""))
    if "Chk_Errors" in names:
        rows.append(LRow("contents/checks", "heading", "Checks"))
        rows.append(LRow("contents/errors", "text", "Error checks failing", indent=1, style="check",
                         cells={TOTAL_COL: "=Chk_Errors"}))
        rows.append(LRow("contents/alerts", "text", "Alerts raised", indent=1, cells={TOTAL_COL: "=Chk_Alerts"}))
    return rows


# --------------------------------------------------------------------------
# Rendering markers into cell formulas
# --------------------------------------------------------------------------

_MARK = re.compile(r"«([RP])\|([^»]+)»")


def _sheet_prefix(sheet: str, dialect: str) -> str:
    plain = re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]*", sheet) is not None
    if dialect == "uno":
        return f"${sheet}." if plain else f"$'{sheet}'."
    return f"{sheet}!" if plain else f"'{sheet}'!"


def render_formula(f: str, sheet: str, col: int, pos: dict, dialect: str = "excel") -> str:
    def ref(m):
        kind, rid = m.group(1), m.group(2)
        s, r = pos[rid]
        c = col if kind == "R" else col - 1
        if c < FIRST_PERIOD_COL:
            return "0"  # the period before the first one
        a1 = f"{col_letter(c)}{r}"
        return a1 if s == sheet else _sheet_prefix(s, dialect) + a1
    out = _MARK.sub(ref, f)
    if dialect == "uno":
        out = out.replace(",", ";")
    return out


def row_cells(layout: Layout, sheet: str, row: LRow, rownum: int, pos: dict, dialect: str = "excel") -> dict:
    """Every cell the engine owns in one row, rendered for one dialect."""
    cells: dict[int, object] = {}
    if row.kind == "blank":
        return cells
    cells[LABEL_COLS[min(row.indent, 2)]] = row.label
    if row.unit:
        cells[UNIT_COL] = row.unit
    if row.kind == "setting":
        cells[TOTAL_COL] = row.value
    if row.kind == "series":
        last = FIRST_PERIOD_COL + layout.periods - 1
        for c in range(FIRST_PERIOD_COL, last + 1):
            tpl = row.first if (c == FIRST_PERIOD_COL and row.first) else row.formula
            cells[c] = render_formula(tpl, sheet, c, pos, dialect)
        if row.total == "sum":
            cells[TOTAL_COL] = f"=SUM({col_letter(FIRST_PERIOD_COL)}{rownum}:{col_letter(last)}{rownum})"
        elif row.total == "last":
            cells[TOTAL_COL] = f"={col_letter(last)}{rownum}"
    for c, v in row.cells.items():
        cells[c] = v
    return cells


# --------------------------------------------------------------------------
# Change plan (the diff between two layouts)
# --------------------------------------------------------------------------

@dataclass
class Plan:
    ops: list[dict]
    preview: list[str]


def _runs(nums: list[int]) -> list[tuple[int, int]]:
    runs = []
    for n in nums:
        if runs and n == runs[-1][0] + runs[-1][1]:
            runs[-1] = (runs[-1][0], runs[-1][1] + 1)
        else:
            runs.append((n, 1))
    return runs


def plan_change(old: Layout, new: Layout, dialect: str = "excel") -> Plan:
    ops: list[dict] = []
    old_sheets, new_sheets = old.sheet_rows(), new.sheet_rows()
    pos = new.positions()
    preview: list[str] = []

    for nm in sorted(set(old.names) - set(new.names)):
        ops.append({"op": "delete_name", "name": nm})
    for s in old_sheets:
        if s not in new_sheets:
            ops.append({"op": "delete_sheet", "sheet": s})
            preview.append(f"{s}: sheet removed (no module left on it).")
    for idx, (s, rows) in enumerate(new.sheets):
        if s not in old_sheets:
            ops.append({"op": "add_sheet", "sheet": s, "index": idx, "periods": new.periods})
            preview.append(f"{s}: new sheet.")

    changed_rows: dict[str, list[tuple[int, LRow, str]]] = {}
    for s, rows in new.sheets:
        new_ids = [r.id for r in rows]
        if s not in old_sheets:
            changed_rows[s] = [(FIRST_ROW + k, r, "new") for k, r in enumerate(rows)]
            continue
        old_rows = old_sheets[s]
        old_ids = [r.id for r in old_rows]
        old_set, new_set = set(old_ids), set(new_ids)
        if [i for i in old_ids if i in new_set] != [i for i in new_ids if i in old_set]:
            raise AssemblyError(f"{s}: rows would change order, which needs a move the plan does not support")
        for start, count in reversed(_runs([FIRST_ROW + k for k, i in enumerate(old_ids) if i not in new_set])):
            ops.append({"op": "delete_rows", "sheet": s, "row": start, "count": count})
        for start, count in _runs([FIRST_ROW + k for k, i in enumerate(new_ids) if i not in old_set]):
            ops.append({"op": "insert_rows", "sheet": s, "row": start, "count": count})
        old_by_id = {r.id: r for r in old_rows}
        lst = []
        for k, r in enumerate(rows):
            if r.id not in old_set:
                lst.append((FIRST_ROW + k, r, "new"))
            elif r.signature() != old_by_id[r.id].signature():
                lst.append((FIRST_ROW + k, r, "rewire"))
        changed_rows[s] = lst

    for s, lst in changed_rows.items():
        for rownum, r, why in lst:
            cells = row_cells(new, s, r, rownum, pos, dialect)
            if why == "rewire" and r.kind == "setting":
                cells.pop(TOTAL_COL, None)  # keep the input someone typed
            ops.append({"op": "write", "sheet": s, "row": rownum, "why": why, "kind": r.kind,
                        "style": r.style, "unit": r.unit, "cells": cells})

    for nm, rid in new.names.items():
        if old.names.get(nm) != rid:
            s, r = pos[rid]
            col = TOTAL_COL
            ops.append({"op": "add_name", "name": nm, "sheet": s, "row": r, "col": col})

    # Preview, in plain words, before anything is touched.
    for s, lst in changed_rows.items():
        new_n = sum(1 for _, r, w in lst if w == "new" and r.kind != "blank")
        rew = [r.label for _, r, w in lst if w == "rewire"]
        if new_n:
            spans = [f"{a}" if n == 1 else f"{a} to {a + n - 1}"
                     for a, n in _runs([n for n, r, w in lst if w == "new"])]
            where = spans[0] if len(spans) == 1 else ", ".join(spans[:-1]) + " and " + spans[-1]
            preview.append(f"{s}: {new_n} new row{'s' if new_n > 1 else ''} (row{'s' if new_n > 1 else ''} {where}).")
        if rew:
            preview.append(f"{s}: {len(rew)} rows rewired ({', '.join(rew)}).")
    for s, rows in old.sheets:
        if s in new_sheets:
            gone = [r for r in rows if r.id not in {x.id for x in new_sheets[s]} and r.kind != "blank"]
            if gone:
                preview.append(f"{s}: {len(gone)} rows removed.")
    key = lambda rec: (rec["link"], rec["from"], rec["to"])  # noqa: E731
    added = [r for r in new.records if key(r) not in {key(x) for x in old.records}]
    removed = [r for r in old.records if key(r) not in {key(x) for x in new.records}]
    for label, recs in (("New links", added), ("Links removed", removed)):
        if recs:
            by_link: dict[str, int] = {}
            for rec in recs:
                by_link[rec["link"]] = by_link.get(rec["link"], 0) + 1
            preview.append(f"{label}: " + ", ".join(f"{k} x{v}" for k, v in sorted(by_link.items())) + ".")
    for w in new.warnings:
        if w not in old.warnings:
            preview.append(f"Warning: {w}.")
    return Plan(ops, preview)


# --------------------------------------------------------------------------
# Package writer: a whole workbook from a layout (openpyxl)
# --------------------------------------------------------------------------

TEXT = "404040"
FONT = "Segoe UI"


def _formats():
    from openpyxl.styles import Border, Font, PatternFill, Side
    return {
        "body": Font(name=FONT, size=9, color=TEXT),
        "bold": Font(name=FONT, size=9, color=TEXT, bold=True),
        "heading": Font(name=FONT, size=10, color="FFFFFF", bold=True),
        "title": Font(name=FONT, size=10, color=TEXT, bold=True),
        "check": Font(name=FONT, size=9, color="9C0006"),
        "input_fill": PatternFill("solid", fgColor="FFF2CC"),
        "heading_fill": PatternFill("solid", fgColor="44546A"),
        "sub_fill": PatternFill("solid", fgColor="E7E6E6"),
        "top": Border(top=Side(style="thin", color=TEXT)),
    }


NUMBER_FORMATS = {"$": '#,##0.00;(#,##0.00);"-"', "%": "0.0%", "flag": '0;-0;"-"'}


def write_workbook(layout: Layout, path: Path, model: Model | None = None) -> Path:
    from openpyxl import Workbook
    from openpyxl.workbook.defined_name import DefinedName

    fm = _formats()
    wb = Workbook()
    wb.remove(wb.active)
    pos = layout.positions()
    for s, rows in layout.sheets:
        ws = wb.create_sheet(s)
        _frame(ws, s, layout.periods, fm)
        for k, r in enumerate(rows):
            rownum = FIRST_ROW + k
            _write_row(ws, rownum, r, row_cells(layout, s, r, rownum, pos), fm, layout.periods)
    for nm, rid in layout.names.items():
        s, r = pos[rid]
        ref = f"{_sheet_prefix(s, 'excel')}${col_letter(TOTAL_COL)}${r}"
        wb.defined_names[nm] = DefinedName(nm, attr_text=ref)
    path = Path(path)
    wb.save(path)
    if model is not None:
        write_metadata(path, model, layout)
    return path


def frame_cells(sheet: str, periods: int) -> dict[tuple[int, int], object]:
    """Header cells every sheet carries, written by both writers."""
    cells = {(1, 2): sheet, (2, 2): "Assembly proof (demo data)"}
    if sheet != CONTENTS:
        cells[(PERIOD_ROW, 2)] = "Month"
        cells[(PERIOD_ROW, TOTAL_COL)] = "Total"
        for p in range(periods):
            cells[(PERIOD_ROW, FIRST_PERIOD_COL + p)] = p + 1
    return cells


def _frame(ws, sheet: str, periods: int, fm) -> None:
    ws.sheet_view.showGridLines = False
    for (r, c), v in frame_cells(sheet, periods).items():
        ws.cell(r, c, v).font = fm["title"] if r == 1 else (fm["body"] if r == 2 else fm["bold"])
    if sheet != CONTENTS:
        ws.freeze_panes = ws.cell(PERIOD_ROW + 1, FIRST_PERIOD_COL)
    for c in range(2, 7):
        ws.column_dimensions[col_letter(c)].width = 2.5
    ws.column_dimensions["A"].width = 2.5
    ws.column_dimensions["G"].width = 34
    ws.column_dimensions["H"].width = 14 if sheet == CONTENTS else 6
    ws.column_dimensions["I"].width = 12
    for p in range(periods):
        ws.column_dimensions[col_letter(FIRST_PERIOD_COL + p)].width = 10
    for r in range(1, FIRST_ROW):
        ws.row_dimensions[r].height = 15


def _write_row(ws, rownum: int, r: LRow, cells: dict, fm, periods: int) -> None:
    end = FIRST_PERIOD_COL + periods
    ws.row_dimensions[rownum].height = 15
    for c, v in cells.items():
        cell = ws.cell(rownum, c, v)
        cell.font = fm["body"]
        if c >= TOTAL_COL and r.unit in NUMBER_FORMATS:
            cell.number_format = NUMBER_FORMATS[r.unit]
    if r.kind == "heading":
        for c in range(2, end):
            ws.cell(rownum, c).fill = fm["heading_fill"]
        ws.cell(rownum, 2).font = fm["heading"]
    elif r.kind in ("subheading", "section"):
        ws.cell(rownum, LABEL_COLS[min(r.indent, 2)]).font = fm["bold"]
        if r.kind == "subheading":
            for c in range(3, end):
                ws.cell(rownum, c).fill = fm["sub_fill"]
    elif r.kind == "setting":
        ws.cell(rownum, TOTAL_COL).fill = fm["input_fill"]
        ws.cell(rownum, TOTAL_COL).font = fm["body"]
    if r.style == "total":
        for c in cells:
            ws.cell(rownum, c).font = fm["bold"]
            if c >= TOTAL_COL:
                ws.cell(rownum, c).border = fm["top"]
    if r.style == "check":
        for c in cells:
            ws.cell(rownum, c).font = fm["check"]


# --------------------------------------------------------------------------
# Model metadata in a custom XML part (openpyxl drops these, so zipfile)
# --------------------------------------------------------------------------

def _meta_xml(model: Model, layout: Layout) -> str:
    from xml.sax.saxutils import escape
    payload = {"model": model.to_dict(), "records": layout.records,
               "rows": {s: [r.id for r in rows] for s, rows in layout.sheets}}
    return (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
            f'<hfgModel xmlns="{META_NS}" version="0">{escape(json.dumps(payload, separators=(",", ":")))}</hfgModel>')


def write_metadata(path: Path, model: Model, layout: Layout) -> None:
    import shutil
    import tempfile
    import zipfile

    path = Path(path)
    with zipfile.ZipFile(path) as z:
        parts = {n: z.read(n) for n in z.namelist()}
    for n in [n for n in parts if n.startswith("customXml/")]:
        del parts[n]
    parts["customXml/item1.xml"] = _meta_xml(model, layout).encode()
    parts["customXml/itemProps1.xml"] = (
        '<?xml version="1.0" encoding="UTF-8" standalone="no"?>\n'
        '<ds:datastoreItem ds:itemID="{6B1F3C2A-0D7E-4C55-9A61-2F0E5B7C8D11}" '
        'xmlns:ds="http://schemas.openxmlformats.org/officeDocument/2006/customXml">'
        f'<ds:schemaRefs><ds:schemaRef ds:uri="{META_NS}"/></ds:schemaRefs></ds:datastoreItem>').encode()
    parts["customXml/_rels/item1.xml.rels"] = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/customXmlProps" '
        'Target="itemProps1.xml"/></Relationships>').encode()
    ct = parts["[Content_Types].xml"].decode()
    if "customXml/itemProps1.xml" not in ct:
        ct = ct.replace("</Types>", '<Override PartName="/customXml/itemProps1.xml" '
                        'ContentType="application/vnd.openxmlformats-officedocument.customXmlProperties+xml"/></Types>')
    if 'Extension="xml"' not in ct:
        ct = ct.replace("</Types>", '<Default Extension="xml" ContentType="application/xml"/></Types>')
    parts["[Content_Types].xml"] = ct.encode()
    rels = parts["xl/_rels/workbook.xml.rels"].decode()
    rels = re.sub(r'<Relationship [^>]*customXml/item1\.xml"[^>]*/>', "", rels)
    rels = rels.replace("</Relationships>", '<Relationship Id="rIdHfgMeta" '
                        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/customXml" '
                        'Target="../customXml/item1.xml"/></Relationships>')
    parts["xl/_rels/workbook.xml.rels"] = rels.encode()
    fd, tmp = tempfile.mkstemp(suffix=".xlsx")
    import os
    os.close(fd)
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as z:
        for n, data in parts.items():
            z.writestr(n, data)
    shutil.move(tmp, path)


def read_metadata(path: Path) -> dict:
    import zipfile
    from xml.etree import ElementTree as ET
    with zipfile.ZipFile(path) as z:
        for n in z.namelist():
            if n.startswith("customXml/item") and not n.startswith("customXml/itemProps") and n.endswith(".xml"):
                root = ET.fromstring(z.read(n))
                if root.tag == f"{{{META_NS}}}hfgModel":
                    return json.loads(root.text)
    raise AssemblyError(f"{path} has no HFG model metadata")


def open_model(path: Path, library: Library) -> tuple[Model, dict]:
    meta = read_metadata(path)
    return Model.from_dict(library, meta["model"]), meta
