"""Saved versions: budgets and monthly reforecasts kept as values, and the comparisons that read them.

Phase 0 proof for the spec's Versions feature. Kelvin (6 October 2026): build a budget and save it
as Modano does, and add monthly reforecasts as the comparisons in budget versus actual models.

Modano keeps one budget as typed values in a budget module and treats the live model as the
reforecast, so last month's reforecast is gone once the model rolls forward. Here every saved
version is a full copy of the statements as values, in a version store:

- Versions (register): one row per version with its type (Budget, Reforecast or Other), label,
  the year a budget is for, the month a reforecast is as at, status, lock, source, when it was
  saved, where its rows sit in the store, and a checksum taken when it was saved.
- Version store: the values, one row per version and statement line, keyed "id|line", for every
  month of the timeline. Lines are keyed, so a version saved before a category was added still
  reads correctly and a new line shows blank for older versions.

Budget versus actual modules get a "Compared with" selection: the approved budget for the year,
last month's reforecast, the latest reforecast, the budget still being built, or any saved
version by name. The Version comparison module (Dashboards) compares the month, year to date and
full year against two versions and charts the full-year outturn by reforecast, what each
reforecast expected for the month shown, and the walk from the comparison to the outturn.

Everything reads the store with INDEX and MATCH, so the workbook works without the add-in.
Saving is the add-in's job: save_live() writes a version through LibreOffice as Save version
will through Office.js; the package writer here writes the demo's history from the reference.
"""

from __future__ import annotations

import datetime as dt

from openpyxl.chart import BarChart, LineChart
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

REG_SHEET, STORE_SHEET, CMP_SHEET = "Versions", "Version store", "Version comparison"
CAP_VERSIONS = 60                  # register rows written with formulas; Save version extends the names when full
CAP_ROWS = 3000                    # store rows covered by Ver_Keys and Ver_Values
REG_FIRST = 7                      # first register row
STORE_FIRST = 10                   # first stored row
FIXED_CHOICES = ["Budget", "Last month's reforecast", "Latest reforecast", "Budget being built"]
COST_LINES = {"cogs", "sal", "opx", "netopex", "dep", "int", "tax"}
VAR = '+#,##0;-#,##0;"-"'
RECIPES = ("trend", "versions", "accuracy", "walk")          # the Version comparison module's charts (C96 to C99)

# Fictional history. Each reforecast's view of the months still to come, as factors on the smooth
# trend: (revenue, with cost of sales following it; overheads, salaries and other operating expenses).
# The September 2026 reforecast is the model as it stands now.
REVISIONS = {13: (0.95, 1.03), 14: (0.96, 1.03), 15: (0.97, 1.02), 16: (0.975, 1.02), 17: (0.985, 1.015),
             18: (0.99, 1.01), 19: (1.0, 1.01), 20: (1.01, 1.0), 21: (1.0, 1.0), 22: (0.995, 1.0), 23: (1.0, 0.995),
             24: (0.97, 1.02), 25: (0.96, 1.02), 26: (0.975, 1.015), 27: (0.99, 1.01), 28: (1.015, 1.0),
             29: (1.005, 1.0), 30: (1.0, 1.0)}
# Budgets: saved as at the March before their year, approved; (year number, seasonality history years included then).
BUDGETS = {12: (2, (1, 0)), 24: (3, (1, 1))}


# ------------------------------------------------------------------------------------ keys and labels
def store_keys(B) -> list[str]:
    """Lines a forecast version keeps: every statement line except headings and the balance check."""
    return [k for k in B.BLOCK if not k.startswith("@") and k != "bal_chk"]


def budget_keys(B) -> list[str]:
    return [k for k, _ in B.BUDGET_LINES]


def month_end(B, t: int) -> dt.date:
    y, m = B.FIRST_FY - 1 + (3 + t - 1) // 12, (3 + t - 1) % 12 + 1
    nxt = dt.date(y + (m == 12), m % 12 + 1, 1)
    return nxt - dt.timedelta(days=1)


def cycle(t: int) -> str:
    n = (t - 1) % 12 + 1
    return f"{n}+{12 - n}"


def suggest_label(B, kind: str, asat: int, year: int | None = None) -> str:
    """The label Save version suggests: "Budget FY2027" or "Reforecast Sep 2026 (6+6)"."""
    if kind == "Budget":
        return f"Budget {B.fy_label(year)}"
    end = month_end(B, asat)
    return f"Reforecast {end:%b} {end.year} ({cycle(asat)})"


# ------------------------------------------------------------------------------------ reference
def inputs_at(B, inp: dict, trend: dict, k: int) -> dict:
    """The inputs as they stood when the reforecast as at month k was saved: actuals to k, then that month's forecast."""
    f_rev, f_ovh = REVISIONS[k]
    out = dict(inp)
    for g, f in (("rev", f_rev), ("cogs", f_rev), ("sal", f_ovh), ("opx", f_ovh)):
        out[g] = [[v if t < k else round(tr[t] * f, 1) for t, v in enumerate(series)]
                  for series, tr in zip(inp[g], trend[g])]
    return out


def history(B, inp: dict) -> list[dict]:
    """The demo's saved versions in the order they were saved, with their values from the reference."""
    trend = B.demo_inputs(last_actual=0)
    out = []
    for k in range(12, B.LAST_ACTUAL + 1):
        if k in REVISIONS:
            st = B.reference(inputs_at(B, inp, trend, k), 1, last_actual=k)
            out.append({"kind": "Reforecast", "asat": k, "year": None, "source": "Forecast (Base)",
                        "values": {key: st[key] for key in store_keys(B)}})
        if k in BUDGETS:
            y, inc = BUDGETS[k]
            bud = B.reference_budget(inp, include=inc)
            out.append({"kind": "Budget", "asat": k, "year": y, "source": "Budget being built",
                        "values": {key: bud[key] for key in budget_keys(B)}})
    first = 1
    latest = max(v["asat"] for v in out if v["kind"] == "Reforecast")
    for i, v in enumerate(out, start=1):
        v["id"] = i
        v["label"] = suggest_label(B, v["kind"], v["asat"], v["year"])
        v["status"] = "Approved" if v["kind"] == "Budget" else "Saved"
        v["locked"] = "No" if (v["kind"] == "Reforecast" and v["asat"] == latest) else "Yes"
        end = month_end(B, v["asat"])
        v["saved_on"] = (end.replace(day=20) if v["kind"] == "Budget" else end + dt.timedelta(days=12))
        v["first"], v["rows"] = first, len(v["values"])
        v["checksum"] = checksum(B, v["values"].values())
        first += v["rows"]
    return out


def checksum(B, rows) -> float:
    """Sum of each stored value's size times its column number, so a value changed or moved between months shows."""
    return sum(abs(x) * (B.FIRST_COL + t) for vals in rows for t, x in enumerate(vals))


def find(hist: list[dict], kind: str, asat: int | None = None, year: int | None = None) -> dict:
    return next(v for v in hist if v["kind"] == kind and (asat is None or v["asat"] == asat) and (year is None or v["year"] == year))


# ------------------------------------------------------------------------------------ package writer
REG_COLS = [("Id", 5), ("Type", 11), ("Label", 30), ("For year", 9), ("As at", 8), ("Status", 11), ("Locked", 7),
            ("Source", 18), ("Saved on", 12), ("First row", 9), ("Rows", 7), ("Checksum when saved", 20),
            ("Checksum now", 16), ("Year number", 12), ("As at month", 12), ("Lookup key", 18)]
SUM_FMT = '#,##0.00;-#,##0.00;""'



def versions_sheet(wb, B, hist: list[dict]):
    """The register: what each version is, when it was saved, and whether its values have changed since."""
    ws = B.frame_sheet(wb, REG_SHEET, "Saved budgets and reforecasts: what each version is, when it was saved, and whether it has changed")
    for j, (_, w) in enumerate(REG_COLS):
        ws.column_dimensions[L(2 + j)].width = w
    B.section(ws, 5, "Saved versions")
    for j, (h, _) in enumerate(REG_COLS):
        B.put(ws, 6, 2 + j, h, B.F_BOLD, border=B.UNDER, align=Alignment(horizontal="right") if 8 <= j <= 14 else None)
    last = REG_FIRST + CAP_VERSIONS - 1
    by_row = {REG_FIRST + i: v for i, v in enumerate(hist)}
    for r in range(REG_FIRST, last + 1):
        v = by_row.get(r)
        if v:
            values = [v["id"], v["kind"], v["label"], B.fy_label(v["year"]) if v["year"] else None, B.month_label(v["asat"]),
                      v["status"], v["locked"], v["source"], v["saved_on"], v["first"], v["rows"], v["checksum"]]
            fmts = ["0", None, None, None, None, None, None, None, "d mmm yyyy", "0", "0", SUM_FMT]
            for j, (x, f) in enumerate(zip(values, fmts)):
                B.put(ws, r, 2 + j, x, fmt=f)
        block = f"INDEX(Ver_Values,$K{r},0):INDEX(Ver_Values,$K{r}+$L{r}-1,0)"
        B.put(ws, r, 14, f'=IF($B{r}="",0,SUMPRODUCT(ABS({block})*COLUMN({block})))', fmt=SUM_FMT)
        B.put(ws, r, 15, f'=IF($E{r}="","",MATCH($E{r},LU_Years,0))', fmt="0")
        B.put(ws, r, 16, f'=IF($F{r}="","",MATCH($F{r},LU_Months,0))', fmt="0")
        B.put(ws, r, 17, f'=IF($B{r}="","",IF($C{r}="Budget",IF($G{r}="Approved","Budget|"&$O{r},""),'
                         f'IF($C{r}="Reforecast","Reforecast|"&$P{r},"")))', align=Alignment(indent=1))
    for j, nm in enumerate(["Ver_Ids", "Ver_Types", "Ver_Labels", None, None, "Ver_Status", "Ver_Locked", None, None, "Ver_First",
                            "Ver_Rows", "Ver_Sum_Saved", "Ver_Sum_Now", "Ver_Year", "Ver_Asat", "Ver_Key"]):
        if nm:
            B.define(wb, nm, B.absref(ws, 2 + j, REG_FIRST, 2 + j, last))
    r = last + 2
    B.put(ws, r, 3, "Versions saved", B.F_BOLD)
    B.put(ws, r, 9, "=COUNT(Ver_Ids)", B.F_BOLD, "0")
    B.define(wb, "Ver_Count", B.absref(ws, 9, r))
    B.put(ws, r + 1, 3, "Latest reforecast is as at month", B.F_BOLD)
    B.put(ws, r + 1, 9, '=_xlfn.MAXIFS(Ver_Asat,Ver_Types,"Reforecast")', B.F_BOLD, "0")
    B.put(ws, r + 1, 10, '=IF(I{0}=0,"none saved",INDEX(LU_Months,I{0}))'.format(r + 1))
    B.define(wb, "Ver_Latest_Asat", B.absref(ws, 9, r + 1))
    B.put(ws, r + 2, 3, "Next free row in the store", B.F_BOLD)
    B.put(ws, r + 2, 9, "=IF(Ver_Count=0,1,SUMPRODUCT(MAX(Ver_First+Ver_Rows)))", B.F_BOLD, "0")
    B.define(wb, "Ver_Next_Row", B.absref(ws, 9, r + 2))
    notes = ["Save version (add-in) writes a version's values to the Version store and adds its row here; nothing here is typed by hand.",
             "A budget is approved for one year; the reports' Budget choice reads the approved budget for the year shown.",
             "A reforecast is saved after each month's close, as at the last actual month (6+6 is six months of actuals and six of forecast).",
             "Locked versions cannot be replaced or deleted. The checksum shows whether stored values have changed since they were saved."]
    for i, text in enumerate(notes):
        B.put(ws, r + 4 + i, 3, text, B.F_NOTE)
    ws.freeze_panes = f"C{REG_FIRST}"
    return ws


def store_sheet(wb, B, hist: list[dict]):
    """The values of every saved version, one row per version and line, keyed "id|line"."""
    ws = B.frame_sheet(wb, STORE_SHEET, "Every saved version's values by line and month; written by Save version, never edited by hand",
                       timeline=True)
    for c, w in (("C", 5), ("D", 9), ("E", 12), ("F", 2.5), ("G", 30)):
        ws.column_dimensions[c].width = w
    B.section(ws, 8, "Stored values")
    for c, h in ((3, "Id"), (4, "Line"), (5, "Key"), (7, "Line label"), (8, "Unit")):
        B.put(ws, 9, c, h, B.F_BOLD)
    r = STORE_FIRST
    for v in hist:
        for key, vals in v["values"].items():
            B.put(ws, r, 3, v["id"], fmt="0")
            B.put(ws, r, 4, key)
            B.put(ws, r, 5, f"{v['id']}|{key}")
            B.put(ws, r, 7, B.block_label(key))
            B.put(ws, r, 8, "%" if key == "gm_pct" else "$000")
            for t, x in enumerate(vals):
                ws.cell(r, B.FIRST_COL + t, x).number_format = B.PCT if key == "gm_pct" else B.NUM
            r += 1
    last = STORE_FIRST + CAP_ROWS - 1
    B.define(wb, "Ver_Keys", B.absref(ws, 5, STORE_FIRST, 5, last))
    B.define(wb, "Ver_Values", B.absref(ws, B.FIRST_COL, STORE_FIRST, B.LAST_COL, last))
    ws.freeze_panes = f"J{STORE_FIRST}"
    return ws


def lookups(wb, B, ws) -> None:
    """Lists behind the comparison drop-downs: the fixed choices then every saved version's label."""
    ws.column_dimensions["L"].width = 28
    ws.column_dimensions["N"].width = 26
    ws.column_dimensions["O"].width = 9
    B.put(ws, 5, 12, "Compared with", B.F_BOLD)
    for i, text in enumerate(FIXED_CHOICES):
        B.put(ws, 6 + i, 12, text)
    first = 6 + len(FIXED_CHOICES)
    for k in range(1, CAP_VERSIONS + 1):
        B.put(ws, first + k - 1, 12, f'=IF(INDEX(Ver_Labels,{k})="","",INDEX(Ver_Labels,{k}))')
    last = first + CAP_VERSIONS - 1
    B.define(wb, "LU_Cmp_Fixed", B.absref(ws, 12, 6, 12, first - 1))
    B.define(wb, "LU_Compare", f"{B.q(ws)}!$L$6:INDEX({B.q(ws)}!$L$6:$L${last},{len(FIXED_CHOICES)}+Ver_Count)")
    B.put(ws, 5, 14, "Version lines", B.F_BOLD)
    B.put(ws, 5, 15, "Key", B.F_BOLD)
    for i, (key, label) in enumerate(B.BUDGET_LINES):
        B.put(ws, 6 + i, 14, label)
        B.put(ws, 6 + i, 15, key)
    n = len(B.BUDGET_LINES)
    B.define(wb, "LU_Ver_Lines", B.absref(ws, 14, 6, 14, 5 + n))
    B.define(wb, "LU_Ver_Keys", B.absref(ws, 15, 6, 15, 5 + n))


# ------------------------------------------------------------------------------------ comparisons
def resolve(sel: str, year: str, ref: str) -> str:
    """Formula for the version id a selection means: -1 is the budget being built, #N/A is none saved."""
    before = 'MAXIFS(Ver_Asat,Ver_Types,"Reforecast",Ver_Asat,"<"&' + ref + ")"
    latest = 'MAXIFS(Ver_Asat,Ver_Types,"Reforecast")'
    rf = lambda m: 'INDEX(Ver_Ids,MATCH("Reforecast|"&_xlfn.' + m + ",Ver_Key,0))"      # noqa: E731
    return ("=IFERROR(CHOOSE(MATCH(" + sel + ",LU_Cmp_Fixed,0),"
            + 'INDEX(Ver_Ids,MATCH("Budget|"&' + year + ",Ver_Key,0)),"
            + rf(before) + "," + rf(latest) + ",-1),"
            + "INDEX(Ver_Ids,MATCH(" + sel + ",Ver_Labels,0)))")


def compare_block(wb, B, ws, code: str, row: int, *, year: str, ref: str, suffix: str = "Cmp",
                  label: str = "Compared with", default: str = "Budget", title: str = "") -> list:
    """A "Compared with" drop-down, the id it resolves to and its label; returns the module's checks for it."""
    B.put(ws, row, 3, label, B.F_BOLD)
    B.put(ws, row, 8, default, fill=B.FILL_IN, align=Alignment(horizontal="right"))
    sel = f"DD_{code}_{suffix}"
    B.define(wb, sel, B.absref(ws, 8, row))
    dv = DataValidation(type="list", formula1="LU_Compare", allow_blank=False)
    dv.error, dv.errorTitle = "Pick a choice or a saved version from the list", label
    ws.add_data_validation(dv)
    dv.add(f"H{row}")
    B.put(ws, row, 9, resolve(sel, year, ref), fmt="0")
    idn = f"{code}_{suffix}"
    B.define(wb, idn, B.absref(ws, 9, row))
    B.put(ws, row, 10, f'=IF(ISERROR({idn}),"none saved",IF({idn}=-1,"Budget being built (not saved)",'
                       f'INDEX(Ver_Labels,MATCH({idn},Ver_Ids,0))))', B.F_NOTE)
    B.define(wb, f"{idn}_Label", B.absref(ws, 10, row))
    return [(f"{title}: {label.lower()} is not a listed choice or a saved version",
             f"=IF(ISNA(MATCH({sel},LU_Compare,0)),1,0)", "error"),
            (f"{title}: nothing is saved for {label.lower()} (no approved budget for the year, or no reforecast yet), so it is blank",
             f"=IF(ISNA(MATCH({sel},LU_Compare,0)),0,IF(ISERROR({idn}),1,0))", "alert")]


def cmp_value(idn: str, key: str, i: str, live: str) -> str:
    """A comparison's value for a line (key is a quoted literal or a name) at month number i."""
    return (f"IF(OR({i}<1,{i}>Ts_Periods),NA(),IF({idn}=-1,{live},"
            f'IFERROR(INDEX(Ver_Values,MATCH({idn}&"|"&{key},Ver_Keys,0),{i}),NA())))')


def cmp_span(idn: str, key: str, a: str, b: str, live_nm: str | None) -> str:
    """A comparison's total for a line from month a to month b."""
    row = f'MATCH({idn}&"|"&{key},Ver_Keys,0)'
    live = f"SUM(INDEX({live_nm},{a}):INDEX({live_nm},{b}))" if live_nm else "NA()"
    return f"IF({idn}=-1,{live},IFERROR(SUM(INDEX(Ver_Values,{row},{a}):INDEX(Ver_Values,{row},{b})),NA()))"


def stored(id_cell: str, key: str, a: str, b: str | None = None) -> str:
    """A saved version's value at month a (or its total from a to b), blank when the slot has no version."""
    row = f'MATCH({id_cell}&"|"&{key},Ver_Keys,0)'
    val = f"INDEX(Ver_Values,{row},{a})" if b is None else f"SUM(INDEX(Ver_Values,{row},{a}):INDEX(Ver_Values,{row},{b}))"
    return f'=IF({id_cell}="","",IFERROR({val},""))'


def model_checks() -> list[tuple[str, str, str]]:
    return [("A saved version's values have changed since it was saved (checksum on Versions)",
             "=IF(SUMPRODUCT(ABS(Ver_Sum_Now-Ver_Sum_Saved))>0.01,1,0)", "error"),
            ("Two versions answer the same question: two approved budgets for a year, or two reforecasts as at one month",
             '=IF(SUMPRODUCT((Ver_Key<>"")*(COUNTIF(Ver_Key,Ver_Key)>1))>0,1,0)', "error"),
            ("This month's reforecast is not saved yet: the latest reforecast is before the last actual month",
             "=IF(Ver_Latest_Asat<Ts_Last_Actual,1,0)", "alert")]


# ------------------------------------------------------------------------------------ Version comparison module
def _live(B, prefix: str) -> str:
    """The line shown as a reference into the statements (St) or the budget being built (Bud)."""
    return "CHOOSE(VerS_Line," + ",".join(B.nm_of(prefix, k) for k in budget_keys(B)) + ")"


def comparison_sheet(wb, B, module: dict, charts: list[dict]):
    """Month, year to date and full year against two versions, and four charts that come with the module."""
    code, title = module["code"], module["title"]
    ws = wb.create_sheet(title)
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 90
    for c in range(1, 60):
        ws.column_dimensions[L(c)].width = B.REPORT_WIDTHS.get(L(c), B.VALUE_WIDTH)
    B.put(ws, 1, 2, title, B.F_HEAD)
    B.put(ws, 2, 2, "The month, the year to date and the full year against two saved versions; the outturn is actuals to date then "
                    "the current forecast for the active scenario")
    B.put(ws, 3, 2, "Demo Building Co (fictional numbers, $000)", B.F_NOTE)
    rep = B.Report(wb, ws, code, title, False, True, False)
    checks = rep.checks
    # Selections
    B.put(ws, 4, 3, "Line shown", B.F_BOLD)
    B.put(ws, 4, 8, B.BUDGET_LINES[0][1], fill=B.FILL_IN, align=Alignment(horizontal="right"))
    B.define(wb, "DD_VerS_Line", B.absref(ws, 8, 4))
    B.put(ws, 4, 9, "=MATCH(H4,LU_Ver_Lines,0)", fmt="0")
    B.define(wb, "VerS_Line", B.absref(ws, 9, 4))
    B.put(ws, 4, 10, "=INDEX(LU_Ver_Keys,VerS_Line)", B.F_NOTE)
    B.define(wb, "VerS_Key", B.absref(ws, 10, 4))
    dv = DataValidation(type="list", formula1="LU_Ver_Lines", allow_blank=False)
    dv.error, dv.errorTitle = "Pick a line from the list", "Line shown"
    ws.add_data_validation(dv)
    dv.add("H4")
    B.put(ws, 5, 3, "Month shown", B.F_BOLD)
    B.put(ws, 5, 8, B.month_label(B.MONTH_SHOWN), fill=B.FILL_IN, align=Alignment(horizontal="right"))
    B.define(wb, "DD_VerS_Month", B.absref(ws, 8, 5))
    B.put(ws, 5, 9, "=MATCH(H5,LU_Months,0)", fmt="0")
    B.define(wb, "VerS_Month", B.absref(ws, 9, 5))
    dv = DataValidation(type="list", formula1="LU_Months", allow_blank=False)
    dv.error, dv.errorTitle = "Pick a month from the list", "Month shown"
    ws.add_data_validation(dv)
    dv.add("H5")
    B.put(ws, 5, 10, '=IF(VerS_Month<=Ts_Last_Actual,"actual","forecast")', B.F_NOTE)
    year = "VerS_FY"
    checks += compare_block(wb, B, ws, code, 6, year=year, ref="VerS_Month", title=title)
    checks += compare_block(wb, B, ws, code, 7, year=year, ref="VerS_Month", suffix="Cmp2", label="Second comparison",
                            default=FIXED_CHOICES[1], title=title)
    checks.append((f"{title}: the line shown is not in the list", "=IF(ISERROR(VerS_Line),1,0)", "error"))
    checks.append((f"{title}: the month shown is not in the list", "=IF(ISERROR(VerS_Month),1,0)", "error"))

    # Variance table
    B.section(ws, 9, "Variance against the two comparisons (favourable when positive)", toc="Variance")
    groups = [("month", '="Month: "&DD_VerS_Month'), ("ytd", '="Year to date: "&INDEX(LU_Months,VerS_YTD_Start)&" to "&DD_VerS_Month'),
              ("fy", '="Full year: "&INDEX(LU_Years,VerS_FY)')]
    heads = {"month": '=IF(VerS_Month<=Ts_Last_Actual,"Actual","Forecast")',
             "ytd": '=IF(VerS_Month<=Ts_Last_Actual,"Actual","To date")', "fy": "Outturn"}
    right = Alignment(horizontal="right")
    for g, (key, text) in enumerate(groups):
        c0 = B.FIRST_COL + 5 * g
        B.put(ws, 10, c0, text, B.F_BOLD)
        for j, h in enumerate([heads[key], "Compared", "Variance", "Second", "Variance"]):
            B.put(ws, 11, c0 + j, h, B.F_BOLD, border=B.UNDER, align=right)
    B.put(ws, 11, 7, "Line", B.F_BOLD, border=B.UNDER)
    B.put(ws, 11, 8, "", border=B.UNDER)
    m, ys, ye = "VerS_Month", "VerS_YTD_Start", "VerS_Y_End"
    top = 12
    for n, (key, label) in enumerate(B.BUDGET_LINES):
        r = top + n
        bold = key in ("rev", "gm", "ebitda", "npat")
        B.put(ws, r, 7, label, B.F_BOLD if bold else B.F_BODY)
        B.put(ws, r, 8, "$000")
        st, bud = B.nm_of("St", key), B.nm_of("Bud", key)
        k = f'"{key}"'
        cells = {"month": (f"INDEX({st},{m})", [cmp_value(f"{code}_{s}", k, m, f"INDEX({bud},{m})") for s in ("Cmp", "Cmp2")]),
                 "ytd": (f"SUM(INDEX({st},{ys}):INDEX({st},{m}))", [cmp_span(f"{code}_{s}", k, ys, m, bud) for s in ("Cmp", "Cmp2")]),
                 "fy": (f"SUM(INDEX({st},{ys}):INDEX({st},{ye}))", [cmp_span(f"{code}_{s}", k, ys, ye, bud) for s in ("Cmp", "Cmp2")])}
        for g, (gk, _) in enumerate(groups):
            c0 = B.FIRST_COL + 5 * g
            act, (c1, c2) = cells[gk]
            B.put(ws, r, c0, "=" + act, B.F_BOLD if bold else B.F_BODY, B.NUM)
            for j, cf in ((1, c1), (3, c2)):
                a, cc = f"{L(c0)}{r}", f"{L(c0 + j)}{r}"
                B.put(ws, r, c0 + j, "=" + cf, B.F_BOLD if bold else B.F_BODY, B.NUM)
                var = f"={cc}-{a}" if key in COST_LINES else f"={a}-{cc}"
                B.put(ws, r, c0 + j + 1, var, B.F_BOLD if bold else B.F_BODY, VAR)
    bottom = top + len(B.BUDGET_LINES) - 1
    for g in range(3):
        for j in (2, 4):
            col = L(B.FIRST_COL + 5 * g + j)
            rng = f"{col}{top}:{col}{bottom}"
            ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThan", formula=["0.5"], font=Font(color=B.GOOD)))
            ws.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=["-0.5"], font=Font(color=B.BAD)))
    B.put(ws, bottom + 1, 7, '="Compared: "&VerS_Cmp_Label&". Second: "&VerS_Cmp2_Label&". For costs, favourable is spending less."',
          B.F_NOTE)

    # Chart grid and chart data
    grid_top = bottom + 5
    B.section(ws, grid_top - 1, "Charts")
    grid_rows = -(-len(charts) // B.GRID_COLS)
    r = grid_top + grid_rows * B.GRID_ROWS + 1
    B.section(ws, r, "Chart data: every number is a formula on the statements or the version store; the charts read these rows",
              toc="Chart data")
    r += 2
    helpers = [("Year of the month shown (year number)", "=INT((VerS_Month-1)/12)+1", "VerS_FY"),
               ("Year of the month shown starts at month", "=(VerS_FY-1)*12+1", "VerS_YTD_Start"),
               ("Year of the month shown ends at month", "=VerS_YTD_Start+11", "VerS_Y_End")]
    for label, f, nm in helpers:
        B.put(ws, r, 7, label)
        B.put(ws, r, 8, f, fmt="0")
        B.define(wb, nm, B.absref(ws, 8, r))
        r += 1
    put_pos = r
    B.put(ws, r, 7, "Position in the year")
    for j in range(12):
        B.put(ws, r, B.FIRST_COL + j, j + 1, fmt="0")
    rep.r_pos = put_pos
    B.put(ws, r + 1, 7, "Month of the financial year")
    for j in range(12):
        B.put(ws, r + 1, B.FIRST_COL + j, f"=INDEX(LU_Month_Names,{L(B.FIRST_COL + j)}${put_pos})")
    rep.r_mname = r + 1
    B.put(ws, r + 2, 7, "Year shown: month number")
    for j in range(12):
        B.put(ws, r + 2, B.FIRST_COL + j, f"=VerS_YTD_Start-1+{L(B.FIRST_COL + j)}${put_pos}", fmt="0")
    rep.r_yidx = r + 2
    rep.row = r + 4
    builders = {"trend": _trend, "versions": _versions, "accuracy": _accuracy, "walk": _walk}
    for k, ch in enumerate(charts):
        chart = builders[ch["recipe"]](B, rep, ch)
        chart.anchor = B.chart_anchor(k, grid_top)
        ws.add_chart(chart)
    for rr in range(1, rep.row + 1):
        ws.row_dimensions[rr].height = 15
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_area = f"A1:{L(max(B.grid_last_col(), B.FIRST_COL + 14))}{grid_top + grid_rows * B.GRID_ROWS}"
    return rep


FY_LABEL = '"Full year "&INDEX(LU_Years,VerS_FY)'
LINE_FY = 'DD_VerS_Line&", "&INDEX(LU_Years,VerS_FY)'


def _trend(B, rep, ch):
    """Full-year outturn by version: the approved budget, each reforecast saved in the year, and the current forecast."""
    t = B.Table(rep, ch, LINE_FY, f"{B.BUDGET_LINES[0][1]}, {B.fy_label(B.YEAR_SHOWN)}")
    slots = ["Budget"] + [cycle(n) for n in range(1, 12)] + ["Current"]
    t.header(slots, "Version")
    ws, n = rep.ws, len(slots)
    ra = t.add("Reforecast as at month", [""] + [f"=VerS_YTD_Start-1+{j}" for j in range(1, 12)] + [""], unit="", fmt="0")
    ids = ['=IFERROR(INDEX(Ver_Ids,MATCH("Budget|"&VerS_FY,Ver_Key,0)),"")']
    ids += [f'=IF({L(B.FIRST_COL + j)}{ra}>VerS_Month,"",IFERROR(INDEX(Ver_Ids,MATCH("Reforecast|"&{L(B.FIRST_COL + j)}{ra},Ver_Key,0)),""))'
            for j in range(1, 12)]
    ri = t.add("Version id", ids + [""], unit="", fmt="0")
    live = _live(B, "St")
    vals = [stored(f"{L(B.FIRST_COL + j)}${ri}", "VerS_Key", "VerS_YTD_Start", "VerS_Y_End")[1:] for j in range(n - 1)]
    vals.append(f"SUM(INDEX({live},VerS_YTD_Start):INDEX({live},VerS_Y_End))")
    rb = t.add("Approved budget", ["=" + vals[0]] + [""] * (n - 1))
    rr = t.add("Reforecasts", [""] + ["=" + v for v in vals[1:-1]] + [""])
    rc = t.add("Current forecast", [""] * (n - 1) + ["=" + vals[-1]], bold=True)
    t.note("Each reforecast's total for the year: actuals to the month it was saved, then its forecast. Blank slots are not saved yet.")
    t.done()
    c = BarChart()
    c.type, c.grouping, c.overlap, c.gapWidth = "col", "stacked", 100, 45
    for r, colour in ((rb, B.PALETTE[0]), (rr, B.GREY), (rc, B.INK)):
        s = t.ser(r)
        B.style_fill(s, colour)
        c.series.append(s)
    B.finish_axes(c)
    c.y_axis.scaling.min = 0
    B.finish_chart(c, t.title(), t.title_text)
    return c


def _versions(B, rep, ch):
    """Actual (solid) and forecast (hatched) by month for the year, with both comparisons as lines."""
    t = B.Table(rep, ch, LINE_FY, f"{B.BUDGET_LINES[0][1]}, {B.fy_label(B.YEAR_SHOWN)}")
    t.header(B.month_cats(rep, "year"), "Series", "Total")
    tot = lambda r: f"=SUM(J{r}:U{r})"                                                     # noqa: E731
    ix = [f"{L(B.FIRST_COL + j)}${rep.r_yidx}" for j in range(12)]
    live, bud = _live(B, "St"), _live(B, "Bud")
    ra = t.add("Actual", [f"=IF(INDEX(Ts_Actual,{i})=1,INDEX({live},{i}),0)" for i in ix], total=tot)
    rf = t.add("Forecast", [f"=IF(INDEX(Ts_Actual,{i})=1,0,INDEX({live},{i}))" for i in ix], total=tot)
    r1 = t.add("=VerS_Cmp_Label", ["=" + cmp_value("VerS_Cmp", "VerS_Key", i, f"INDEX({bud},{i})") for i in ix], total=tot, bold=True)
    r2 = t.add("=VerS_Cmp2_Label", ["=" + cmp_value("VerS_Cmp2", "VerS_Key", i, f"INDEX({bud},{i})") for i in ix], total=tot)
    t.done()
    c = BarChart()
    c.type, c.grouping, c.overlap, c.gapWidth = "col", "stacked", 100, 55
    sa, sf = t.ser(ra), t.ser(rf)
    B.style_fill(sa, B.INK)
    B.style_fill(sf, B.INK, hatch=True)
    c.series += [sa, sf]
    ln = LineChart()
    s1, s2 = t.ser(r1), t.ser(r2)
    B.style_line(s1, B.PALETTE[0], 2.25, None, True)
    B.style_line(s2, B.PALETTE[2], 1.75, "dash", True)
    ln.series += [s1, s2]
    c += ln
    B.finish_axes(c)
    B.finish_chart(c, t.title(), t.title_text)
    return c


def _accuracy(B, rep, ch):
    """What the budget and each of the twelve reforecasts before it expected for the month shown, against the actual."""
    t = B.Table(rep, ch, 'DD_VerS_Line&", "&DD_VerS_Month', f"{B.BUDGET_LINES[0][1]}, {B.month_label(B.MONTH_SHOWN)}")
    n = 13
    asat_row = t.next
    cats = ["Budget"] + [f'=IF({L(B.FIRST_COL + j)}{asat_row}<1,"",INDEX(LU_Months,{L(B.FIRST_COL + j)}{asat_row}))' for j in range(1, n)]
    t.header(cats, "Saved as at")
    ra = t.add("Reforecast as at month", [""] + [f"=VerS_Month-{n - j}" for j in range(1, n)], unit="", fmt="0")
    assert ra == asat_row
    ids = ['=IFERROR(INDEX(Ver_Ids,MATCH("Budget|"&VerS_FY,Ver_Key,0)),"")']
    ids += [f'=IF({L(B.FIRST_COL + j)}{ra}<1,"",IFERROR(INDEX(Ver_Ids,MATCH("Reforecast|"&{L(B.FIRST_COL + j)}{ra},Ver_Key,0)),""))'
            for j in range(1, n)]
    ri = t.add("Version id", ids, unit="", fmt="0")
    vals = [stored(f"{L(B.FIRST_COL + j)}${ri}", "VerS_Key", "VerS_Month") for j in range(n)]
    rb = t.add("Approved budget", [vals[0]] + [""] * (n - 1))
    rr = t.add("Reforecasts", [""] + vals[1:])
    live = _live(B, "St")
    rc = t.add('=IF(VerS_Month<=Ts_Last_Actual,"Actual","Current forecast")', [f"=INDEX({live},VerS_Month)"] * n, bold=True)
    t.note("Each bar is what that version expected for the month shown; the line is what happened (or the current forecast).")
    t.done()
    c = BarChart()
    c.type, c.grouping, c.overlap, c.gapWidth = "col", "stacked", 100, 45
    for r, colour in ((rb, B.PALETTE[0]), (rr, B.GREY)):
        s = t.ser(r)
        B.style_fill(s, colour)
        c.series.append(s)
    ln = LineChart()
    s = t.ser(rc)
    B.style_line(s, B.INK, 2.25)
    ln.series.append(s)
    c += ln
    B.finish_axes(c)
    c.y_axis.scaling.min = 0
    B.finish_chart(c, t.title(), t.title_text)
    return c


def _walk(B, rep, ch):
    """A waterfall from the comparison's profit after tax for the year to the outturn, one step per line's variance."""
    t = B.Table(rep, ch, FY_LABEL, f"Full year {B.fy_label(B.YEAR_SHOWN)}")
    ys, ye = "VerS_YTD_Start", "VerS_Y_End"
    out = lambda k: f"SUM(INDEX({B.nm_of('St', k)},{ys}):INDEX({B.nm_of('St', k)},{ye}))"            # noqa: E731
    cmp = lambda k: cmp_span("VerS_Cmp", f'"{k}"', ys, ye, B.nm_of("Bud", k))                       # noqa: E731
    items = [("=VerS_Cmp_Label", "=" + cmp("npat"), "start")]
    for k, lab in (("rev", "Revenue"), ("cogs", "Cost of sales"), ("sal", "Salaries"), ("opx", "Other expenses"),
                   ("dep", "Depreciation"), ("int", "Interest"), ("tax", "Tax")):
        items.append((lab, f"={cmp(k)}-{out(k)}" if k in COST_LINES else f"={out(k)}-{cmp(k)}", "step"))
    items.append(("Outturn", None, "end"))
    t.header([lab for lab, _, _ in items], "Walk")
    vr = t.next
    rr = vr + 1
    rows = {k: [] for k in ("value", "run", "base", "total", "inc", "dec")}
    for j, (_, value, kind) in enumerate(items):
        c, pc = L(B.FIRST_COL + j), (L(B.FIRST_COL + j - 1) if j else None)
        if kind == "end":
            value = f"={pc}{rr}"
        rows["value"].append(value)
        if kind == "step":
            rows["run"].append(f"={pc}{rr}+{c}{vr}")
            rows["base"].append(f"=MIN({pc}{rr},{c}{rr})")
            rows["total"].append("=0")
            rows["inc"].append(f"=MAX({c}{vr},0)")
            rows["dec"].append(f"=MAX(-{c}{vr},0)")
        else:
            rows["run"].append(f"={c}{vr}")
            rows["base"].append("=0")
            rows["total"].append(f"={c}{vr}")
            rows["inc"].append("=0")
            rows["dec"].append("=0")
    t.add("Value", rows["value"])
    t.add("Running total", rows["run"])
    br, tr, ir, dr = (t.add(label, rows[key]) for key, label in
                      (("base", "Base (not drawn)"), ("total", "Total"), ("inc", "Favourable"), ("dec", "Unfavourable")))
    t.note("Steps are favourable variances (for costs, spending less than the comparison); the walk ends at the outturn.")
    t.done()
    end = f"{B.q(t.ws)}!{L(B.FIRST_COL + len(items) - 1)}{vr}"
    rep.checks.append((f"{ch['id']} {ch['title']}: the walk does not reach the outturn's profit after tax",
                       f"=IF(ISERROR({end}),0,IF(ABS({end}-{out('npat')})>0.001,1,0))", "error"))
    c = BarChart()
    c.type, c.grouping, c.overlap, c.gapWidth = "col", "stacked", 100, 35
    for r, colour, fmt, pos in ((br, None, None, None), (tr, "808080", "#,##0;-#,##0;", "inEnd"),
                                (ir, B.GOOD, '"+"#,##0;;', "ctr"), (dr, B.BAD, '"-"#,##0;;', "ctr")):
        s = t.ser(r)
        B.style_fill(s, colour, outline="404040" if colour else None)
        if fmt:
            s.dLbls = B.labels(fmt, pos)
        c.series.append(s)
    B.finish_axes(c, value_axis=False)
    B.finish_chart(c, t.title(), t.title_text, legend=None)
    return c


# ------------------------------------------------------------------------------------ live writer (Save version)
def _serial(d: dt.date) -> float:
    return float((d - dt.date(1899, 12, 30)).days)


def save_live(doc, B, kind: str, *, label: str | None = None, year: int | None = None, source: str = "forecast",
              status: str | None = None, locked: str = "No", saved_on: dt.date | None = None) -> int:
    """Save version, as the add-in will: read the source block's values, append them to the store, add the register row.

    source "forecast" saves the statements for the active scenario (every line); "budget" saves the budget being built
    (the income statement). Approving a budget supersedes the year's previous approved budget. Returns the new id.
    """
    names = doc.NamedRanges

    def cells(nm):
        return names.getByName(nm).getReferredCells()

    asat = int(round(cells("Ts_Last_Actual").getCellByPosition(0, 0).getValue()))
    keys, prefix = (store_keys(B), "St") if source == "forecast" else (budget_keys(B), "Bud")
    rows = [[float(x) for x in cells(B.nm_of(prefix, k)).getDataArray()[0]] for k in keys]
    reg = doc.Sheets.getByName(REG_SHEET)
    data = reg.getCellRangeByPosition(1, REG_FIRST - 1, 16, REG_FIRST + CAP_VERSIONS - 2).getDataArray()
    used = [r for r in data if r[0] not in ("", None)]
    new_id = (max(int(r[0]) for r in used) + 1) if used else 1
    first = (max(int(r[9] + r[10]) for r in used)) if used else 1
    status = status or ("Approved" if kind == "Budget" else "Saved")
    label = label or suggest_label(B, kind, asat, year)
    if kind == "Budget" and status == "Approved":
        for i, r in enumerate(data):
            if r[0] not in ("", None) and r[1] == "Budget" and r[3] == B.fy_label(year) and r[5] == "Approved":
                reg.getCellByPosition(6, REG_FIRST - 1 + i).setString("Superseded")
    store = doc.Sheets.getByName(STORE_SHEET)
    block = []
    for k, vals in zip(keys, rows):
        block.append((float(new_id), k, f"{new_id}|{k}", "", B.block_label(k), "%" if k == "gm_pct" else "$000", "") + tuple(vals))
    top = STORE_FIRST + first - 1
    store.getCellRangeByPosition(2, top - 1, 2 + len(block[0]) - 1, top - 1 + len(block) - 1).setDataArray(tuple(block))
    entry = (float(new_id), kind, label, B.fy_label(year) if year else "", B.month_label(asat), status, locked,
             "Forecast (" + cells("DD_Scenario").getCellByPosition(0, 0).getString() + ")" if source == "forecast" else "Budget being built",
             _serial(saved_on or dt.date.today()), float(first), float(len(keys)), checksum(B, rows))
    row = REG_FIRST + len(used)
    reg.getCellRangeByPosition(1, row - 1, 12, row - 1).setDataArray((entry,))
    doc.calculateAll()
    return new_id


def stored_live(doc, version_id: int) -> dict[str, list[float]]:
    """A saved version's values by line, read back from the store."""
    store = doc.Sheets.getByName(STORE_SHEET)
    data = store.getCellRangeByPosition(2, STORE_FIRST - 1, 56, STORE_FIRST + CAP_ROWS - 2).getDataArray()
    return {r[1]: list(r[7:]) for r in data if r[0] not in ("", None) and int(r[0]) == version_id}


def set_inputs_live(doc, B, inp: dict) -> None:
    """Write the category inputs (revenue, cost of sales, salaries, other operating expenses) onto Inputs by label."""
    sh = doc.Sheets.getByName("Inputs")
    labels = [r[0] for r in sh.getCellRangeByPosition(2, 0, 2, 200).getDataArray()]
    for g in ("rev", "cogs", "sal", "opx"):
        for k, name in enumerate(B.GROUPS[g][1]):
            r = labels.index(name)
            sh.getCellRangeByPosition(B.FIRST_COL - 1, r, B.LAST_COL - 1, r).setDataArray((tuple(float(x) for x in inp[g][k]),))
