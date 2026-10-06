"""Phase 0 proof: every chart the summary and report modules bring, rebuilt as native charts.

Builds a workbook with a small fictional three-statement model (four financial years, actuals
to September 2026 then forecast, a budget and three scenarios) and the nine summary and report
modules, each with its charts. The charts come from register.yaml: each names a recipe, and
the builder expands the recipe into formula rows on the module's sheet and a native chart over
them. Nothing is typed into a chart table; every number reads the statements through INDEX,
so the charts follow the model and the selections on their sheet (year shown, month shown)
and the active scenario.

    python prototypes/reports/build.py [out.xlsx]

The workbook works without the add-in. Fonts and number formats follow the frame standard;
category colours use the HF chart palette and comparisons follow IBCS (prior grey, actual
solid, forecast hatched).
"""

from __future__ import annotations

import calendar
import os
import random
import re
import shutil
import sys
import tempfile
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

import yaml
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference, Series
from openpyxl.chart.axis import ChartLines
from openpyxl.chart.data_source import AxDataSource, StrData, StrRef, StrVal
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.chart.marker import DataPoint, Marker
from openpyxl.chart.series import SeriesLabel
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText, Text
from openpyxl.chart.title import Title
from openpyxl.drawing.fill import ColorChoice, PatternFillProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.text import CharacterProperties, Font as DrawingFont, Paragraph, ParagraphProperties
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.utils.units import cm_to_EMU, pixels_to_EMU
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.hyperlink import Hyperlink

HERE = Path(__file__).resolve().parent
REGISTER = HERE / "register.yaml"

# Look ---------------------------------------------------------------------------------------
TEXT = "404040"
INK, GREY, LIGHT, GOOD, BAD = "262626", "A6A6A6", "D9D9D9", "8CB400", "FF0000"
PALETTE = ["679DB5", "09122C", "566A89", "90B6C8", "3C4E60", "CEDCE5"]   # HF chart sequence (HFG brand guidelines)
F_BODY = Font(name="Segoe UI", size=9, color=TEXT)
F_BOLD = Font(name="Segoe UI", size=9, color=TEXT, bold=True)
F_HEAD = Font(name="Segoe UI", size=10, color=TEXT, bold=True)
F_NOTE = Font(name="Segoe UI", size=9, color="808080", italic=True)
F_CHECK = Font(name="Segoe UI", size=9, color="9C0006")
F_LINK = Font(name="Segoe UI", size=9, color="0563C1", underline="single")
FILL_IN = PatternFill("solid", fgColor="FFF2CC")
TOP = Border(top=Side(style="thin", color=TEXT))
UNDER = Border(bottom=Side(style="thin", color=TEXT))
NUM = '#,##0;(#,##0);"-"'
PCT = '0.0%'
CHART_W, CHART_H = 12.2, 7.6          # cm
GRID_ROWS = 15                         # sheet rows per row of charts (rows are 15 points high)
GRID_COLS = 3

# Timeline -----------------------------------------------------------------------------------
FIRST_FY = 2025                        # FY2025 runs April 2024 to March 2025
YEARS = 4
PERIODS = 12 * YEARS
LAST_ACTUAL = 30                       # September 2026
YEAR_SHOWN = 3                         # FY2027
MONTH_SHOWN = 30                       # Sep 26
FIRST_COL = 10                         # column J: the first month on time series sheets
LAST_COL = FIRST_COL + PERIODS - 1
MONTH_NAMES = ["Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec", "Jan", "Feb", "Mar"]
TAX_RATE = 0.28


def month_label(t: int) -> str:
    m = (3 + t - 1) % 12 + 1
    y = FIRST_FY - 1 + (3 + t - 1) // 12
    return f"{calendar.month_abbr[m]} {y % 100:02d}"


def fy_label(y: int) -> str:
    return f"FY{FIRST_FY + y - 1}"


# Fictional entity: a building and maintenance business, numbers in $000 ---------------------
GROUPS = {
    "rev": ("Revenue lines", ["Build contracts", "Renovations", "Maintenance", "Design fees", "Other revenue"]),
    "cogs": ("Cost of sales lines", ["Build materials", "Subcontractors", "Renovation materials", "Maintenance parts", "Other direct costs"]),
    "sal": ("Salary and wage lines", ["Site crews", "Project managers", "Office staff", "Senior leaders", "Leave, KiwiSaver and ACC"]),
    "opx": ("Operating expense lines", ["Vehicles and plant", "Rent and occupancy", "IT and software", "Insurance", "Marketing",
                                         "Professional fees", "General and other"]),
}
SCENARIOS = [  # name, revenue factor, unit cost factor, overheads factor (forecast months only)
    ("Base", 1.00, 1.00, 1.00),
    ("Upside", 1.08, 0.98, 1.02),
    ("Downside", 0.88, 1.04, 0.98),
]
ASSUMPTIONS = {  # name: (label, value, unit)
    "Asm_Tax_Rate": ("Company tax rate", TAX_RATE, "%"),
    "Asm_Debtor_Days": ("Debtor days (on revenue)", 38, "days"),
    "Asm_Stock_Days": ("Inventory days (on cost of sales)", 22, "days"),
    "Asm_Creditor_Days": ("Creditor days (on other operating expenses)", 30, "days"),
    "Asm_Supplier_Days": ("Inventory payable days (on cost of sales)", 42, "days"),
    "Asm_Open_Cash": ("Opening cash", 850, "$000"),
    "Asm_Open_NCA": ("Opening non-current assets", 3300, "$000"),
    "Asm_Open_Loan": ("Opening borrowings", 2400, "$000"),
    "Asm_Open_Share_Capital": ("Opening share capital", 1500, "$000"),
    "Asm_Reserves": ("Reserves", 180, "$000"),
}
SEASON = [1.00, 0.97, 0.88, 0.86, 0.92, 1.00, 1.06, 1.10, 0.96, 0.82, 1.08, 1.15]   # April to March


def demo_inputs(seed: int = 7) -> dict:
    """Typed monthly inputs for the fictional entity. Actual months carry noise; forecast months are smooth."""
    rng = random.Random(seed)

    def series(base, growth=0.06, season=True, noise=0.07, step=None):
        out = []
        for t in range(1, PERIODS + 1):
            v = base * (1 + growth) ** ((t - 1) / 12)
            if season:
                v *= SEASON[(t - 1) % 12]
            if step:
                v *= step(t)
            if t <= LAST_ACTUAL:
                v *= 1 + rng.uniform(-noise, noise)
            out.append(round(v, 1))
        return out

    rev = [series(600), series(240, 0.04), series(118, 0.03, season=False), series(132, 0.09), series(24, 0.0, noise=0.2)]
    build, reno, maint = rev[0], rev[1], rev[2]
    cogs = [[round(0.36 * b * (1 + rng.uniform(-0.03, 0.03) if t < LAST_ACTUAL else 1), 1) for t, b in enumerate(build)],
            [round(0.24 * (b + r), 1) for b, r in zip(build, reno)],
            [round(0.30 * r, 1) for r in reno],
            [round(0.22 * m, 1) for m in maint],
            series(12, 0.0, season=False, noise=0.25)]
    sal = [series(125, 0.04, season=False, noise=0.03), series(72, 0.05, season=False, noise=0.03),
           series(41, 0.03, season=False, noise=0.03), series(36, 0.04, season=False, noise=0.02)]
    sal.append([round(0.08 * sum(s[t] for s in sal), 1) for t in range(PERIODS)])
    audit = lambda t: 2.6 if (t - 1) % 12 == 2 else 1.0                  # noqa: E731  audit fees in June
    opx = [series(22, 0.03), series(18, 0.03, season=False, noise=0.01), series(12, 0.04, season=False, noise=0.05),
           series(9, 0.05, season=False, noise=0.02), series(9, 0.32, noise=0.15), series(6, 0.03, season=False, noise=0.1, step=audit),
           series(6, 0.02, season=False, noise=0.15)]
    capex = [20.0] * PERIODS
    capex[16] += 300          # August 2025: plant
    capex[43] += 450          # November 2027: yard
    draws = [0.0] * PERIODS
    draws[39] = 600           # July 2027
    repay = [20.0] * PERIODS
    issues = [0.0] * PERIODS
    div = [200.0 if (t - 1) % 12 in (5, 11) else 0.0 for t in range(1, PERIODS + 1)]   # September and March
    oca = [round(150 + 12 * ((t % 12) / 12), 1) for t in range(1, PERIODS + 1)]
    ocl = [round(215 + 20 * (((t + 4) % 12) / 12), 1) for t in range(1, PERIODS + 1)]
    loan, dep, interest = Asm("Asm_Open_Loan"), [], []
    nca = Asm("Asm_Open_NCA")
    for t in range(PERIODS):
        interest.append(round(loan * 0.072 / 12, 1))
        loan += draws[t] - repay[t]
        dep.append(round(nca * 0.0065, 1))
        nca += capex[t] - dep[t]
    smooth = lambda base, growth, season=True: [round(base * (1 + growth) ** (t / 12) * (SEASON[t % 12] if season else 1), 1)  # noqa: E731
                                                for t in range(PERIODS)]
    rev_b = [a + b + c + d + e for a, b, c, d, e in zip(smooth(600, .06), smooth(240, .04), smooth(118, .03, False), smooth(132, .09), smooth(24, 0, False))]
    budget = {
        "rev": [round(v * 1.04, 1) for v in rev_b],
        "cogs": [round(v * 0.475, 1) for v in rev_b],
        "sal": [round(v, 1) for v in (a + b + c + d for a, b, c, d in zip(smooth(125, .04, False), smooth(72, .05, False), smooth(41, .03, False), smooth(36, .04, False)))],
        "opx": [round(v, 1) for v in (a + b for a, b in zip(smooth(70, .05, False), smooth(12, .1)))],
        "dep": dep[:],
        "int": interest[:],
    }
    budget["sal"] = [round(v * 1.08, 1) for v in budget["sal"]]
    inp = {"rev": rev, "cogs": cogs, "sal": sal, "opx": opx, "dep": dep, "int": interest, "capex": capex, "draws": draws,
           "repay": repay, "issues": issues, "div": div, "oca": oca, "ocl": ocl, "budget": budget}
    # Opening working capital as if the first month had run for a while, and retained earnings
    # as the balancing item, so the opening balance sheet balances.
    r1, c1, o1 = sum(s[0] for s in rev), sum(s[0] for s in cogs), sum(s[0] for s in opx)
    inp["open"] = {
        "deb": round(r1 * Asm("Asm_Debtor_Days") / 30 * 0.97, 1),
        "inv": round(c1 * Asm("Asm_Stock_Days") / 30 * 0.97, 1),
        "oca": oca[0],
        "cred": round(o1 * Asm("Asm_Creditor_Days") / 30 * 0.97, 1),
        "invpay": round(c1 * Asm("Asm_Supplier_Days") / 30 * 0.97, 1),
        "ocl": ocl[0],
    }
    o = inp["open"]
    assets = Asm("Asm_Open_Cash") + o["deb"] + o["inv"] + o["oca"] + Asm("Asm_Open_NCA")
    liabilities = o["cred"] + o["invpay"] + o["ocl"] + Asm("Asm_Open_Loan")
    o["re"] = round(assets - liabilities - Asm("Asm_Open_Share_Capital") - Asm("Asm_Reserves"), 1)
    return inp


def Asm(name: str) -> float:
    return ASSUMPTIONS[name][1]


# Statement lines: key, label, kind (flow, balance or ratio) -----------------------------------
LINES = [
    ("rev", "Revenue", "flow"), ("cogs", "Cost of sales", "flow"), ("gm", "Gross margin", "flow"),
    ("gm_pct", "Gross margin %", "ratio"), ("sal", "Salaries and wages", "flow"), ("opx", "Other operating expenses", "flow"),
    ("netopex", "Operating expenses", "flow"), ("ebitda", "EBITDA", "flow"), ("dep", "Depreciation", "flow"),
    ("ebit", "EBIT", "flow"), ("int", "Interest", "flow"), ("npbt", "Net profit before tax", "flow"), ("tax", "Tax", "flow"),
    ("npat", "Net profit after tax", "flow"),
    ("cash", "Cash", "balance"), ("deb", "Debtors", "balance"), ("inv", "Inventory", "balance"),
    ("oca", "Other current assets", "balance"), ("ca", "Current assets", "balance"), ("nca", "Non-current assets", "balance"),
    ("ta", "Total assets", "balance"), ("cred", "Creditors", "balance"), ("invpay", "Inventory payables", "balance"),
    ("ocl", "Other current liabilities", "balance"), ("cl", "Current liabilities", "balance"),
    ("loan", "Borrowings (non-current)", "balance"), ("tl", "Total liabilities", "balance"), ("na", "Net assets", "balance"),
    ("sc", "Share capital", "balance"), ("res", "Reserves", "balance"), ("re", "Retained earnings", "balance"),
    ("te", "Total equity", "balance"), ("bal_chk", "Net assets less total equity (must be nil)", "check"),
    ("open_cash", "Opening cash", "flow"), ("receipts", "Receipts from customers", "flow"),
    ("payments", "Payments to suppliers and staff", "flow"), ("othop", "Interest and tax paid", "flow"),
    ("opcf", "Operating cash flow", "flow"), ("invcf", "Investing cash flows", "flow"), ("fincf", "Financing cash flows", "flow"),
    ("chg", "Change in cash", "flow"), ("invfin", "Investing and financing", "flow"),
    ("wc_deb", "Decrease in debtors", "flow"), ("wc_inv", "Decrease in inventory", "flow"),
    ("wc_cred", "Increase in creditors", "flow"), ("wc_invpay", "Increase in inventory payables", "flow"),
    ("wc_net", "Working capital movement", "flow"),
]
LABEL = {k: lab for k, lab, _ in LINES}
KIND = {k: kind for k, _, kind in LINES}
BUDGET_LINES = [("rev", "Revenue"), ("cogs", "Cost of sales"), ("gm", "Gross margin"), ("sal", "Salaries and wages"),
                ("opx", "Other operating expenses"), ("netopex", "Operating expenses"), ("ebitda", "EBITDA"), ("dep", "Depreciation"),
                ("ebit", "EBIT"), ("int", "Interest"), ("npbt", "Net profit before tax"), ("tax", "Tax"), ("npat", "Net profit after tax")]


def reference(inp: dict, scenario: int, last_actual: int = LAST_ACTUAL) -> dict[str, list[float]]:
    """The statements for one scenario (1 Base, 2 Upside, 3 Downside), as the Statements sheet calculates them."""
    _, f_rev, f_unit, f_ovh = SCENARIOS[scenario - 1]
    out = {k: [] for k in LABEL}
    for g, (_, names) in GROUPS.items():
        for k in range(len(names)):
            out[f"{g}_{k + 1}"] = []
    prev = dict(inp["open"], cash=Asm("Asm_Open_Cash"), nca=Asm("Asm_Open_NCA"), loan=Asm("Asm_Open_Loan"),
                sc=Asm("Asm_Open_Share_Capital"))
    for t in range(PERIODS):
        fc = t + 1 > last_actual
        v = {}
        for g, factor in (("rev", f_rev), ("cogs", f_rev * f_unit), ("sal", f_ovh), ("opx", f_ovh)):
            vals = [inp[g][k][t] * (factor if fc else 1) for k in range(len(GROUPS[g][1]))]
            for k, x in enumerate(vals):
                v[f"{g}_{k + 1}"] = x
            v[g] = sum(vals)
        v["gm"] = v["rev"] - v["cogs"]
        v["gm_pct"] = v["gm"] / v["rev"] if v["rev"] else 0
        v["netopex"] = v["sal"] + v["opx"]
        v["ebitda"] = v["gm"] - v["netopex"]
        v["dep"] = inp["dep"][t]
        v["ebit"] = v["ebitda"] - v["dep"]
        v["int"] = inp["int"][t]
        v["npbt"] = v["ebit"] - v["int"]
        v["tax"] = v["npbt"] * TAX_RATE
        v["npat"] = v["npbt"] - v["tax"]
        v["deb"] = v["rev"] * Asm("Asm_Debtor_Days") / 30
        v["inv"] = v["cogs"] * Asm("Asm_Stock_Days") / 30
        v["oca"] = inp["oca"][t]
        v["nca"] = prev["nca"] + inp["capex"][t] - v["dep"]
        v["cred"] = v["opx"] * Asm("Asm_Creditor_Days") / 30
        v["invpay"] = v["cogs"] * Asm("Asm_Supplier_Days") / 30
        v["ocl"] = inp["ocl"][t]
        v["loan"] = prev["loan"] + inp["draws"][t] - inp["repay"][t]
        v["sc"] = prev["sc"] + inp["issues"][t]
        v["res"] = Asm("Asm_Reserves")
        v["re"] = prev["re"] + v["npat"] - inp["div"][t]
        v["open_cash"] = prev["cash"]
        v["receipts"] = v["rev"] - (v["deb"] - prev["deb"])
        v["payments"] = -(v["cogs"] + v["sal"] + v["opx"] + (v["inv"] - prev["inv"]) + (v["oca"] - prev["oca"])
                          - (v["cred"] - prev["cred"]) - (v["invpay"] - prev["invpay"]) - (v["ocl"] - prev["ocl"]))
        v["othop"] = -(v["int"] + v["tax"])
        v["opcf"] = v["receipts"] + v["payments"] + v["othop"]
        v["invcf"] = -inp["capex"][t]
        v["fincf"] = inp["draws"][t] - inp["repay"][t] + inp["issues"][t] - inp["div"][t]
        v["chg"] = v["opcf"] + v["invcf"] + v["fincf"]
        v["invfin"] = v["invcf"] + v["fincf"]
        v["cash"] = prev["cash"] + v["chg"]
        v["wc_deb"] = -(v["deb"] - prev["deb"])
        v["wc_inv"] = -(v["inv"] - prev["inv"])
        v["wc_cred"] = v["cred"] - prev["cred"]
        v["wc_invpay"] = v["invpay"] - prev["invpay"]
        v["wc_net"] = v["wc_deb"] + v["wc_inv"] + v["wc_cred"] + v["wc_invpay"]
        v["ca"] = v["cash"] + v["deb"] + v["inv"] + v["oca"]
        v["ta"] = v["ca"] + v["nca"]
        v["cl"] = v["cred"] + v["invpay"] + v["ocl"]
        v["tl"] = v["cl"] + v["loan"]
        v["na"] = v["ta"] - v["tl"]
        v["te"] = v["sc"] + v["res"] + v["re"]
        v["bal_chk"] = v["na"] - v["te"]
        for k in out:
            out[k].append(v[k])
        prev = v
    return out


def reference_budget(inp: dict) -> dict[str, list[float]]:
    b = {k: list(v) for k, v in inp["budget"].items()}
    b["gm"] = [r - c for r, c in zip(b["rev"], b["cogs"])]
    b["netopex"] = [s + o for s, o in zip(b["sal"], b["opx"])]
    b["ebitda"] = [g - n for g, n in zip(b["gm"], b["netopex"])]
    b["ebit"] = [e - d for e, d in zip(b["ebitda"], b["dep"])]
    b["npbt"] = [e - i for e, i in zip(b["ebit"], b["int"])]
    b["tax"] = [n * TAX_RATE for n in b["npbt"]]
    b["npat"] = [n - x for n, x in zip(b["npbt"], b["tax"])]
    return b


# Workbook helpers -----------------------------------------------------------------------------
def q(ws) -> str:
    return "'" + ws.title.replace("'", "''") + "'"


def absref(ws, c1: int, r1: int, c2: int | None = None, r2: int | None = None) -> str:
    a = f"{q(ws)}!${L(c1)}${r1}"
    if c2 is None:
        return a
    return a + f":${L(c2)}${r2 if r2 is not None else r1}"


def put(ws, r: int, c: int, value, font=None, fmt=None, fill=None, border=None, align=None):
    cell = ws.cell(r, c)
    cell.value = value
    cell.font = font or F_BODY
    if fmt:
        cell.number_format = fmt
    if fill:
        cell.fill = fill
    if border:
        cell.border = border
    if align:
        cell.alignment = align
    return cell


def define(wb, nm: str, ref: str) -> None:
    wb.defined_names[nm] = DefinedName(nm, attr_text=ref)


def nm_of(prefix: str, key: str) -> str:
    return prefix + "_" + "_".join(p.capitalize() for p in key.split("_"))


def frame_sheet(wb, title: str, purpose: str, timeline: bool = False):
    ws = wb.create_sheet(title)
    ws.sheet_view.showGridLines = False
    put(ws, 1, 2, title, F_HEAD)
    put(ws, 2, 2, purpose)
    put(ws, 3, 2, "Demo Building Co (fictional numbers, $000)", F_NOTE)
    put(ws, 1, 1, '=HYPERLINK("#Contents!A1","<")', F_LINK)
    ws.column_dimensions["A"].width = 2.5
    for c in "BCDEF":
        ws.column_dimensions[c].width = 2.5
    ws.column_dimensions["G"].width = 34
    ws.column_dimensions["H"].width = 9
    ws.column_dimensions["I"].width = 10
    if timeline:
        for c in range(FIRST_COL, LAST_COL + 1):
            ws.column_dimensions[L(c)].width = 8
        put(ws, 5, 3, "Month")
        put(ws, 6, 3, "Month ending")
        put(ws, 7, 3, "Forecast (1) or actual (0)")
        put(ws, 5, 9, "Total", F_BOLD)
        for t in range(1, PERIODS + 1):
            c = L(FIRST_COL + t - 1)
            put(ws, 5, FIRST_COL + t - 1, f"=Time!{c}12", F_BOLD, "0")
            put(ws, 6, FIRST_COL + t - 1, f"=Time!{c}14", F_BOLD)
            put(ws, 7, FIRST_COL + t - 1, f"=Time!{c}18", F_BODY, "0")
        ws.freeze_panes = "J8"
    return ws


def section(ws, r: int, text: str) -> None:
    put(ws, r, 2, text, F_HEAD, border=UNDER)
    for c in range(3, 10):
        ws.cell(r, c).border = UNDER


# Model sheets ---------------------------------------------------------------------------------
OPENING = {"cash": "Asm_Open_Cash", "nca": "Asm_Open_NCA", "loan": "Asm_Open_Loan", "sc": "Asm_Open_Share_Capital",
           "deb": "Asm_Open_Debtors", "inv": "Asm_Open_Inventory", "oca": "Asm_Open_OCA", "cred": "Asm_Open_Creditors",
           "invpay": "Asm_Open_Inv_Payables", "ocl": "Asm_Open_OCL", "re": "Asm_Open_Retained"}
OPENING_LABELS = {"deb": "Opening debtors", "inv": "Opening inventory", "oca": "Opening other current assets",
                  "cred": "Opening creditors", "invpay": "Opening inventory payables", "ocl": "Opening other current liabilities",
                  "re": "Opening retained earnings (balances the opening position)"}


def time_sheet(wb):
    ws = frame_sheet(wb, "Time", "Timeline: four financial years to March, actuals to the last actual month")
    for c in range(FIRST_COL, LAST_COL + 1):
        ws.column_dimensions[L(c)].width = 8
    rows = [("Ts_Start_Date", "First month", "=DATE(2024,4,1)", "d mmm yyyy", True),
            ("Ts_Periods", "Months in the timeline", PERIODS, "0", False),
            ("Ts_First_FY", "First financial year (year ending March)", FIRST_FY, "0", False),
            ("Ts_Last_Actual", "Last actual month (month number)", LAST_ACTUAL, "0", True),
            ("Ts_Years", "Financial years", "=Ts_Periods/12", "0", False)]
    for i, (nm, label, value, fmt, is_input) in enumerate(rows):
        r = 5 + i
        put(ws, r, 2, label)
        put(ws, r, 8, value, fmt=fmt, fill=FILL_IN if is_input else None)
        define(wb, nm, absref(ws, 8, r))
    dv = DataValidation(type="whole", operator="between", formula1="1", formula2="48", allow_blank=False)
    dv.error, dv.errorTitle = "Pick a month number from 1 to 48", "Last actual month"
    ws.add_data_validation(dv)
    dv.add("H8")
    put(ws, 8, 9, "=INDEX(Ts_Labels,Ts_Last_Actual)")
    timeline = [("Ts_Index", "Month number", lambda c, p: "1" if p is None else f"={p}12+1", "0"),
                ("Ts_Month_End", "Month ending", lambda c, p: f"=EOMONTH(Ts_Start_Date,{c}12-1)", "d mmm yyyy"),
                ("Ts_Labels", "Month", lambda c, p: f'=TEXT({c}13,"mmm yy")', None),
                ("Ts_FY", "Financial year number", lambda c, p: f"=INT(({c}12-1)/12)+1", "0"),
                ("Ts_FY_Label", "Financial year", lambda c, p: f'="FY"&(Ts_First_FY+{c}15-1)', None),
                ("Ts_Actual", "Actual month (1) or not (0)", lambda c, p: f"=IF({c}12<=Ts_Last_Actual,1,0)", "0"),
                ("Ts_Forecast", "Forecast month (1) or not (0)", lambda c, p: f"=1-{c}17", "0"),
                ("Ts_AF", "Actual or forecast", lambda c, p: f'=IF({c}17=1,"A","F")', None)]
    section(ws, 11, "Timeline")
    for i, (nm, label, fn, fmt) in enumerate(timeline):
        r = 12 + i
        put(ws, r, 3, label)
        for t in range(1, PERIODS + 1):
            col = FIRST_COL + t - 1
            f = fn(L(col), None if t == 1 else L(col - 1))
            put(ws, r, col, int(f) if f.isdigit() else f, fmt=fmt)
        define(wb, nm, absref(ws, FIRST_COL, r, LAST_COL, r))
    ws.freeze_panes = "J12"
    return ws


def assumptions_sheet(wb, inp):
    ws = frame_sheet(wb, "Assumptions", "Rates, working capital days and the opening balance sheet")
    ws.column_dimensions["G"].width = 44
    section(ws, 5, "Assumptions")
    r = 6
    items = [(nm, label, value, unit) for nm, (label, value, unit) in ASSUMPTIONS.items()]
    items += [(OPENING[k], OPENING_LABELS[k], inp["open"][k], "$000") for k in ("deb", "inv", "oca", "cred", "invpay", "ocl", "re")]
    for nm, label, value, unit in items:
        put(ws, r, 3, label)
        put(ws, r, 8, value, fmt=PCT if unit == "%" else ("0" if unit == "days" else NUM), fill=FILL_IN)
        put(ws, r, 9, unit)
        define(wb, nm, absref(ws, 8, r))
        r += 1
    put(ws, r + 1, 3, "The opening position balances: retained earnings are what the other opening balances leave.", F_NOTE)
    return ws


def scenarios_sheet(wb):
    ws = frame_sheet(wb, "Scenarios", "Scenario factors, applied to forecast months only; the active scenario feeds the reports")
    ws.column_dimensions["G"].width = 20
    for c in "HIJ":
        ws.column_dimensions[c].width = 12
    put(ws, 5, 3, "Active scenario", F_BOLD)
    put(ws, 5, 8, SCENARIOS[0][0], fill=FILL_IN)
    define(wb, "DD_Scenario", absref(ws, 8, 5))
    put(ws, 5, 9, "=MATCH(DD_Scenario,LU_Scenarios,0)", fmt="0")
    define(wb, "Scn_Active", absref(ws, 9, 5))
    dv = DataValidation(type="list", formula1="LU_Scenarios", allow_blank=False)
    dv.error, dv.errorTitle = "Pick a scenario from the list", "Active scenario"
    ws.add_data_validation(dv)
    dv.add("H5")
    section(ws, 7, "Scenario factors")
    for c, h in zip("GHIJ", ["Scenario", "Revenue", "Unit cost", "Overheads"]):
        put(ws, 8, "GHIJ".index(c) + 7, h, F_BOLD)
    for s, (name, f_rev, f_unit, f_ovh) in enumerate(SCENARIOS, start=1):
        r = 8 + s
        put(ws, r, 7, name, fill=FILL_IN)
        for c, v, nm in ((8, f_rev, "Scn_Rev"), (9, f_unit, "Scn_Unit"), (10, f_ovh, "Scn_Ovh")):
            put(ws, r, c, v, fmt="0.00", fill=FILL_IN)
            define(wb, f"{nm}_{s}", absref(ws, c, r))
    define(wb, "LU_Scenarios", absref(ws, 7, 9, 7, 11))
    put(ws, 13, 3, "Revenue scales sales; unit cost scales cost of sales per dollar of sales; overheads scale salaries and other operating expenses.", F_NOTE)
    return ws


def inputs_sheet(wb, inp):
    ws = frame_sheet(wb, "Inputs", "Monthly inputs: actuals to the last actual month, then the base forecast; and the budget", timeline=True)
    rows = {}
    r = 9
    blocks = [(g, GROUPS[g][0], [(f"{g}_{k + 1}", n, inp[g][k]) for k, n in enumerate(GROUPS[g][1])]) for g in GROUPS]
    blocks.append(("other", "Other monthly inputs", [("dep", "Depreciation", inp["dep"]), ("int", "Interest", inp["int"]),
                                                     ("capex", "Capital expenditure", inp["capex"]), ("draws", "Borrowings drawn", inp["draws"]),
                                                     ("repay", "Borrowings repaid", inp["repay"]), ("issues", "Shares issued", inp["issues"]),
                                                     ("div", "Dividends paid", inp["div"]), ("oca", "Other current assets", inp["oca"]),
                                                     ("ocl", "Other current liabilities", inp["ocl"])]))
    blocks.append(("budget", "Budget", [(f"bud_{k}", lab, inp["budget"][k]) for k, lab in
                                       (("rev", "Revenue"), ("cogs", "Cost of sales"), ("sal", "Salaries and wages"),
                                        ("opx", "Other operating expenses"), ("dep", "Depreciation"), ("int", "Interest"))]))
    for _, title, items in blocks:
        section(ws, r, title)
        r += 1
        for key, label, values in items:
            put(ws, r, 3, label)
            put(ws, r, 8, "$000")
            put(ws, r, 9, f"=SUM(J{r}:{L(LAST_COL)}{r})", F_BOLD, NUM)
            for t, v in enumerate(values):
                put(ws, r, FIRST_COL + t, v, fmt=NUM, fill=FILL_IN)
            rows[key] = r
            r += 1
        r += 1
    for g, (_, names) in GROUPS.items():
        define(wb, f"LU_{g.capitalize()}_Lines", absref(ws, 3, rows[f"{g}_1"], 3, rows[f"{g}_{len(names)}"]))
    return ws, rows


BLOCK = (["@Income"] + [f"rev_{k + 1}" for k in range(5)] + ["rev"] + [f"cogs_{k + 1}" for k in range(5)] + ["cogs", "gm", "gm_pct"]
         + [f"sal_{k + 1}" for k in range(5)] + ["sal"] + [f"opx_{k + 1}" for k in range(7)]
         + ["opx", "netopex", "ebitda", "dep", "ebit", "int", "npbt", "tax", "npat",
            "@Balance sheet", "cash", "deb", "inv", "oca", "ca", "nca", "ta", "cred", "invpay", "ocl", "cl", "loan", "tl", "na",
            "sc", "res", "re", "te", "bal_chk",
            "@Cash flow", "open_cash", "receipts", "payments", "othop", "opcf", "invcf", "fincf", "chg", "invfin",
            "wc_deb", "wc_inv", "wc_cred", "wc_invpay", "wc_net"])


def block_label(key: str) -> str:
    g, _, k = key.partition("_")
    if g in GROUPS and k.isdigit():
        return GROUPS[g][1][int(k) - 1]
    return LABEL[key]


def block_kind(key: str) -> str:
    g, _, k = key.partition("_")
    if g in GROUPS and k.isdigit():
        return "flow"
    return "opening" if key == "open_cash" else KIND[key]


def statements_sheet(wb, inrows):
    ws = frame_sheet(wb, "Statements", "Income statement, balance sheet and cash flow for the active scenario, the budget and each scenario",
                     timeline=True)
    r = 9
    rows: dict[str, dict[str, int]] = {}

    def plan(tag):
        nonlocal r
        rows[tag] = {}
        for key in BLOCK:
            if key.startswith("@"):
                r += 1
                rows[tag][key] = r
                r += 1
                continue
            rows[tag][key] = r
            r += 1
        r += 1

    starts = {}
    for tag in ("active", "budget", "s1", "s2", "s3"):
        starts[tag] = r
        r += 1
        if tag == "budget":
            rows["budget"] = {}
            for key, _ in BUDGET_LINES:
                rows["budget"][key] = r
                r += 1
            r += 1
        else:
            plan(tag)

    titles = {"active": "Active scenario (the reports read this block)", "budget": "Budget",
              "s1": f"Scenario 1: {SCENARIOS[0][0]}", "s2": f"Scenario 2: {SCENARIOS[1][0]}", "s3": f"Scenario 3: {SCENARIOS[2][0]}"}
    for tag, start in starts.items():
        section(ws, start, titles[tag])

    def total(key, rr, tag_rows):
        kind = block_kind(key)
        if kind == "flow":
            return f"=SUM(J{rr}:{L(LAST_COL)}{rr})"
        if kind == "balance":
            return f"={L(LAST_COL)}{rr}"
        if kind == "ratio":
            return f"=IF(I{tag_rows['rev']}=0,0,I{tag_rows['gm']}/I{tag_rows['rev']})"
        if kind == "check":
            return f"=SUMPRODUCT(ABS(J{rr}:{L(LAST_COL)}{rr}))"
        return f"=J{rr}"

    def scenario_formula(s, key, c, pc, R):
        cur = lambda k: f"{c}{R[k]}"                                          # noqa: E731
        prev = lambda k: f"{pc}{R[k]}" if pc else OPENING[k]                  # noqa: E731
        g, _, k = key.partition("_")
        fc = f"{c}$7=1"
        if g in GROUPS and k.isdigit():
            factor = {"rev": f"Scn_Rev_{s}", "cogs": f"Scn_Rev_{s}*Scn_Unit_{s}", "sal": f"Scn_Ovh_{s}", "opx": f"Scn_Ovh_{s}"}[g]
            return f"=Inputs!{c}{inrows[key]}*IF({fc},{factor},1)"
        if key in GROUPS:
            n = len(GROUPS[key][1])
            return f"=SUM({cur(key + '_1')}:{cur(f'{key}_{n}')})"
        simple = {
            "gm": f"={cur('rev')}-{cur('cogs')}",
            "gm_pct": f"=IF({cur('rev')}=0,0,{cur('gm')}/{cur('rev')})",
            "netopex": f"={cur('sal')}+{cur('opx')}",
            "ebitda": f"={cur('gm')}-{cur('netopex')}",
            "dep": f"=Inputs!{c}{inrows['dep']}",
            "ebit": f"={cur('ebitda')}-{cur('dep')}",
            "int": f"=Inputs!{c}{inrows['int']}",
            "npbt": f"={cur('ebit')}-{cur('int')}",
            "tax": f"={cur('npbt')}*Asm_Tax_Rate",
            "npat": f"={cur('npbt')}-{cur('tax')}",
            "cash": f"={prev('cash')}+{cur('chg')}",
            "deb": f"={cur('rev')}*Asm_Debtor_Days/30",
            "inv": f"={cur('cogs')}*Asm_Stock_Days/30",
            "oca": f"=Inputs!{c}{inrows['oca']}",
            "ca": f"={cur('cash')}+{cur('deb')}+{cur('inv')}+{cur('oca')}",
            "nca": f"={prev('nca')}+Inputs!{c}{inrows['capex']}-{cur('dep')}",
            "ta": f"={cur('ca')}+{cur('nca')}",
            "cred": f"={cur('opx')}*Asm_Creditor_Days/30",
            "invpay": f"={cur('cogs')}*Asm_Supplier_Days/30",
            "ocl": f"=Inputs!{c}{inrows['ocl']}",
            "cl": f"={cur('cred')}+{cur('invpay')}+{cur('ocl')}",
            "loan": f"={prev('loan')}+Inputs!{c}{inrows['draws']}-Inputs!{c}{inrows['repay']}",
            "tl": f"={cur('cl')}+{cur('loan')}",
            "na": f"={cur('ta')}-{cur('tl')}",
            "sc": f"={prev('sc')}+Inputs!{c}{inrows['issues']}",
            "res": "=Asm_Reserves",
            "re": f"={prev('re')}+{cur('npat')}-Inputs!{c}{inrows['div']}",
            "te": f"={cur('sc')}+{cur('res')}+{cur('re')}",
            "bal_chk": f"=ROUND({cur('na')}-{cur('te')},6)",
            "open_cash": f"={prev('cash')}",
            "receipts": f"={cur('rev')}-({cur('deb')}-{prev('deb')})",
            "payments": (f"=-({cur('cogs')}+{cur('sal')}+{cur('opx')}+({cur('inv')}-{prev('inv')})+({cur('oca')}-{prev('oca')})"
                         f"-({cur('cred')}-{prev('cred')})-({cur('invpay')}-{prev('invpay')})-({cur('ocl')}-{prev('ocl')}))"),
            "othop": f"=-({cur('int')}+{cur('tax')})",
            "opcf": f"={cur('receipts')}+{cur('payments')}+{cur('othop')}",
            "invcf": f"=-Inputs!{c}{inrows['capex']}",
            "fincf": f"=Inputs!{c}{inrows['draws']}-Inputs!{c}{inrows['repay']}+Inputs!{c}{inrows['issues']}-Inputs!{c}{inrows['div']}",
            "chg": f"={cur('opcf')}+{cur('invcf')}+{cur('fincf')}",
            "invfin": f"={cur('invcf')}+{cur('fincf')}",
            "wc_deb": f"=-({cur('deb')}-{prev('deb')})",
            "wc_inv": f"=-({cur('inv')}-{prev('inv')})",
            "wc_cred": f"={cur('cred')}-{prev('cred')}",
            "wc_invpay": f"={cur('invpay')}-{prev('invpay')}",
            "wc_net": f"={cur('wc_deb')}+{cur('wc_inv')}+{cur('wc_cred')}+{cur('wc_invpay')}",
        }
        return simple[key]

    def write_block(tag):
        R = rows[tag]
        for key in BLOCK:
            rr = R[key]
            if key.startswith("@"):
                put(ws, rr, 3, key[1:], F_BOLD)
                continue
            kind = block_kind(key)
            indent = 4 if key.partition("_")[0] in GROUPS and key.partition("_")[2].isdigit() else 3
            bold = key in ("rev", "cogs", "sal", "opx", "gm", "ebitda", "npat", "ca", "ta", "cl", "tl", "na", "te", "opcf", "chg", "wc_net")
            put(ws, rr, indent, block_label(key), F_BOLD if bold else F_BODY)
            put(ws, rr, 8, "%" if kind == "ratio" else ("check" if kind == "check" else "$000"))
            fmt = PCT if kind == "ratio" else NUM
            put(ws, rr, 9, total(key, rr, R), F_BOLD, fmt, border=TOP if bold else None)
            for t in range(1, PERIODS + 1):
                col = FIRST_COL + t - 1
                c, pc = L(col), (L(col - 1) if t > 1 else None)
                if tag == "active":
                    f = f"=CHOOSE(Scn_Active,{c}{rows['s1'][key]},{c}{rows['s2'][key]},{c}{rows['s3'][key]})"
                else:
                    f = scenario_formula(int(tag[1]), key, c, pc, R)
                put(ws, rr, col, f, F_CHECK if kind == "check" else (F_BOLD if bold else F_BODY), fmt,
                    border=TOP if bold else None)
            prefix = "St" if tag == "active" else f"Sc{tag[1]}"
            define(wb, nm_of(prefix, key), absref(ws, FIRST_COL, rr, LAST_COL, rr))
        prefix = "St" if tag == "active" else f"Sc{tag[1]}"
        for g, (_, names) in GROUPS.items():
            define(wb, f"{prefix}_{g.capitalize()}_Lines", absref(ws, FIRST_COL, R[f"{g}_1"], LAST_COL, R[f"{g}_{len(names)}"]))

    for tag in ("active", "s1", "s2", "s3"):
        write_block(tag)
    B = rows["budget"]
    for key, label in BUDGET_LINES:
        rr = B[key]
        bold = key in ("rev", "gm", "ebitda", "npat")
        put(ws, rr, 3, label, F_BOLD if bold else F_BODY)
        put(ws, rr, 8, "$000")
        put(ws, rr, 9, f"=SUM(J{rr}:{L(LAST_COL)}{rr})", F_BOLD, NUM)
        for t in range(1, PERIODS + 1):
            c = L(FIRST_COL + t - 1)
            cur = lambda k: f"{c}{B[k]}"                                       # noqa: E731
            f = {"rev": f"=Inputs!{c}{inrows['bud_rev']}", "cogs": f"=Inputs!{c}{inrows['bud_cogs']}",
                 "gm": f"={cur('rev')}-{cur('cogs')}", "sal": f"=Inputs!{c}{inrows['bud_sal']}",
                 "opx": f"=Inputs!{c}{inrows['bud_opx']}", "netopex": f"={cur('sal')}+{cur('opx')}",
                 "ebitda": f"={cur('gm')}-{cur('netopex')}", "dep": f"=Inputs!{c}{inrows['bud_dep']}",
                 "ebit": f"={cur('ebitda')}-{cur('dep')}", "int": f"=Inputs!{c}{inrows['bud_int']}",
                 "npbt": f"={cur('ebit')}-{cur('int')}", "tax": f"={cur('npbt')}*Asm_Tax_Rate",
                 "npat": f"={cur('npbt')}-{cur('tax')}"}[key]
            put(ws, rr, FIRST_COL + t - 1, f, F_BOLD if bold else F_BODY, NUM)
        define(wb, nm_of("Bud", key), absref(ws, FIRST_COL, rr, LAST_COL, rr))
    return ws, rows


def lookups_sheet(wb):
    ws = frame_sheet(wb, "Lookups", "Lists behind the drop-downs and the chart labels")
    ws.column_dimensions["G"].width = 12
    put(ws, 5, 3, "Financial years", F_BOLD)
    for y in range(1, YEARS + 1):
        put(ws, 5 + y, 3, f"=INDEX(Ts_FY_Label,{(y - 1) * 12 + 1})")
    define(wb, "LU_Years", absref(ws, 3, 6, 3, 5 + YEARS))
    put(ws, 5, 8, "Months", F_BOLD)
    for t in range(1, PERIODS + 1):
        put(ws, 5 + t, 8, f"=INDEX(Ts_Labels,{t})")
    define(wb, "LU_Months", absref(ws, 8, 6, 8, 5 + PERIODS))
    put(ws, 5, 10, "Month names", F_BOLD)
    ws.column_dimensions["J"].width = 12
    for p, mname in enumerate(MONTH_NAMES, start=1):
        put(ws, 5 + p, 10, mname)
    define(wb, "LU_Month_Names", absref(ws, 10, 6, 10, 17))
    put(ws, 20, 3, "Category lists (LU_Rev_Lines and the others) are the label cells on Inputs, so they grow with the categories.", F_NOTE)
    return ws


# Chart helpers --------------------------------------------------------------------------------
def txt(size=8, bold=False, colour=TEXT) -> RichText:
    props = CharacterProperties(sz=int(size * 100), b=bold, solidFill=colour, latin=DrawingFont(typeface="Segoe UI"))
    return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=props), r=[], endParaRPr=CharacterProperties())])


def line_props(colour, width=1.5, dash=None):
    return LineProperties(solidFill=colour, w=int(width * 12700), prstDash=dash)


def style_line(ser, colour, width=1.5, dash=None, marker=False):
    ser.graphicalProperties = GraphicalProperties(ln=line_props(colour, width, dash))
    ser.marker = Marker(symbol="circle" if marker else "none")
    if marker:
        ser.marker.size = 4
        ser.marker.graphicalProperties = GraphicalProperties(solidFill=colour, ln=LineProperties(solidFill=colour))
    ser.smooth = False


def style_fill(ser, colour=None, hatch=False, outline=None):
    gp = GraphicalProperties()
    if hatch:
        gp.pattFill = PatternFillProperties(prst="upDiag", fgClr=ColorChoice(srgbClr=colour or INK), bgClr=ColorChoice(srgbClr="FFFFFF"))
        gp.line = LineProperties(solidFill=colour or INK, w=6350)
    elif colour is None:
        gp.noFill = True
        gp.line = LineProperties(noFill=True)
    else:
        gp.solidFill = colour
        gp.line = LineProperties(solidFill=outline, w=6350) if outline else LineProperties(noFill=True)
    ser.graphicalProperties = gp


def labels(fmt=NUM, pos=None, val=True, pct=False, cat=False, size=7):
    dl = DataLabelList()
    dl.numFmt = fmt
    dl.txPr = txt(size)
    dl.showVal, dl.showPercent, dl.showCatName = val, pct, cat
    dl.showSerName = dl.showLegendKey = dl.showBubbleSize = False
    if pos:
        dl.position = pos
    return dl


def series(ws, r: int, c1: int, c2: int, label_ref: str, cats: str) -> Series:
    ser = Series(Reference(ws, min_col=c1, max_col=c2, min_row=r))
    ser.tx = SeriesLabel(strRef=StrRef(label_ref))
    ser.cat = AxDataSource(strRef=StrRef(cats))
    return ser


def finish_axes(ch, y_fmt=NUM, value_axis=True, reverse=False):
    ch.x_axis.delete = False
    ch.x_axis.txPr = txt(7.5)
    ch.x_axis.majorTickMark = "none"
    ch.x_axis.tickLblPos = "low"
    ch.x_axis.spPr = GraphicalProperties(ln=LineProperties(solidFill="808080", w=6350))
    if reverse:
        ch.x_axis.scaling.orientation = "maxMin"
    ch.y_axis.delete = not value_axis
    ch.y_axis.numFmt = y_fmt
    ch.y_axis.txPr = txt(7.5)
    ch.y_axis.majorTickMark = "none"
    ch.y_axis.spPr = GraphicalProperties(ln=LineProperties(noFill=True))
    ch.y_axis.majorGridlines = ChartLines(spPr=GraphicalProperties(ln=LineProperties(solidFill="E7E6E6", w=6350))) if value_axis else None


def finish_chart(ch, title_ref: str, title_text: str, legend="b"):
    ref = StrRef(title_ref)
    ref.strCache = StrData(ptCount=1, pt=[StrVal(idx=0, v=title_text)])
    ch.title = Title(tx=Text(strRef=ref), overlay=False, txPr=txt(9.5, True))
    if legend:
        ch.legend.position = legend
        ch.legend.txPr = txt(7.5)
        ch.legend.overlay = False
    else:
        ch.legend = None
    ch.graphical_properties = GraphicalProperties(ln=LineProperties(solidFill="D9D9D9", w=6350))
    ch.width, ch.height = CHART_W, CHART_H
    ch.style = None


# Report modules -------------------------------------------------------------------------------
def col_px(width: float) -> int:
    """Excel's pixel width for a column width in characters (Calibri 11, the workbook default)."""
    return int(((256 * width + int(128 / 7)) / 256) * 7)


REPORT_WIDTHS = {"A": 2.5, "B": 2.5, "C": 2.5, "D": 2.5, "E": 2.5, "F": 2.5, "G": 34, "H": 9, "I": 10}
VALUE_WIDTH = 9
CHART_PX = round(CHART_W / 2.54 * 96)
GAP_PX = 14
GRID_TOP = 9


def grid_last_col() -> int:
    """The last column the chart grid reaches, for the print area."""
    x, c = 0, 0
    right = col_px(REPORT_WIDTHS["A"]) + GRID_COLS * CHART_PX + (GRID_COLS - 1) * GAP_PX
    while x < right:
        c += 1
        x += col_px(REPORT_WIDTHS.get(L(c), VALUE_WIDTH))
    return c


def chart_anchor(k: int) -> OneCellAnchor:
    widths = [REPORT_WIDTHS.get(L(c), VALUE_WIDTH) for c in range(1, 60)]
    starts = [0]
    for w in widths:
        starts.append(starts[-1] + col_px(w))
    x = starts[1] + (k % GRID_COLS) * (CHART_PX + GAP_PX)
    col = max(i for i in range(len(widths)) if starts[i] <= x)
    row = GRID_TOP - 1 + (k // GRID_COLS) * GRID_ROWS
    return OneCellAnchor(_from=AnchorMarker(col=col, colOff=pixels_to_EMU(x - starts[col]), row=row, rowOff=pixels_to_EMU(4)),
                         ext=XDRPositiveSize2D(cm_to_EMU(CHART_W), cm_to_EMU(CHART_H)))


FIXED = {
    "current_assets": {"stacks": [("cash", 1), ("deb", 1), ("inv", 1), ("oca", 1)], "total": "ca"},
    "equity": {"stacks": [("sc", 1), ("res", 1), ("re", 1)], "total": "te"},
    "net_assets": {"stacks": [("ta", 1), ("tl", -1)], "total": "na", "current": "na"},
    "bs_parts": {"items": [("cash", "Cash"), ("deb", "Debtors"), ("inv", "Inventory"), ("oca", "Other current assets"),
                           ("nca", "Non-current assets"), ("cl", "Current liabilities"), ("loan", "Non-current liabilities")]},
    "cash_parts": {"items": [("receipts", "Receipts"), ("payments", "Payments"), ("othop", "Interest and tax"),
                             ("invcf", "Investing"), ("fincf", "Financing")]},
    "bs_lines": {"items": [("cash", "Cash"), ("deb", "Debtors"), ("inv", "Inventory"), ("oca", "Other current assets"),
                           ("nca", "Non-current assets"), ("cred", "Creditors"), ("invpay", "Inventory payables"),
                           ("ocl", "Other current liabilities"), ("loan", "Borrowings"), ("sc", "Share capital"),
                           ("res", "Reserves"), ("re", "Retained earnings")]},
    "bs_totals": {"items": [("ta", "Total assets"), ("tl", "Total liabilities"), ("te", "Total equity")]},
}


def uses(chart: dict) -> set[str]:
    """Which selections a chart reads: the year shown, the month shown, or both."""
    r, frame = chart["recipe"], chart.get("frame")
    if r in ("budget",) or (r == "scenario" and chart.get("by") == "month"):
        return {"year"}
    if r == "scenario":
        return set()
    if r == "movement" or frame in ("rolling", "at", "ytd"):
        return {"month"}
    return {"year"}


@dataclass
class Report:
    wb: object
    ws: object
    code: str
    title: str
    year: bool
    month: bool
    rolling: bool
    row: int = 0
    r_pos: int = 0
    r_mname: int = 0
    r_yidx: int = 0
    r_ridx: int = 0
    r_rlab: int = 0
    checks: list = field(default_factory=list)
    tables: dict = field(default_factory=dict)
    expected: dict = field(default_factory=dict)

    # selections
    def n(self, what: str) -> str:
        return f"{self.code}_{what}"

    def bounds(self, frame: str) -> tuple[str, str]:
        return {"year": (self.n("Y_Start"), self.n("Y_End")), "rolling": (f"{self.n('Month')}-11", self.n("Month")),
                "ytd": (self.n("YTD_Start"), self.n("Month"))}[frame]

    def default_label(self, frame: str) -> str:
        """The window label as the default selections show it (chart title caches for apps that ignore linked titles)."""
        return {"year": fy_label(YEAR_SHOWN), "rolling": f"12 months to {month_label(MONTH_SHOWN)}", "at": f"at {month_label(MONTH_SHOWN)}",
                "ytd": f"year to {month_label(MONTH_SHOWN)}", "movement": f"{month_label(MONTH_SHOWN)} against a year earlier",
                "each": "each year"}[frame]

    def window_label(self, frame: str) -> str:
        return {"year": f"DD_{self.code}_Year", "rolling": f'"12 months to "&DD_{self.code}_Month',
                "at": f'"at "&DD_{self.code}_Month', "ytd": f'"year to "&DD_{self.code}_Month'}[frame]

    def idx(self, frame: str, col: int, shift: int = 0) -> str:
        base = f"{L(col)}${self.r_yidx if frame == 'year' else self.r_ridx}"
        return base if shift == 0 else f"({base}{shift:+d})"

    def cats_row(self, frame: str) -> int:
        return self.r_mname if frame == "year" else self.r_rlab

    def period_label(self, frame: str, which: str) -> str:
        if which == "shown":
            return "=" + self.window_label(frame)
        if frame == "year":
            return f'="FY"&(Ts_First_FY+{self.n("Year")}{"-2" if which == "prior" else ""})'
        return '="Prior 12 months"' if which == "prior" else '="Next 12 months"'


def pick(nm: str, i: str) -> str:
    return f"IF(OR({i}<1,{i}>Ts_Periods),NA(),INDEX({nm},{i}))"


def pick2(nm: str, k: str, i: str) -> str:
    return f"IF(OR({i}<1,{i}>Ts_Periods),NA(),INDEX({nm},{k},{i}))"


def span(nm: str, a: str, b: str) -> str:
    return f"SUM(INDEX({nm},{a}):INDEX({nm},{b}))"


def span2(nm: str, k: str, a: str, b: str) -> str:
    return f"SUM(INDEX({nm},{k},{a}):INDEX({nm},{k},{b}))"


def cum(nm: str, i: str, p: int) -> str:
    start = i if p == 1 else f"({i}-{p - 1})"
    return f"IF(OR({start}<1,{i}>Ts_Periods),NA(),SUM(INDEX({nm},{start}):INDEX({nm},{i})))"


class Table:
    """One chart's rows: a heading, the title the chart shows, categories across from column J, series below."""

    def __init__(self, rep: Report, chart: dict, label_expr: str, default: str = ""):
        self.rep, self.ws, self.chart = rep, rep.ws, chart
        self.title_text = chart["title"] + (f", {default}" if default else "")
        r = rep.row
        put(self.ws, r, 2, f"{chart['id']}  {chart['title']}", F_HEAD, border=UNDER)
        for c in range(3, 22):
            self.ws.cell(r, c).border = UNDER
        put(self.ws, r + 1, 3, "Chart title", F_NOTE)
        self.title_row = r + 1
        self.title_formula = f'="{chart["title"]}"' + (f"&\", \"&{label_expr}" if label_expr else "")
        put(self.ws, r + 1, 7, self.title_formula, F_BOLD)
        self.head = r + 2
        self.next = r + 3
        self.rows: list[int] = []
        rep.tables[chart["id"]] = r

    def header(self, cells: list, label="", total=""):
        put(self.ws, self.head, 7, label, F_BOLD)
        put(self.ws, self.head, 9, total, F_BOLD, align=Alignment(horizontal="right"))
        for j, v in enumerate(cells):
            put(self.ws, self.head, FIRST_COL + j, v, F_BOLD, align=Alignment(horizontal="right"))
        self.ncat = len(cells)

    def add(self, label, values: list, unit="$000", total=None, fmt=NUM, bold=False) -> int:
        r = self.next
        put(self.ws, r, 7, label, F_BOLD if bold else F_BODY)
        put(self.ws, r, 8, unit)
        if callable(total):
            total = total(r)
        if total is not None:
            put(self.ws, r, 9, total, F_BOLD, fmt)
        for j, v in enumerate(values):
            put(self.ws, r, FIRST_COL + j, v, fmt=fmt)
        self.rows.append(r)
        self.next += 1
        return r

    def note(self, text):
        put(self.ws, self.next, 7, text, F_NOTE)
        self.next += 1

    def done(self):
        self.rep.row = max(self.next, getattr(self, "after", 0)) + 1

    # references for the chart
    def cats(self) -> str:
        return absref(self.ws, FIRST_COL, self.head, FIRST_COL + self.ncat - 1, self.head)

    def ser(self, r: int, n: int | None = None) -> Series:
        n = n or self.ncat
        return series(self.ws, r, FIRST_COL, FIRST_COL + n - 1, absref(self.ws, 7, r), self.cats())

    def title(self) -> str:
        return absref(self.ws, 7, self.title_row)


def total_if_full(r: int, n: int = 12) -> str:
    a, b = L(FIRST_COL), L(FIRST_COL + n - 1)
    return f'=IF(COUNT({a}{r}:{b}{r})={n},SUM({a}{r}:{b}{r}),"")'


def ranking(t: Table, g: str, frame: str, start: int) -> list[int]:
    """Rows from `start` that rank a category group by its total for the period; returns the row of each rank."""
    rep, ws = t.rep, t.ws
    names = GROUPS[g][1]
    a, b = rep.bounds(frame)
    put(ws, start, 7, "Ranking for the period", F_NOTE)
    for c, h in ((9, "Total"), (10, "Rank key"), (11, "Line at rank")):
        put(ws, start, c, h, F_NOTE, align=Alignment(horizontal="right"))
    first = start + 1
    last = first + len(names) - 1
    out = []
    for k in range(1, len(names) + 1):
        r = first + k - 1
        put(ws, r, 7, f"=INDEX(LU_{g.capitalize()}_Lines,{k})", F_NOTE)
        put(ws, r, 9, "=" + span2(f"St_{g.capitalize()}_Lines", str(k), a, b), F_NOTE, NUM)
        put(ws, r, 10, f"=I{r}-{k}/1000000000", F_NOTE, "#,##0.000")
        put(ws, r, 11, f"=MATCH(LARGE($J${first}:$J${last},{k}),$J${first}:$J${last},0)", F_NOTE, "0")
        out.append(r)
    t.rank_first, t.rank_last = first, last
    t.after = last + 1
    return out


def members(t: Table, chart: dict, frame: str, rank_rows: list[int] | None):
    """(label formula, line-number expression) for each member shown, ranked or in model order."""
    g = chart["group"]
    n = len(GROUPS[g][1])
    top = chart.get("top")
    if top:
        return [(f"=INDEX(LU_{g.capitalize()}_Lines,$K${rank_rows[k]})", f"$K${rank_rows[k]}") for k in range(min(top, n))]
    return [(f"=INDEX(LU_{g.capitalize()}_Lines,{k})", str(k)) for k in range(1, n + 1)]


def has_other(chart: dict) -> bool:
    return bool(chart.get("top")) and chart["top"] < len(GROUPS[chart["group"]][1])


def other_label(g: str) -> str:
    return {"rev": "Other revenue lines", "cogs": "Other cost of sales", "sal": "Other salaries and wages",
            "opx": "Remaining operating expenses"}[g]


# Recipes ----------------------------------------------------------------------------------------
def month_cats(rep: Report, frame: str) -> list[str]:
    return [f"={L(FIRST_COL + j)}${rep.cats_row(frame)}" for j in range(12)]


def flow_total(line: str):
    return (lambda r: total_if_full(r)) if KIND.get(line, "flow") == "flow" else None


def r_compare(rep: Report, ch: dict):
    frame, line = ch.get("frame", "year"), ch["line"]
    nm, cumulative = nm_of("St", line), ch.get("cumulative", False)
    pct = KIND[line] == "ratio"
    t = Table(rep, ch, rep.window_label(frame), rep.default_label(frame))
    t.header(month_cats(rep, frame), "Series", "" if cumulative else "Total")
    styles = {"prior": (GREY, 1.25, None, False), "shown": (INK, 2.25, None, True), "next": (PALETTE[0], 1.5, "dash", False)}
    fills = {"prior": GREY, "shown": INK, "next": PALETTE[0]}
    rows = {}
    for which in ch.get("periods", ["prior", "shown", "next"]):
        shift = {"prior": -12, "shown": 0, "next": 12}[which]
        vals = []
        for p in range(1, 13):
            i = rep.idx(frame, FIRST_COL + p - 1, shift)
            vals.append("=" + (cum(nm, i, p) if cumulative else pick(nm, i)))
        rows[which] = t.add(rep.period_label(frame, which), vals, "%" if pct else "$000",
                            None if cumulative else flow_total(line), PCT if pct else NUM, bold=which == "shown")
    t.done()
    if ch.get("kind") == "column":
        c = BarChart()
        c.type, c.grouping, c.gapWidth, c.overlap = "col", "clustered", 60, -10
        for which, r in rows.items():
            s = t.ser(r)
            style_fill(s, fills[which])
            c.series.append(s)
    else:
        c = LineChart()
        for which, r in rows.items():
            s = t.ser(r)
            style_line(s, *styles[which])
            c.series.append(s)
    finish_axes(c, "0%" if pct else "#,##0")
    finish_chart(c, t.title(), t.title_text)
    return c


def r_mix(rep: Report, ch: dict):
    frame, g = ch.get("frame", "year"), ch["group"]
    compare = ch.get("compare", True)
    t = Table(rep, ch, rep.window_label(frame), rep.default_label(frame))
    t.header(month_cats(rep, frame), "Series", "Total")
    stack_rows, line_rows = [], []
    colours = iter(PALETTE * 2)
    if g in GROUPS:
        G = g.capitalize()
        n_show = min(ch.get("top") or 99, len(GROUPS[g][1]))
        n_rows = (1 if ch.get("lead") else 0) + n_show + (1 if has_other(ch) else 0) + (2 if compare else 0)
        rank_rows = ranking(t, g, frame, t.next + n_rows + 1) if ch.get("top") else None
        if ch.get("lead"):
            lead = ch["lead"]
            vals = ["=" + pick(nm_of("St", lead), rep.idx(frame, FIRST_COL + j)) for j in range(12)]
            stack_rows.append((t.add(LABEL[lead], vals, total=flow_total(lead)), PALETTE[1]))
            colours = iter([PALETTE[0]] + PALETTE[2:] + PALETTE)
        first = t.next
        for label, k in members(t, ch, frame, rank_rows):
            vals = ["=" + pick2(f"St_{G}_Lines", k, rep.idx(frame, FIRST_COL + j)) for j in range(12)]
            stack_rows.append((t.add(label, vals, total=flow_total(g)), next(colours)))
        last = t.next - 1
        if has_other(ch):
            vals = ["=" + pick(nm_of("St", g), rep.idx(frame, FIRST_COL + j)) + f"-SUM({L(FIRST_COL + j)}{first}:{L(FIRST_COL + j)}{last})"
                    for j in range(12)]
            stack_rows.append((t.add(other_label(g), vals, total=flow_total(g)), LIGHT))
        total_key = ch.get("total", g)
    else:
        spec = FIXED[g]
        for key, sign in spec["stacks"]:
            vals = [("=-" if sign < 0 else "=") + pick(nm_of("St", key), rep.idx(frame, FIRST_COL + j)) for j in range(12)]
            stack_rows.append((t.add(LABEL[key] + (" (shown below zero)" if sign < 0 else ""), vals), next(colours)))
        total_key = spec["total"]
        if spec.get("current"):
            vals = ["=" + pick(nm_of("St", spec["current"]), rep.idx(frame, FIRST_COL + j)) for j in range(12)]
            line_rows.append((t.add(LABEL[spec["current"]], vals, bold=True), (INK, 2.25, None, True)))
    if compare:
        for which, style in (("prior", (GREY, 1.5, None, False)), ("next", (INK, 1.5, "dash", False))):
            shift = -12 if which == "prior" else 12
            vals = ["=" + pick(nm_of("St", total_key), rep.idx(frame, FIRST_COL + j, shift)) for j in range(12)]
            line_rows.append((t.add(rep.period_label(frame, which), vals, total=flow_total(total_key)), style))
    t.done()
    if ch.get("top") and frame == "year":
        a, b = rep.bounds(frame)
        rows = [r for r, _ in stack_rows]
        rep.checks.append((f"{ch['id']} {ch['title']}: the lines shown do not add to the total",
                           f"=IF(ABS(SUM({q(t.ws)}!I{rows[0]}:I{rows[-1]})-{span(nm_of('St', total_key), a, b)})>0.001,1,0)", "error"))
    c = BarChart()
    c.type, c.grouping, c.overlap, c.gapWidth = "col", "stacked", 100, 55
    for r, colour in stack_rows:
        s = t.ser(r)
        style_fill(s, colour)
        c.series.append(s)
    if line_rows:
        ln = LineChart()
        for r, style in line_rows:
            s = t.ser(r)
            style_line(s, *style)
            ln.series.append(s)
        c += ln
    finish_axes(c)
    finish_chart(c, t.title(), t.title_text)
    return c


def r_depth(rep: Report, ch: dict):
    frame, g = ch.get("frame", "year"), ch["group"]
    G = g.capitalize()
    a, b = rep.bounds(frame)
    t = Table(rep, ch, rep.window_label(frame), rep.default_label(frame))
    t.header(["=" + rep.window_label(frame)], "Line")
    n_show = min(ch.get("top") or 99, len(GROUPS[g][1]))
    rank_rows = ranking(t, g, frame, t.next + n_show + (1 if has_other(ch) else 0) + 1) if ch.get("top") else None
    rows, first = [], t.next
    for label, k in members(t, ch, frame, rank_rows):
        v = f"=INDEX($I${t.rank_first}:$I${t.rank_last},{k})" if ch.get("top") else "=" + span2(f"St_{G}_Lines", k, a, b)
        rows.append(t.add(label, [v]))
    if has_other(ch):
        rows.append(t.add(other_label(g), ["=" + span(nm_of("St", g), a, b) + f"-SUM(J{first}:J{t.next - 1})"]))
        rep.checks.append((f"{ch['id']} {ch['title']}: Other is below zero, so the ranking is wrong",
                           f"=IF({q(t.ws)}!J{rows[-1]}<-0.001,1,0)", "error"))
    t.done()
    c = BarChart()
    c.type, c.grouping, c.gapWidth, c.overlap = "col", "clustered", 40, -5
    for j, r in enumerate(rows):
        s = t.ser(r, 1)
        style_fill(s, LIGHT if (has_other(ch) and j == len(rows) - 1) else PALETTE[j % len(PALETTE)])
        s.dLbls = labels("#,##0", "outEnd")
        c.series.append(s)
    finish_axes(c)
    finish_chart(c, t.title(), t.title_text)
    return c


def r_pie(rep: Report, ch: dict):
    frame, g = ch["frame"], ch["group"]
    t = Table(rep, ch, rep.window_label(frame), rep.default_label(frame))
    if g in GROUPS:
        G = g.capitalize()
        a, b = rep.bounds(frame)
        n_show = min(ch.get("top") or 99, len(GROUPS[g][1]))
        rank_rows = ranking(t, g, frame, t.next + 2)
        mem = members(t, ch, frame, rank_rows)
        cats = [lab for lab, _ in mem] + ([f'="{other_label(g)}"'] if has_other(ch) else [])
        t.header(cats, "Line")
        vals = [f"=INDEX($I${t.rank_first}:$I${t.rank_last},{k})" for _, k in mem]
        if has_other(ch):
            vals.append("=" + span(nm_of("St", g), a, b) + f"-SUM(J{t.next}:{L(FIRST_COL + n_show - 1)}{t.next})")
        vr = t.add("Total for the period", vals, total=lambda r: f"=SUM(J{r}:{L(FIRST_COL + len(vals) - 1)}{r})")
        colours = PALETTE[:n_show] + ([LIGHT] if has_other(ch) else [])
    elif g == "bs_parts":
        items = FIXED[g]["items"]
        t.header([lab for _, lab in items], "Line")
        vr = t.add("Balance at the month shown", ["=" + pick(nm_of("St", k), rep.n("Month")) for k, _ in items])
        colours = PALETTE[:5] + [GREY, "595959"]
    else:  # cash_parts
        items = FIXED[g]["items"]
        a, b = rep.bounds(frame)
        sr = t.next
        t.header([f'="{lab}"&IF({L(FIRST_COL + j)}{sr}<0," (out)"," (in)")' for j, (_, lab) in enumerate(items)], "Flow")
        t.add("Net flow, year to the month shown", ["=" + span(nm_of("St", k), a, b) for k, _ in items])
        vr = t.add("Size (the pie shows sizes)", [f"=ABS({L(FIRST_COL + j)}{sr})" for j in range(len(items))])
        colours = PALETTE[:5]
    t.done()
    c = PieChart()
    c.varyColors = True
    s = t.ser(vr)
    s.dPt = [DataPoint(idx=j, spPr=GraphicalProperties(solidFill=col, ln=LineProperties(solidFill="FFFFFF", w=9525)))
             for j, col in enumerate(colours)]
    s.dLbls = labels("0%", "bestFit", val=False, pct=True)
    c.series.append(s)
    finish_chart(c, t.title(), t.title_text, legend="r")
    return c


def r_combo(rep: Report, ch: dict):
    frame = ch.get("frame", "year")
    t = Table(rep, ch, rep.window_label(frame), rep.default_label(frame))
    t.header(month_cats(rep, frame), "Series", "Total")
    bar_rows = [t.add(LABEL[k], ["=" + pick(nm_of("St", k), rep.idx(frame, FIRST_COL + j)) for j in range(12)], total=flow_total(k))
                for k in ch["bars"]]
    line_rows = [t.add(LABEL[k], ["=" + pick(nm_of("St", k), rep.idx(frame, FIRST_COL + j)) for j in range(12)], total=flow_total(k), bold=True)
                 for k in ch["lines"]]
    t.done()
    c = BarChart()
    c.type, c.grouping = "col", ch["grouping"]
    c.overlap, c.gapWidth = (100, 55) if ch["grouping"] == "stacked" else (-10, 60)
    for j, r in enumerate(bar_rows):
        s = t.ser(r)
        style_fill(s, PALETTE[j % len(PALETTE)])
        c.series.append(s)
    if line_rows:
        ln = LineChart()
        for j, r in enumerate(line_rows):
            s = t.ser(r)
            style_line(s, INK if j == 0 else GREY, 2.25, None, True)
            ln.series.append(s)
        c += ln
    finish_axes(c)
    finish_chart(c, t.title(), t.title_text)
    return c


def r_budget(rep: Report, ch: dict):
    line = ch["line"]
    st, bud = nm_of("St", line), nm_of("Bud", line)
    t = Table(rep, ch, rep.window_label("year"), rep.default_label("year"))
    t.header(month_cats(rep, "year"), "Series", "Total")
    tot = lambda r: f"=SUM(J{r}:U{r})"                                                     # noqa: E731
    ix = [rep.idx("year", FIRST_COL + j) for j in range(12)]
    ra = t.add("Actual", [f"=IF(INDEX(Ts_Actual,{i})=1,INDEX({st},{i}),0)" for i in ix], total=tot)
    rf = t.add("Forecast", [f"=IF(INDEX(Ts_Actual,{i})=1,0,INDEX({st},{i}))" for i in ix], total=tot)
    rb = t.add("Budget", [f"=INDEX({bud},{i})" for i in ix], total=tot, bold=True)
    t.add("Actual and forecast less budget", [f"={L(FIRST_COL + j)}{ra}+{L(FIRST_COL + j)}{rf}-{L(FIRST_COL + j)}{rb}" for j in range(12)],
          total=tot, fmt='+#,##0;-#,##0;"-"')
    t.done()
    c = BarChart()
    c.type, c.grouping, c.overlap, c.gapWidth = "col", "stacked", 100, 55
    sa, sf = t.ser(ra), t.ser(rf)
    style_fill(sa, INK)
    style_fill(sf, INK, hatch=True)
    c.series += [sa, sf]
    ln = LineChart()
    sb = t.ser(rb)
    style_line(sb, PALETTE[0], 2.25, None, True)
    ln.series.append(sb)
    c += ln
    finish_axes(c)
    finish_chart(c, t.title(), t.title_text)
    return c


def r_scenario(rep: Report, ch: dict):
    line, by = ch["line"], ch["by"]
    t = Table(rep, ch, rep.window_label("year") if by == "month" else '"each year"', rep.default_label("year" if by == "month" else "each"))
    rows = []
    if by == "month":
        t.header(month_cats(rep, "year"), "Scenario", "Total")
        for s in range(1, 4):
            vals = ["=" + pick(nm_of(f"Sc{s}", line), rep.idx("year", FIRST_COL + j)) for j in range(12)]
            rows.append(t.add(f"=INDEX(LU_Scenarios,{s})", vals, total=flow_total(line)))
    else:
        t.header([f"=INDEX(LU_Years,{y})" for y in range(1, YEARS + 1)], "Scenario")
        for s in range(1, 4):
            nm = nm_of(f"Sc{s}", line)
            vals = [("=" + span(nm, str((y - 1) * 12 + 1), str(y * 12))) if KIND[line] == "flow" else f"=INDEX({nm},{y * 12})"
                    for y in range(1, YEARS + 1)]
            rows.append(t.add(f"=INDEX(LU_Scenarios,{s})", vals))
        t.note("Flows are totals for the year; balances are at the end of the year.")
    t.done()
    c = LineChart()
    for r, style in zip(rows, [(INK, 2.25, None, by == "year"), (PALETTE[0], 1.75, None, by == "year"),
                               (PALETTE[2], 1.75, "dash", by == "year")]):
        s = t.ser(r)
        style_line(s, *style)
        c.series.append(s)
    finish_axes(c)
    finish_chart(c, t.title(), t.title_text)
    return c


def r_bridge(rep: Report, ch: dict):
    g = ch["group"]
    if g == "net_assets_build":
        frame, M = "at", rep.n("Month")
        items = [("Current assets", "=" + pick("St_Ca", M), "step"), ("Non-current assets", "=" + pick("St_Nca", M), "step"),
                 ("Current liabilities", "=-" + pick("St_Cl", M), "step"),
                 ("Non-current liabilities", "=-" + pick("St_Loan", M), "step"), ("Net assets", None, "end")]
        orient, check_ref = "col", pick("St_Na", M)
    else:
        frame = "year"
        a, b = rep.bounds(frame)
        items = [("Opening cash", f"=IF({a}=1,Asm_Open_Cash,INDEX(St_Cash,{a}-1))", "start")]
        items += [(lab, "=" + span(nm_of("St", k), a, b), "step") for k, lab in FIXED["cash_parts"]["items"]]
        items += [("Closing cash", None, "end")]
        orient, check_ref = "bar", f"INDEX(St_Cash,{b})"
    t = Table(rep, ch, rep.window_label(frame), rep.default_label(frame))
    t.header([lab for lab, _, _ in items], "Bridge")
    vr = t.next
    rr, br, tr, ir, dr = vr + 1, vr + 2, vr + 3, vr + 4, vr + 5
    rows = {k: [] for k in ("value", "run", "base", "total", "inc", "dec")}
    for j, (_, value, kind) in enumerate(items):
        c, pc = L(FIRST_COL + j), (L(FIRST_COL + j - 1) if j else None)
        prev_run = f"{pc}{rr}" if pc else "0"
        if kind == "end":
            value = f"={pc}{rr}"
        rows["value"].append(value)
        if kind == "step":
            rows["run"].append(f"={prev_run}+{c}{vr}")
            rows["base"].append(f"=MIN({prev_run},{c}{rr})")
            rows["total"].append("=0")
            rows["inc"].append(f"=MAX({c}{vr},0)")
            rows["dec"].append(f"=MAX(-{c}{vr},0)")
        else:
            rows["run"].append(f"={c}{vr}")
            rows["base"].append("=0")
            rows["total"].append(f"={c}{vr}")
            rows["inc"].append("=0")
            rows["dec"].append("=0")
    t.add("Value", rows["value"])
    t.add("Running total", rows["run"])
    for key, label in (("base", "Base (not drawn)"), ("total", "Total"), ("inc", "Increase"), ("dec", "Decrease")):
        t.add(label, rows[key])
    t.done()
    end = f"{q(t.ws)}!{L(FIRST_COL + len(items) - 1)}{vr}"
    rep.checks.append((f"{ch['id']} {ch['title']}: the bridge does not reach the statement figure",
                       f"=IF(ABS({end}-{check_ref})>0.001,1,0)", "error"))
    rep.checks.append((f"{ch['id']} {ch['title']}: the running total crosses zero, so a step is drawn from the wrong base",
                       f"=IF(MIN({q(t.ws)}!J{rr}:{L(FIRST_COL + len(items) - 1)}{rr})<0,1,0)", "alert"))
    c = BarChart()
    c.type, c.grouping, c.overlap, c.gapWidth = orient, "stacked", 100, 35
    for r, colour, fmt, pos in ((br, None, None, None), (tr, "808080", "#,##0;-#,##0;", "inEnd"),
                                (ir, GOOD, '"+"#,##0;;', "ctr"), (dr, BAD, '"-"#,##0;;', "ctr")):
        s = t.ser(r)
        style_fill(s, colour, outline="404040" if colour else None)
        if fmt:
            s.dLbls = labels(fmt, pos)
        c.series.append(s)
    finish_axes(c, value_axis=False, reverse=orient == "bar")
    finish_chart(c, t.title(), t.title_text, legend=None)
    return c


def r_movement(rep: Report, ch: dict):
    items = FIXED[ch["group"]]["items"]
    M = rep.n("Month")
    t = Table(rep, ch, f'DD_{rep.code}_Month&" against a year earlier"', rep.default_label("movement"))
    t.header([lab for _, lab in items], "Line")
    n = len(items)
    r1 = t.add("Month shown", ["=" + pick(nm_of("St", k), M) for k, _ in items])
    r2 = t.add("A year earlier", ["=" + pick(nm_of("St", k), f"({M}-12)") for k, _ in items])
    r3 = t.add("Change", [f"={L(FIRST_COL + j)}{r1}-{L(FIRST_COL + j)}{r2}" for j in range(n)], fmt='+#,##0;-#,##0;"-"')
    ri = t.add("Increase", [f"=IF({L(FIRST_COL + j)}{r3}>=0,{L(FIRST_COL + j)}{r3},0)" for j in range(n)])
    rd = t.add("Decrease", [f"=IF({L(FIRST_COL + j)}{r3}<0,{L(FIRST_COL + j)}{r3},0)" for j in range(n)])
    t.done()
    c = BarChart()
    c.type, c.grouping, c.overlap, c.gapWidth = "bar", "clustered", 100, 40
    for r, colour in ((ri, PALETTE[0]), (rd, PALETTE[2])):
        s = t.ser(r)
        style_fill(s, colour)
        s.dLbls = labels('+#,##0;-#,##0;;', "outEnd")
        c.series.append(s)
    finish_axes(c, value_axis=False, reverse=True)
    finish_chart(c, t.title(), t.title_text, legend=None)
    return c


RECIPES = {"compare": r_compare, "mix": r_mix, "depth": r_depth, "pie": r_pie, "combo": r_combo, "budget": r_budget,
           "scenario": r_scenario, "bridge": r_bridge, "movement": r_movement}


def report_sheet(wb, module: dict, charts: list[dict]) -> Report:
    code = module["code"]
    ws = wb.create_sheet(module["title"])
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 90
    for c in range(1, 60):
        ws.column_dimensions[L(c)].width = REPORT_WIDTHS.get(L(c), VALUE_WIDTH)
    put(ws, 1, 1, '=HYPERLINK("#Contents!A1","<")', F_LINK)
    put(ws, 1, 2, module["title"], F_HEAD)
    put(ws, 2, 2, f"{len(charts)} charts that come with the module; pick the year or month shown and every chart follows")
    put(ws, 3, 2, "Demo Building Co (fictional numbers, $000)", F_NOTE)
    needs = set().union(*(uses(c) for c in charts))
    rep = Report(wb, ws, code, module["title"], "year" in needs, "month" in needs,
                 any(c.get("frame") == "rolling" for c in charts))
    if rep.year:
        put(ws, 4, 3, "Year shown", F_BOLD)
        put(ws, 4, 8, fy_label(YEAR_SHOWN), fill=FILL_IN, align=Alignment(horizontal="right"))
        put(ws, 4, 9, "=MATCH(H4,LU_Years,0)", fmt="0")
        define(wb, f"DD_{code}_Year", absref(ws, 8, 4))
        define(wb, rep.n("Year"), absref(ws, 9, 4))
        dv = DataValidation(type="list", formula1="LU_Years", allow_blank=False)
        dv.error, dv.errorTitle = "Pick a financial year from the list", "Year shown"
        ws.add_data_validation(dv)
        dv.add("H4")
        rep.checks.append((f"{module['title']}: the year shown is not in the list", f"=IF(ISERROR({rep.n('Year')}),1,0)", "error"))
        if any(c["recipe"] in ("compare", "mix") and c.get("frame", "year") == "year" for c in charts):
            rep.checks.append((f"{module['title']}: the year before or after the year shown is outside the timeline, so those series are blank",
                               f"=IF(ISERROR({rep.n('Year')}),0,IF(OR({rep.n('Year')}=1,{rep.n('Year')}=Ts_Years),1,0))", "alert"))
    if rep.month:
        put(ws, 5, 3, "Month shown", F_BOLD)
        put(ws, 5, 8, month_label(MONTH_SHOWN), fill=FILL_IN, align=Alignment(horizontal="right"))
        put(ws, 5, 9, "=MATCH(H5,LU_Months,0)", fmt="0")
        define(wb, f"DD_{code}_Month", absref(ws, 8, 5))
        define(wb, rep.n("Month"), absref(ws, 9, 5))
        dv = DataValidation(type="list", formula1="LU_Months", allow_blank=False)
        dv.error, dv.errorTitle = "Pick a month from the list", "Month shown"
        ws.add_data_validation(dv)
        dv.add("H5")
        rep.checks.append((f"{module['title']}: the month shown is not in the list", f"=IF(ISERROR({rep.n('Month')}),1,0)", "error"))
        low = 24 if rep.rolling else 13
        rep.checks.append((f"{module['title']}: a comparison for the month shown falls outside the timeline, so it is blank",
                           f"=IF(ISERROR({rep.n('Month')}),0,IF(OR({rep.n('Month')}<{low},{rep.n('Month')}>Ts_Periods-12),1,0))", "alert"))
    put(ws, 6, 3, "Scenario shown", F_BOLD)
    put(ws, 6, 8, "=DD_Scenario", align=Alignment(horizontal="right"))
    put(ws, 6, 10, "Change the scenario on the Scenarios sheet.", F_NOTE)
    section(ws, 8, "Charts")
    grid_rows = -(-len(charts) // GRID_COLS)
    r = GRID_TOP + grid_rows * GRID_ROWS + 1
    section(ws, r, "Chart data: every number is a formula on the statements; the charts read these rows")
    r += 2
    put(ws, r, 7, "Position in the period")
    for j in range(12):
        put(ws, r, FIRST_COL + j, j + 1, fmt="0")
    rep.r_pos = r
    put(ws, r + 1, 7, "Month of the financial year")
    for j in range(12):
        put(ws, r + 1, FIRST_COL + j, f"=INDEX(LU_Month_Names,{L(FIRST_COL + j)}${r})")
    rep.r_mname = r + 1
    r += 2
    if rep.year:
        put(ws, r, 7, "Year shown: month number")
        for j in range(12):
            put(ws, r, FIRST_COL + j, f"=({rep.n('Year')}-1)*12+{L(FIRST_COL + j)}${rep.r_pos}", fmt="0")
        rep.r_yidx = r
        put(ws, r + 1, 7, "Year shown starts at month")
        put(ws, r + 1, 8, f"=({rep.n('Year')}-1)*12+1", fmt="0")
        define(wb, rep.n("Y_Start"), absref(ws, 8, r + 1))
        put(ws, r + 2, 7, "Year shown ends at month")
        put(ws, r + 2, 8, f"={rep.n('Y_Start')}+11", fmt="0")
        define(wb, rep.n("Y_End"), absref(ws, 8, r + 2))
        r += 3
    if rep.month:
        put(ws, r, 7, "Financial year of the month shown starts at month")
        put(ws, r, 8, f"=INT(({rep.n('Month')}-1)/12)*12+1", fmt="0")
        define(wb, rep.n("YTD_Start"), absref(ws, 8, r))
        r += 1
        if rep.rolling:
            put(ws, r, 7, "12 months to the month shown: month number")
            for j in range(12):
                put(ws, r, FIRST_COL + j, f"={rep.n('Month')}-12+{L(FIRST_COL + j)}${rep.r_pos}", fmt="0")
            rep.r_ridx = r
            put(ws, r + 1, 7, "12 months to the month shown")
            for j in range(12):
                c = L(FIRST_COL + j)
                put(ws, r + 1, FIRST_COL + j, f'=IF(OR({c}{r}<1,{c}{r}>Ts_Periods),"",INDEX(LU_Months,{c}{r}))')
            rep.r_rlab = r + 1
            r += 2
    rep.row = r + 1
    for k, ch in enumerate(charts):
        chart = RECIPES[ch["recipe"]](rep, ch)
        chart.anchor = chart_anchor(k)
        ws.add_chart(chart)
    for rr in range(1, rep.row + 1):
        ws.row_dimensions[rr].height = 15
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_area = f"A1:{L(grid_last_col())}{GRID_TOP + grid_rows * GRID_ROWS}"
    return rep


def register_sheet(wb, register: dict, reports: dict[str, Report]):
    ws = wb.create_sheet("Chart register")
    ws.sheet_view.showGridLines = False
    put(ws, 1, 1, '=HYPERLINK("#Contents!A1","<")', F_LINK)
    put(ws, 1, 2, "Chart register", F_HEAD)
    put(ws, 2, 2, "Every chart the summary and report modules bring with them; the id links to the chart's rows")
    put(ws, 3, 2, "Demo Building Co (fictional numbers, $000)", F_NOTE)
    heads = ["Id", "Module", "Chart", "Recipe", "Reads", "Period"]
    for j, (h, w) in enumerate(zip(heads, [7, 18, 44, 11, 40, 12])):
        put(ws, 5, 2 + j, h, F_BOLD, border=UNDER)
        ws.column_dimensions[L(2 + j)].width = w
    ws.column_dimensions["A"].width = 2.5
    titles = {m["key"]: m["title"] for m in register["modules"]}
    for i, ch in enumerate(register["charts"]):
        r = 6 + i
        sheet = titles[ch["module"]]
        row = reports[ch["module"]].tables[ch["id"]]
        put(ws, r, 2, f'=HYPERLINK("#\'{sheet}\'!B{row}","{ch["id"]}")', F_LINK)
        put(ws, r, 3, sheet)
        put(ws, r, 4, ch["title"])
        put(ws, r, 5, ch["recipe"])
        reads = ch.get("line") or ch.get("group") or ", ".join(ch.get("bars", []) + ch.get("lines", []))
        if ch.get("top"):
            reads += f", top {ch['top']}"
        if ch.get("lead"):
            reads = f"{ch['lead']} and {reads}"
        put(ws, r, 6, reads)
        put(ws, r, 7, ch.get("frame") or ("month shown" if ch["recipe"] == "movement" else
                                          ("year shown" if ch.get("by") == "month" or ch["recipe"] == "budget" else "each year")))
    ws.freeze_panes = "B6"
    return ws


def checks_sheet(wb, items: list[tuple[str, str, str]]):
    ws = frame_sheet(wb, "Checks", "Error checks must be nil; alerts are for review")
    ws.column_dimensions["G"].width = 90
    r = 5
    totals = {}
    for severity, title in (("error", "Error checks"), ("alert", "Alerts")):
        section(ws, r, title)
        r += 1
        first = r
        for label, formula, sev in items:
            if sev != severity:
                continue
            put(ws, r, 3, label)
            put(ws, r, 8, formula, F_CHECK if severity == "error" else F_BODY, "0")
            r += 1
        put(ws, r, 3, f"{title} raised", F_BOLD)
        put(ws, r, 8, f"=SUM(H{first}:H{r - 1})", F_BOLD, "0", border=TOP)
        totals[severity] = r
        define(wb, "Chk_Errors" if severity == "error" else "Chk_Alerts", absref(ws, 8, r))
        r += 2
    return ws


def model_checks() -> list[tuple[str, str, str]]:
    out = []
    for prefix, name in (("St", "active scenario"), ("Sc1", SCENARIOS[0][0]), ("Sc2", SCENARIOS[1][0]), ("Sc3", SCENARIOS[2][0])):
        out.append((f"Balance sheet does not balance ({name})", f"=IF(SUMPRODUCT(ABS({prefix}_Bal_Chk))>0.001,1,0)", "error"))
    out.append(("Closing cash is not opening cash plus the change in cash",
                "=IF(ABS(INDEX(St_Cash,Ts_Periods)-Asm_Open_Cash-SUM(St_Chg))>0.001,1,0)", "error"))
    out.append(("Active scenario is not in the list", "=IF(ISERROR(Scn_Active),1,0)", "error"))
    out.append(("Last actual month is outside the timeline", "=IF(OR(Ts_Last_Actual<1,Ts_Last_Actual>Ts_Periods),1,0)", "error"))
    out.append(("Cash goes below zero in the active scenario", "=IF(MIN(St_Cash)<0,1,0)", "alert"))
    return out


def contents_sheet(wb, register: dict, reports: dict[str, Report]):
    ws = wb["Contents"]
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 2.5
    ws.column_dimensions["B"].width = 4
    ws.column_dimensions["C"].width = 22
    ws.column_dimensions["D"].width = 8
    ws.column_dimensions["E"].width = 90
    put(ws, 1, 2, "Contents", F_HEAD)
    put(ws, 2, 2, "Report charts proof: the 95 charts the summary and report modules bring, as native charts on a live model")
    put(ws, 3, 2, "Demo Building Co (fictional numbers, $000)", F_NOTE)
    r = 5
    groups = [("Reports", [(m["title"], f"{len(reports[m['key']].tables)} charts") for m in register["modules"]]),
              ("Model", [("Time", "Timeline and the last actual month"), ("Assumptions", "Rates, days and the opening balance sheet"),
                         ("Scenarios", "Scenario factors and the active scenario"), ("Inputs", "Monthly inputs and the budget"),
                         ("Statements", "Statements for the active scenario, the budget and each scenario")]),
              ("Appendices", [("Chart register", "Every chart, its recipe and what it reads, with links"),
                              ("Lookups", "Lists behind the drop-downs"), ("Checks", "Error checks and alerts")])]
    n = 0
    for title, items in groups:
        put(ws, r, 2, title, F_HEAD, border=UNDER)
        for c in range(3, 6):
            ws.cell(r, c).border = UNDER
        r += 1
        for sheet, desc in items:
            n += 1
            put(ws, r, 2, n, fmt="0")
            put(ws, r, 3, f'=HYPERLINK("#\'{sheet}\'!A1","{sheet}")', F_LINK)
            put(ws, r, 5, desc)
            r += 1
        r += 1
    put(ws, r, 3, "Error checks raised", F_BOLD)
    put(ws, r, 4, "=Chk_Errors", F_CHECK, "0")
    put(ws, r + 1, 3, "Alerts raised", F_BOLD)
    put(ws, r + 1, 4, "=Chk_Alerts", F_BODY, "0")
    put(ws, r + 3, 3, "Pick the year or month shown at the top of each report sheet. Change the scenario on the Scenarios sheet "
                      "and the last actual month on the Time sheet.", F_NOTE)
    return ws


def load_register(path: Path = REGISTER) -> dict:
    return yaml.safe_load(Path(path).read_text())


def build(path: Path, register_path: Path = REGISTER) -> Path:
    register = load_register(register_path)
    inp = demo_inputs()
    wb = Workbook()
    wb.active.title = "Contents"
    reports: dict[str, Report] = {}
    for module in register["modules"]:
        charts = [c for c in register["charts"] if c["module"] == module["key"]]
        reports[module["key"]] = report_sheet(wb, module, charts)
    time_sheet(wb)
    assumptions_sheet(wb, inp)
    scenarios_sheet(wb)
    _, inrows = inputs_sheet(wb, inp)
    statements_sheet(wb, inrows)
    register_sheet(wb, register, reports)
    lookups_sheet(wb)
    items = model_checks()
    for rep in reports.values():
        items += rep.checks
    checks_sheet(wb, items)
    contents_sheet(wb, register, reports)
    for ws in wb.worksheets:
        if ws.title in [m["title"] for m in register["modules"]]:
            continue
        for rr in range(1, ws.max_row + 1):
            ws.row_dimensions[rr].height = 15
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    names = {m["title"]: [f"{c['id']} {c['title']}" for c in register["charts"] if c["module"] == m["key"]]
             for m in register["modules"]}
    finish(path, names)
    return path


NA_AS_BLANK = ('<extLst><ext uri="{56B9EC1D-385E-4148-901F-78D8002777C0}" '
               'xmlns:c16r3="http://schemas.microsoft.com/office/drawing/2017/03/chart">'
               '<c16r3:dataDisplayOptions16><c16r3:dispNaAsBlank val="1"/></c16r3:dataDisplayOptions16></ext></extLst>')


def finish(path: Path, chart_names: dict[str, list[str]]) -> None:
    """What openpyxl does not write: chart names (the register id and title, so the add-in can find a module's
    charts), Excel's "Show #N/A as an empty cell", and label formats that do not follow the cells."""
    with zipfile.ZipFile(path) as z:
        parts = {n: z.read(n) for n in z.namelist()}
    wbx = parts["xl/workbook.xml"].decode()
    rels = parts["xl/_rels/workbook.xml.rels"].decode()
    for sheet, names in chart_names.items():
        rid = re.search(rf'<sheet [^>]*name="{re.escape(sheet)}"[^>]*r:id="(rId\d+)"', wbx).group(1)
        target = re.search(rf'<Relationship [^>]*Id="{rid}"[^>]*Target="([^"]+)"', rels) or \
            re.search(rf'<Relationship [^>]*Target="([^"]+)"[^>]*Id="{rid}"', rels)
        sheet_part = "xl/" + target.group(1).lstrip("/").replace("xl/", "")
        srels = parts[sheet_part.replace("worksheets/", "worksheets/_rels/") + ".rels"].decode()
        drawing = re.search(r'Target="([^"]*drawing\d+\.xml)"', srels).group(1)
        dpart = "xl/drawings/" + drawing.split("/")[-1]
        x = parts[dpart].decode()
        it = iter(names)
        x = re.sub(r'name="Chart \d+"', lambda m: f'name="{next(it)}"', x)
        parts[dpart] = x.encode()
    for n in parts:
        if n.startswith("xl/charts/chart") and n.endswith(".xml"):
            x = parts[n].decode()
            if "dispNaAsBlank" not in x:
                x = re.sub(r"(</(?:c:)?chart>)", NA_AS_BLANK + r"\1", x, count=1)
            x = re.sub(r'<(c:)?numFmt formatCode="([^"]*)"/>', r'<\1numFmt formatCode="\2" sourceLinked="0"/>', x)
            parts[n] = x.encode()
    fd, tmp = tempfile.mkstemp(suffix=".xlsx")
    os.close(fd)
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as z:
        for n, b in parts.items():
            z.writestr(n, b)
    shutil.move(tmp, path)


if __name__ == "__main__":
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else HERE.parent.parent / "build" / "reports" / "reports_demo.xlsx"
    print(build(out))
