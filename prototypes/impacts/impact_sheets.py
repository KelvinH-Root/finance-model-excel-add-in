"""Impacts sheets: what one transaction does to the income statement, balance sheet and cash flow.

The add-in's Impacts command has two modes. The live mode (live.py) changes an input in the
model, recalculates, reads the statements and puts the input back. This module is the other
mode: Add to model writes an Impacts section of formula sheets, one per kind of transaction
the model actually has, labelled with the model's own accounts and entities, so a reviewer or a
new analyst can see the entries and their effect without the add-in.

HFG's own design (clean room): Modano's Financial Statement Impacts Analyser is the idea, not
the source. Each sheet holds:

- inputs (amounts, the GST rate) and switches (GST registered, paid in the period, capitalised);
- the entries, debit positive, each with the entity it sits in and, for cash, its cash flow class;
- the effect on the statements, one column per entity and, for intergroup items, the
  eliminations and the group: surplus carried to retained surplus, net cash flow equal to the
  change in cash, and a balance check.

Each item names the account roles it needs; an item is offered only when the model's context
has all of them. Every formula is reproduced by `panels` (Python), which the tests compare.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

TEXT = "404040"
F_BODY = Font(name="Segoe UI", size=9, color=TEXT)
F_BOLD = Font(name="Segoe UI", size=9, color=TEXT, bold=True)
F_HEAD = Font(name="Segoe UI", size=10, color=TEXT, bold=True)
F_NOTE = Font(name="Segoe UI", size=9, color="808080", italic=True)
F_ARROW = Font(name="Segoe UI", size=9, color="679DB5", bold=True)
FILL_IN = PatternFill("solid", fgColor="FFF2CC")
UNDER = Border(bottom=Side(style="thin", color=TEXT))
TOP = Border(top=Side(style="thin", color=TEXT))
NUM = '#,##0.0;(#,##0.0);"-"'
RIGHT = Alignment(horizontal="right")
CENTRE = Alignment(horizontal="center")
IND = Alignment(horizontal="left", indent=1)

CLASSES = ("asset", "liability", "equity", "revenue", "expense")
CASH_CLASSES = ("Operating", "Investing", "Financing")
# Order the roles appear in on a sheet
ROLE_ORDER = ["cash", "receivables", "gst", "wip", "inv_property", "payables", "netting", "loans", "capital",
              "revenue", "property_sales", "construction_revenue", "cos", "construction_costs", "opex", "interest"]


@dataclass
class Context:
    """What the current model hands the Impacts sheets: its accounts by role and its entities."""
    model: str
    accounts: dict[str, tuple[str, str]]          # role -> (label as the model shows it, class)
    entities: dict[str, str]                      # slot -> entity label
    gst_rate: float = 0.15


@dataclass
class Item:
    key: str
    title: str
    purpose: str
    needs: list[str]
    inputs: list[tuple[str, str, float]]          # (name, label, default)
    switches: list[tuple[str, str, int]] = field(default_factory=list)    # (name, label, default 1 or 0)
    legs: list[tuple[str, str, str, str | None]] = field(default_factory=list)   # (slot, role, expression, cash class)
    slots: list[str] = field(default_factory=lambda: ["entity"])          # entity slots, in column order
    group: bool = False                                                  # intergroup: eliminations and group columns
    notes: list[str] = field(default_factory=list)

    @property
    def sheet(self) -> str:
        return "Imp " + SHORT.get(self.key, self.title)


ELIM = "Eliminations"
SHORT = {"opex": "Operating cost", "facility": "Facility", "construction": "Construction cost", "capint": "Capitalised interest",
         "homes": "Homes sold to the Fund", "claim": "Construction claim", "oncharge": "On-charge at cost"}

ITEMS: list[Item] = [
    Item("sale", "Sale on credit", "A sale invoiced now and collected later (or in the period)",
         ["revenue", "receivables", "cash"], [("amount", "Sale, excluding GST", 100.0)], [("collected", "Collected in the period", 0)],
         [("entity", "receivables", "amount", None), ("entity", "revenue", "-amount", None),
          ("entity", "cash", "amount*collected", "Operating"), ("entity", "receivables", "-amount*collected", None)],
         notes=["Revenue is earned when invoiced; cash follows when the debtor pays."]),
    Item("opex", "Operating cost paid", "An operating cost paid in cash", ["opex", "cash"],
         [("amount", "Cost", 90.0)], [],
         [("entity", "opex", "amount", None), ("entity", "cash", "-amount", "Operating")]),
    Item("facility", "Facility drawn and repaid", "A loan drawn and partly repaid in the period", ["loans", "cash"],
         [("drawn", "Drawn", 500.0), ("repaid", "Repaid", 200.0)], [],
         [("entity", "cash", "drawn", "Financing"), ("entity", "loans", "-drawn", None),
          ("entity", "loans", "repaid", None), ("entity", "cash", "-repaid", "Financing")],
         notes=["Drawdowns and repayments are financing cash flows; the surplus is untouched."]),
    Item("interest", "Interest paid", "Interest on a loan, expensed and paid", ["interest", "cash"],
         [("amount", "Interest", 40.0)], [],
         [("entity", "interest", "amount", None), ("entity", "cash", "-amount", "Operating")]),
    Item("equity", "Equity injection", "Capital put in by the parent or outside investors", ["capital", "cash"],
         [("amount", "Capital", 1000.0)], [],
         [("entity", "cash", "amount", "Financing"), ("entity", "capital", "-amount", None)]),
    Item("land", "Land purchase", "Land bought for a development, held in work in progress", ["wip", "cash"],
         [("price", "Price", 2000.0)], [],
         [("entity", "wip", "price", None), ("entity", "cash", "-price", "Operating")],
         notes=["Land for a development is inventory, so the payment is an operating cash flow.",
                "Land between GST-registered parties is usually zero-rated, so no GST is shown; check each purchase."]),
    Item("construction", "Construction cost into WIP", "A construction cost capitalised, with GST",
         ["wip", "gst", "cash"], [("amount", "Cost, excluding GST", 1000.0), ("gst_rate", "GST rate", 0.15)],
         [("registered", "GST registered", 1), ("refunded", "GST refund received in the period", 0)],
         [("entity", "wip", "amount+amount*gst_rate*(1-registered)", None),
          ("entity", "gst", "amount*gst_rate*registered", None),
          ("entity", "cash", "-amount*(1+gst_rate)", "Operating"),
          ("entity", "cash", "amount*gst_rate*registered*refunded", "Operating"),
          ("entity", "gst", "-amount*gst_rate*registered*refunded", None)],
         notes=["A registered entity claims the GST back (a debit on the GST account until refunded); "
                "an entity that is not registered capitalises it into WIP."]),
    Item("capint", "Interest capitalised", "Interest on a development facility added to WIP", ["wip", "loans", "cash"],
         [("amount", "Interest", 50.0)], [("paid", "Paid in cash", 0)],
         [("entity", "wip", "amount", None), ("entity", "loans", "-amount*(1-paid)", None),
          ("entity", "cash", "-amount*paid", "Operating")],
         notes=["Capitalised interest goes to WIP, not the income statement; added to the loan it is not a cash flow."]),
    Item("homes", "Homes sold to the Fund", "A completed portfolio sold by a development LP to the Fund (WIP released)",
         ["wip", "inv_property", "property_sales", "cos", "receivables", "payables", "cash"],
         [("price", "Price", 4200.0), ("carrying", "Carrying amount in WIP", 3520.0)], [("settled", "Paid in the period", 1)],
         [("seller", "cash", "price*settled", "Operating"), ("seller", "receivables", "price*(1-settled)", None),
          ("seller", "property_sales", "-price", None), ("seller", "cos", "carrying", None), ("seller", "wip", "-carrying", None),
          ("buyer", "inv_property", "price", None), ("buyer", "cash", "-price*settled", "Investing"),
          ("buyer", "payables", "-price*(1-settled)", None),
          (ELIM, "property_sales", "price", None), (ELIM, "cos", "-carrying", None),
          (ELIM, "inv_property", "-(price-carrying)", None),
          (ELIM, "payables", "price*(1-settled)", None), (ELIM, "receivables", "-price*(1-settled)", None),
          (ELIM, "cash", "-price*settled", "Operating"), (ELIM, "cash", "price*settled", "Investing")],
         slots=["seller", "buyer"], group=True,
         notes=["Inside the group the sale and its margin come out; the homes stay at the group's cost.",
                "The cash moved inside the group, so the group cash flow shows nothing."]),
    Item("claim", "Construction claim with margin", "A construction claim from the builder, capitalised by a development LP",
         ["wip", "construction_revenue", "construction_costs", "receivables", "payables", "cash"],
         [("price", "Claim, excluding GST", 1000.0), ("cost", "Builder's cost", 850.0)], [("settled", "Paid in the period", 1)],
         [("seller", "construction_costs", "cost", None), ("seller", "cash", "-cost", "Operating"),
          ("seller", "construction_revenue", "-price", None), ("seller", "cash", "price*settled", "Operating"),
          ("seller", "receivables", "price*(1-settled)", None),
          ("buyer", "wip", "price", None), ("buyer", "cash", "-price*settled", "Operating"),
          ("buyer", "payables", "-price*(1-settled)", None),
          (ELIM, "construction_revenue", "price", None), (ELIM, "construction_costs", "-cost", None),
          (ELIM, "wip", "-(price-cost)", None),
          (ELIM, "payables", "price*(1-settled)", None), (ELIM, "receivables", "-price*(1-settled)", None)],
         slots=["seller", "buyer"], group=True,
         notes=["The builder's margin is in the LP's WIP; the group takes it out, so the group holds WIP at the builder's cost."]),
    Item("oncharge", "On-charge at cost", "A shared cost paid by the parent and on-charged at cost through its netting account",
         ["netting", "receivables", "payables", "wip", "opex", "cash"],
         [("amount", "Cost on-charged", 120.0)], [("capitalised", "Receiver capitalises it into WIP", 1),
                                                  ("settled", "Receiver pays in the period", 0)],
         [("seller", "netting", "amount", None), ("seller", "cash", "-amount", "Operating"),
          ("seller", "receivables", "amount", None), ("seller", "netting", "-amount", None),
          ("seller", "cash", "amount*settled", "Operating"), ("seller", "receivables", "-amount*settled", None),
          ("buyer", "wip", "amount*capitalised", None), ("buyer", "opex", "amount*(1-capitalised)", None),
          ("buyer", "payables", "-amount*(1-settled)", None), ("buyer", "cash", "-amount*settled", "Operating"),
          (ELIM, "payables", "amount*(1-settled)", None), (ELIM, "receivables", "-amount*(1-settled)", None)],
         slots=["seller", "buyer"], group=True,
         notes=["The supplier's bill and the on-charge invoice are both coded to the netting account, so it stays at nil.",
                "A balance left there at the cut-off is an on-charge not yet raised."]),
]
BY_KEY = {i.key: i for i in ITEMS}
SLOT_DEFAULTS = {"homes": {"seller": "devlp", "buyer": "fund"}, "claim": {"seller": "builder", "buyer": "devlp"},
                 "oncharge": {"seller": "parent", "buyer": "devlp"}}


def offered(ctx: Context) -> list[Item]:
    """The items this model's context supports, in library order."""
    return [i for i in ITEMS if all(r in ctx.accounts for r in i.needs) and
            all(SLOT_DEFAULTS.get(i.key, {}).get(s, s) in ctx.entities for s in i.slots)]


def entity_label(ctx: Context, item: Item, slot: str) -> str:
    return ctx.entities[SLOT_DEFAULTS.get(item.key, {}).get(slot, slot)]


# ---------------------------------------------------------------------------------- reference
_IDENT = re.compile(r"[a-z_][a-z_0-9]*")


def _eval(expr: str, values: dict) -> float:
    return float(eval(expr, {"__builtins__": {}}, dict(values)))      # expressions come from ITEMS only


def defaults(item: Item, ctx: Context) -> dict:
    v = {k: (ctx.gst_rate if k == "gst_rate" else d) for k, _, d in item.inputs}
    v.update({k: d for k, _, d in item.switches})
    return v


def lines(item: Item, ctx: Context) -> dict[str, list[str]]:
    """The statement lines the item touches, by section, in role order."""
    used = [r for r in ROLE_ORDER if any(leg[1] == r for leg in item.legs)]
    cls = lambda r: ctx.accounts[r][1]                                   # noqa: E731
    return {"revenue": [r for r in used if cls(r) == "revenue"], "expense": [r for r in used if cls(r) == "expense"],
            "asset": [r for r in used if cls(r) == "asset"], "liability": [r for r in used if cls(r) == "liability"],
            "equity": [r for r in used if cls(r) == "equity"]}


def columns(item: Item, ctx: Context) -> list[str]:
    cols = [entity_label(ctx, item, s) for s in item.slots]
    return cols + [ELIM, "Group"] if item.group else cols


def panels(item: Item, ctx: Context, values: dict | None = None) -> dict[str, dict[str, float]]:
    """The effect on the statements by column, worked out directly from the entries."""
    v = defaults(item, ctx)
    v.update(values or {})
    legs = [((entity_label(ctx, item, s) if s != ELIM else ELIM), r, _eval(e, v), c) for s, r, e, c in item.legs]
    ln = lines(item, ctx)
    out = {}
    for col in [c for c in columns(item, ctx) if c != "Group"]:
        amt = lambda r: sum(a for e, rr, a, _ in legs if e == col and rr == r)           # noqa: E731
        p = {}
        for r in ln["revenue"]:
            p[r] = -amt(r)
        for r in ln["expense"]:
            p[r] = amt(r)
        p["surplus"] = sum(p[r] for r in ln["revenue"]) - sum(p[r] for r in ln["expense"])
        for r in ln["asset"]:
            p[r] = amt(r)
        p["assets"] = sum(p[r] for r in ln["asset"])
        for r in ln["liability"]:
            p[r] = -amt(r)
        p["liabilities"] = sum(p[r] for r in ln["liability"])
        for r in ln["equity"]:
            p[r] = -amt(r)
        p["retained"] = p["surplus"]
        p["equity_total"] = sum(p[r] for r in ln["equity"]) + p["retained"]
        p["balance"] = p["assets"] - p["liabilities"] - p["equity_total"]
        for cc in CASH_CLASSES:
            p[cc] = sum(a for e, rr, a, c in legs if e == col and rr == "cash" and c == cc)
        p["net_cash"] = sum(p[cc] for cc in CASH_CLASSES)
        p["cash_tie"] = p["net_cash"] - p.get("cash", 0.0)
        out[col] = p
    if item.group:
        out["Group"] = {k: sum(out[c][k] for c in out) for k in next(iter(out.values()))}
    return out


FIXED = {"surplus": "Surplus", "assets": "Assets", "liabilities": "Liabilities", "retained": "Retained surplus: this period's surplus",
         "equity_total": "Equity", "balance": "Assets less liabilities less equity (must be nil)", "net_cash": "Net cash flow",
         "cash_tie": "Net cash flow less the change in cash (must be nil)",
         **{cc: f"{cc} cash flows" for cc in CASH_CLASSES}}


def line_label(ctx: Context, key: str) -> str:
    """The label a statement line carries on the sheet: the model's account name, or a fixed total."""
    return ctx.accounts[key][0] if key in ctx.accounts else FIXED[key]


# ---------------------------------------------------------------------------------- the sheets
def _put(ws, r, c, v, font=None, fmt=None, fill=None, border=None, align=None):
    cell = ws.cell(r, c)
    cell.value = v
    cell.font = font or F_BODY
    if fmt:
        cell.number_format = fmt
    if fill:
        cell.fill = fill
    if border:
        cell.border = border
    if align:
        cell.alignment = align
    return cell


def _q(ws) -> str:
    return "'" + ws.title.replace("'", "''") + "'"


def _section(ws, r, text, last_col):
    _put(ws, r, 2, text, F_HEAD, border=UNDER)
    for c in range(3, last_col + 1):
        ws.cell(r, c).border = UNDER


def write_item(wb, item: Item, ctx: Context, headings: dict | None = None) -> tuple[str, str]:
    """Write one Impacts sheet. Returns (sheet name, the cell holding its checks total)."""
    ws = wb.create_sheet(item.sheet)
    ws.sheet_view.showGridLines = False
    cols = columns(item, ctx)
    c0 = 10                                                          # first statement column (J)
    last = c0 + len(cols)
    for col, w in {"A": 2.5, "B": 24, "C": 36, "D": 2, "E": 2, "F": 2, "G": 2, "H": 12, "I": 11}.items():
        ws.column_dimensions[col].width = w
    for j in range(len(cols)):
        ws.column_dimensions[L(c0 + j)].width = 20
    ws.column_dimensions[L(last)].width = 30
    _put(ws, 1, 2, item.title, F_HEAD)
    _put(ws, 2, 2, item.purpose)
    _put(ws, 3, 2, f"{ctx.model}: accounts and entities from this model ($000)", F_NOTE)
    heads = (headings or {}).setdefault(ws.title, [])
    # Inputs and switches
    r = 5
    _section(ws, r, "Inputs", last)
    heads.append((r, "Inputs"))
    v = defaults(item, ctx)
    cell = {}
    for name, label, _ in item.inputs:
        r += 1
        _put(ws, r, 3, label, align=IND)
        _put(ws, r, 8, v[name], fmt="0%" if name == "gst_rate" else NUM, fill=FILL_IN)
        cell[name] = f"$H${r}"
    for name, label, _ in item.switches:
        r += 1
        _put(ws, r, 3, label, align=IND)
        _put(ws, r, 8, "Yes" if v[name] else "No", fill=FILL_IN, align=RIGHT)
        dv = DataValidation(type="list", formula1='"Yes,No"', allow_blank=False)
        dv.error, dv.errorTitle = "Yes or No", label
        ws.add_data_validation(dv)
        dv.add(f"H{r}")
        cell[name] = f'($H${r}="Yes")'

    def excel(expr: str) -> str:
        return _IDENT.sub(lambda m: cell[m.group(0)], expr)

    # The entries
    r += 2
    _section(ws, r, "Entries (debit positive)", last)
    heads.append((r, "Entries"))
    r += 1
    for c, h in ((2, "Entity"), (3, "Account"), (8, "Amount"), (9, "Cash flow")):
        _put(ws, r, c, h, F_BOLD, border=UNDER, align=RIGHT if c >= 8 else IND)
    j0 = r + 1
    for slot, role, expr, cash in item.legs:
        r += 1
        _put(ws, r, 2, entity_label(ctx, item, slot) if slot != ELIM else ELIM, align=IND)
        _put(ws, r, 3, ctx.accounts[role][0], align=IND)
        _put(ws, r, 8, "=" + excel(expr), fmt=NUM)
        _put(ws, r, 9, cash or "", align=RIGHT)
    j1 = r
    r += 1
    _put(ws, r, 3, "Entries add to nil in each entity (must be nil)", F_NOTE, align=IND)
    ent = f"$B${j0}:$B${j1}"
    amt = f"$H${j0}:$H${j1}"
    nil_terms = [f"ABS(SUMPRODUCT(({ent}=\"{c}\")*{amt}))" for c in cols if c != "Group"]
    _put(ws, r, 8, "=ROUND(" + "+".join(nil_terms) + ",6)", F_NOTE, NUM)
    entries_nil = f"$H${r}"
    # The effect
    r += 2
    _section(ws, r, "Effect on the statements", last)
    heads.append((r, "Effect on the statements"))
    r += 1
    for j, c in enumerate(cols):
        _put(ws, r, c0 + j, c, F_BOLD, border=UNDER, align=RIGHT)
    hdr = r
    acct = f"$C${j0}:$C${j1}"
    cls = f"$I${j0}:$I${j1}"
    rows: dict[str, int] = {}
    ln = lines(item, ctx)

    def value_row(key, label, formula_for, bold=False, border=None, note=None, font=None):
        nonlocal r
        r += 1
        rows[key] = r
        _put(ws, r, 3, label, font or (F_BOLD if bold else F_BODY), align=IND)
        for j, c in enumerate(cols):
            col = L(c0 + j)
            if c == "Group":
                f = "=" + "+".join(f"{L(c0 + k)}{r}" for k in range(len(cols) - 1))
            else:
                f = "=" + formula_for(f"{col}${hdr}")
            _put(ws, r, c0 + j, f, font or (F_BOLD if bold else F_BODY), NUM, border=border)
        if note:
            _put(ws, r, last, note, F_ARROW)

    def sumleg(e, label, sign="", extra=""):
        return f"{sign}SUMPRODUCT(({ent}={e})*({acct}=\"{label}\"){extra}*{amt})"

    def sheet_sub(title):
        nonlocal r
        r += 1
        _put(ws, r, 2, title, F_BOLD)

    sheet_sub("Income statement")
    for role in ln["revenue"]:
        value_row(role, ctx.accounts[role][0], lambda e, rl=role: sumleg(e, ctx.accounts[rl][0], "-"))
    for role in ln["expense"]:
        value_row(role, ctx.accounts[role][0], lambda e, rl=role: sumleg(e, ctx.accounts[rl][0]))
    rev = [rows[x] for x in ln["revenue"]]
    exp = [rows[x] for x in ln["expense"]]

    def total(keys_plus, keys_minus):
        def f(e):
            col = e.split("$")[0]
            plus = "+".join(f"{col}{x}" for x in keys_plus) or "0"
            minus = "".join(f"-{col}{x}" for x in keys_minus)
            return plus + minus
        return f

    value_row("surplus", "Surplus", total(rev, exp), bold=True, border=TOP, note="→ carried to retained surplus")
    sheet_sub("Balance sheet")
    for role in ln["asset"]:
        value_row(role, ctx.accounts[role][0], lambda e, rl=role: sumleg(e, ctx.accounts[rl][0]))
    value_row("assets", "Assets", total([rows[x] for x in ln["asset"]], []), bold=True, border=TOP)
    for role in ln["liability"]:
        value_row(role, ctx.accounts[role][0], lambda e, rl=role: sumleg(e, ctx.accounts[rl][0], "-"))
    value_row("liabilities", "Liabilities", total([rows[x] for x in ln["liability"]], []), bold=True, border=TOP)
    for role in ln["equity"]:
        value_row(role, ctx.accounts[role][0], lambda e, rl=role: sumleg(e, ctx.accounts[rl][0], "-"))
    value_row("retained", "Retained surplus: this period's surplus", lambda e: f"{e.split('$')[0]}{rows['surplus']}",
              note="← from the surplus")
    value_row("equity_total", "Equity", total([rows[x] for x in ln["equity"]] + [rows["retained"]], []), bold=True, border=TOP)
    value_row("balance", "Assets less liabilities less equity (must be nil)",
              lambda e: f"ROUND({e.split('$')[0]}{rows['assets']}-{e.split('$')[0]}{rows['liabilities']}-"
                        f"{e.split('$')[0]}{rows['equity_total']},6)", font=F_NOTE)
    sheet_sub("Cash flow")
    cash_label = ctx.accounts["cash"][0]
    for cc in CASH_CLASSES:
        value_row(cc, f"{cc} cash flows", lambda e, k=cc: sumleg(e, cash_label, "", f"*({cls}=\"{k}\")"))
    value_row("net_cash", "Net cash flow", total([rows[c] for c in CASH_CLASSES], []), bold=True, border=TOP,
              note="→ equals the change in cash")
    value_row("cash_tie", "Net cash flow less the change in cash (must be nil)",
              lambda e: f"ROUND({e.split('$')[0]}{rows['net_cash']}-{e.split('$')[0]}{rows['cash']},6)"
              if "cash" in rows else f"ROUND({e.split('$')[0]}{rows['net_cash']},6)", font=F_NOTE)
    # Notes and the sheet's checks total
    r += 2
    for n in item.notes:
        _put(ws, r, 2, n, F_NOTE)
        r += 1
    r += 1
    _put(ws, r, 3, "Checks on this sheet (must be nil)", F_BOLD, align=IND)
    span = lambda key: f"{L(c0)}{rows[key]}:{L(c0 + len(cols) - 1)}{rows[key]}"     # noqa: E731
    _put(ws, r, 8, f"=ROUND({entries_nil}+SUMPRODUCT(ABS({span('balance')}))+SUMPRODUCT(ABS({span('cash_tie')})),6)",
         F_BOLD, NUM, border=TOP)
    chk = f"{_q(ws)}!$H${r}"
    for rr in range(1, r + 1):
        ws.row_dimensions[rr].height = 15
    ws.freeze_panes = "D5"
    ws._impact_rows = rows                                        # for the tests: where each line sits
    ws._impact_cols = {c: c0 + j for j, c in enumerate(cols)}
    return ws.title, chk


def write_all(wb, ctx: Context, headings: dict | None = None) -> list[str]:
    """Write every Impacts sheet the context supports, each with a name on its checks total (Imp_Chk_1, 2, ...)."""
    names, checks = [], []
    for item in offered(ctx):
        name, chk = write_item(wb, item, ctx, headings)
        names.append(name)
        checks.append(chk)
    for k, chk in enumerate(checks, start=1):              # one name per sheet: a name cannot span sheets
        wb.defined_names[f"Imp_Chk_{k}"] = DefinedName(f"Imp_Chk_{k}", attr_text=chk)
    return names


def checks_formula(n: int) -> str:
    """The error check over every Impacts sheet's total."""
    return "=IF(" + "+".join(f"ABS(Imp_Chk_{k})" for k in range(1, n + 1)) + ">0.001,1,0)" if n else "=0"


# ---------------------------------------------------------------------------------- contexts
def from_consolidation(G) -> Context:
    """The consolidation example's group chart and entities (prototypes/consolidation/group.py)."""
    role_acct = {"cash": G.CASH, "receivables": G.AR, "revenue": G.FEES, "gst": G.GSTPAY, "wip": G.WIP, "inv_property": G.PROP,
                 "payables": G.AP, "netting": G.CLR, "loans": G.LOANPAY, "capital": G.CAPITAL,
                 "property_sales": G.SALES, "construction_revenue": G.CONSTR, "cos": G.COS,
                 "construction_costs": G.CONSTCOST, "opex": G.OPEX, "interest": G.INTEXP}
    accounts = {role: (f"{a} {G.ACCT[a]}", G.CLASS[a]) for role, a in role_acct.items()}
    ent = lambda c: f"{c} {G.NAME[c]}"                                           # noqa: E731
    return Context("Demo Group", accounts, {"entity": ent(9006), "devlp": ent(9006), "builder": ent(9001),
                                            "fund": ent(9008), "parent": ent(9000)})


def from_assembly(layout, model) -> Context:
    """A model built by the assembly engine: its statement lines and module titles."""
    titles = layout.blocks
    first = lambda mod: next((t for b, t in titles.items() if b.startswith(mod + "#") and "/" not in b), None)  # noqa: E731
    accounts = {"cash": ("Cash", "asset")}
    if first("demo.revenue_line"):
        accounts["revenue"] = (first("demo.revenue_line"), "revenue")
    if any(b.startswith("demo.debtors#") for b in titles):
        accounts["receivables"] = ("Debtors", "asset")
    if first("demo.cost_line"):
        accounts["opex"] = (first("demo.cost_line"), "expense")
    if first("demo.facility"):
        accounts["loans"] = (first("demo.facility"), "liability")
        accounts["interest"] = (f"Interest: {first('demo.facility')}", "expense")
    return Context("Assembly demo", accounts, {"entity": "The model's entity"})
