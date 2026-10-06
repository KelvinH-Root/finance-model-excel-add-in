"""Example models: contents, section covers and links, as Modano lays them out, kept current.

The budget and actuals example is the report charts model put into sections. These tests read
its structure with openpyxl, then add a sheet and refresh, as the add-in does after a
structural change. The development example needs Kelvin's template (HFG_TEMPLATE) and
LibreOffice, so it is skipped without them.
"""

import importlib.util
import re
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "prototypes" / "models"))
_spec = importlib.util.spec_from_file_location("reports_build_models", ROOT / "prototypes" / "reports" / "build.py")
B = importlib.util.module_from_spec(_spec)
sys.modules["reports_build_models"] = B
_spec.loader.exec_module(B)

import navigation as N  # noqa: E402

SECTIONS = {"Dashboards": ["Income summary", "Balance summary", "Cash summary", "Budget summary", "Version comparison"],
            "Model": ["Time", "Assumptions", "Scenarios", "Seasonality", "Inputs", "Statements", "Versions"],
            "Reports": ["Income report", "Balance report", "Cash report", "Budget report", "Scenario report"],
            "Appendices": ["Chart register", "Version store", "Lookups", "Checks"]}


@pytest.fixture(scope="module")
def budget(tmp_path_factory):
    path = B.build(tmp_path_factory.mktemp("models") / "budget_example.xlsx")
    nav, order = B.budget_navigation(B.load_register())
    return path, nav


def toc(wb):
    """Contents entries as (level, target name, shown) from the HYPERLINK formulas."""
    out = []
    for row in wb["Contents"].iter_rows(min_row=1):
        for c in row:
            v = c.value
            if isinstance(v, str) and v.startswith('=HYPERLINK("#HL_') and c.column in (3, 4, 5):
                m = re.match(r'=HYPERLINK\("#(HL_\w+)",(.*)\)$', v)
                if m.group(1) != "HL_Err_Chk":                                # the checks block at the end
                    out.append((c.column - 2, m.group(1), m.group(2)))
    return out


def test_sheets_sit_in_sections_behind_their_covers(budget):
    wb = load_workbook(budget[0])
    expected = ["Contents"] + [s for cover, sheets in SECTIONS.items() for s in [cover] + sheets]
    assert wb.sheetnames == expected
    entries = [(lvl, nm) for lvl, nm, _ in toc(wb) if lvl < 3]
    want = [(1 if s in SECTIONS else 2, N.sheet_name(s)) for s in expected[1:]]
    assert entries == want                                                   # every section and every sheet, in order
    third = [shown for lvl, _, shown in toc(wb) if lvl == 3]
    assert '"Seasonality"' not in third and '"Profile"' in third and '"Chart data"' in third
    for s in expected:
        assert wb.defined_names[N.sheet_name(s)].attr_text == N.ref(s, "$A$1")
    assert wb.defined_names["HL_Home"].attr_text == "Contents!$B$1"
    assert wb.defined_names["HL_Err_Chk"].attr_text == "Checks!$B$1"


def test_every_sheet_links_to_the_contents_and_the_checks(budget):
    wb = load_workbook(budget[0])
    for ws in wb.worksheets[1:]:
        assert ws["A1"].value.startswith('=HYPERLINK("#HL_Home"'), ws.title
        assert ws["A2"].value.startswith('=HYPERLINK("#HL_Err_Chk",IF(Chk_Errors=0'), ws.title


def test_covers_number_their_section_and_link_either_side(budget):
    wb = load_workbook(budget[0])
    model = wb["Model"]
    assert model["B9"].value == "Financial Model" and model["B10"].value == "Section 2."
    assert model["B12"].value == '=HYPERLINK("#HL_Home","Go to contents")'
    assert "HL_Sheet_Version_comparison" in model["B13"].value and "'Version comparison'!$B$1" in model["B13"].value
    assert "HL_Sheet_Time" in model["B14"].value
    assert wb["Appendices"]["B10"].value == "Section 4."


def test_a_new_sheet_is_filed_where_it_sits_and_the_contents_follow(budget):
    path, nav = budget
    wb = load_workbook(path)
    ws = wb.create_sheet("Capex plan", wb.sheetnames.index("Inputs") + 1)   # a sheet added after Inputs
    ws["B1"] = "Capex plan"
    wb.move_sheet("Lookups", offset=-(wb.sheetnames.index("Lookups") - wb.sheetnames.index("Statements")) + 1)
    N.refresh(wb, nav)
    model = [nm for lvl, nm, _ in toc(wb) if lvl == 2][5:14]
    assert model == [N.sheet_name(s) for s in ["Time", "Assumptions", "Scenarios", "Seasonality", "Inputs",
                                                 "Capex plan", "Statements", "Lookups", "Versions"]]
    letters = [c.value for c in wb["Contents"]["C"] if isinstance(c.value, str) and c.value.endswith(".")]
    assert letters.count("i.") == 1                                          # nine sheets in Financial Model now
    assert ws["A1"].value.startswith('=HYPERLINK("#HL_Home"') and "HL_Sheet_Capex_plan" in wb.defined_names
    assert "HL_Sheet_Chart_register" in wb["Appendices"]["B14"].value       # Appendices now starts at Chart register
    assert "HL_Sheet_Versions" in wb["Reports"]["B13"].value                # Reports' previous sheet is still Versions
    del wb["Capex plan"]
    N.refresh(wb, nav)
    assert "HL_Sheet_Capex_plan" not in wb.defined_names


# ---------------------------------------------------------------- development example (needs the template)

def test_development_example_has_contents_covers_and_links(tmp_path):
    from conftest import template_or_skip
    template = template_or_skip()
    sys.path.insert(0, str(ROOT / "prototypes" / "models"))
    import development
    path = development.build(tmp_path / "development_example.xlsx", template)
    wb = load_workbook(path)
    assert wb.sheetnames[0] == "Contents"
    covers = [s for s in wb.sheetnames if s in development.COVERS]
    assert covers == list(development.COVERS)
    for ws in wb.worksheets[1:]:
        assert ws["A1"].value.startswith('=HYPERLINK("#HL_Home"'), ws.title
    from hfgmodels.verify import recalculate
    cw = recalculate(path, wb.sheetnames)
    errors = {k: v for k, v in cw.errors.items() if v}
    assert errors in ({}, {"Time": ["Time.H16:H17"]}), errors     # the template's LET formulas: Excel only (LibreOffice 24.8+)
