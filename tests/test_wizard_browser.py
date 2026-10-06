"""The New model wizard in a browser, as the probe's task pane shows it.

Steps through the wizard in Chromium (the probe's styles and the generated wizard bundle),
creates a model outside Excel (so the pane offers a download), and checks the file: the entity,
title, timeline and recipe chosen are in it, and it recalculates in LibreOffice with no errors.
Screenshots of each step go to build/wizard/. Needs Playwright with Chromium; fictional data only.
"""

import base64
import shutil
import sys
from pathlib import Path

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parent.parent
PROBE = ROOT / "addin-probe" / "src"
BUNDLE = PROBE / "wizard.bundle.js"
SHOTS = ROOT / "build" / "wizard"

playwright = pytest.importorskip("playwright.sync_api")
pytestmark = pytest.mark.skipif(not BUNDLE.exists(), reason="run npm run build in addin/ first")

PAGE = """<!doctype html><html><head><meta charset="utf-8"><link rel="stylesheet" href="{css}">
<style>body {{ width: 340px; }}</style></head>
<body><main><div class="view"><h2>New model</h2><div id="app" class="gapp"></div></div></main>
<script src="{bundle}"></script>
<script>window.HfgWizard.render('model-new', document.getElementById('app'), () => {{}});</script></body></html>"""


@pytest.fixture(scope="module")
def page(tmp_path_factory):
    d = tmp_path_factory.mktemp("wizard")
    html = d / "wizard.html"
    html.write_text(PAGE.format(css=(PROBE / "taskpane.css").as_uri(), bundle=BUNDLE.as_uri()))
    with playwright.sync_playwright() as p:
        browser = p.chromium.launch()
        pg = browser.new_page(viewport={"width": 360, "height": 900}, device_scale_factor=2)
        errors = []
        pg.on("pageerror", lambda e: errors.append(str(e)))
        pg.goto(html.as_uri())
        yield pg, errors
        browser.close()


def shot(pg, name):
    SHOTS.mkdir(parents=True, exist_ok=True)
    pg.locator("main").screenshot(path=str(SHOTS / f"{name}.png"))


def test_wizard_steps_and_create(page, tmp_path):
    pg, errors = page
    # 1 Entity: choose Home Capital Partners and name a development LP under its brand.
    assert pg.locator(".wsteps li.on").inner_text() == "1. Entity"
    pg.locator("input[name=wbrand][value=HCP]").check()
    pg.locator(".wfield input[type=text]").fill("Demo Development LP")
    shot(pg, "1_entity")
    pg.get_by_role("button", name="Next").click()
    # 2 Model: Next without a title is refused with a reason.
    pg.get_by_role("button", name="Next").click()
    assert "Give the model a title." in pg.locator(".werr").inner_text()
    pg.get_by_placeholder("For example, Operating budget FY2027").fill("Development budget FY2027")
    pg.locator("input[name=wrecipe][value=demo]").check()
    pg.locator("textarea").fill("Fictional demo data.\nBuilt by the New model wizard.")
    shot(pg, "2_model")
    pg.get_by_role("button", name="Next").click()
    # 3 Timeline: 24 months, six of them actual.
    pg.locator("input[type=number]").nth(0).fill("24")
    pg.locator("input[type=number]").nth(1).fill("6")
    assert "Actuals to September 2026" in pg.locator("main").inner_text()
    shot(pg, "3_timeline")
    pg.get_by_role("button", name="Next").click()
    shot(pg, "4_display")
    pg.get_by_role("button", name="Next").click()
    text = pg.locator("main").inner_text()
    assert "24 months, April 2026 to March 2028" in text and "actuals to September 2026" in text
    assert "Demo Development LP" in text and "Financial Model" in text
    shot(pg, "5_review")
    pg.get_by_role("button", name="Create model").click()
    link = pg.get_by_role("link", name="Download a copy")
    link.wait_for()
    shot(pg, "6_created")
    assert "Excel is not running this pane" in pg.locator("[role=status]").inner_text()
    data = pg.evaluate("""async () => {
        const a = document.querySelector('a[download]');
        const buf = new Uint8Array(await (await fetch(a.href)).arrayBuffer());
        let s = ''; for (const b of buf) s += String.fromCharCode(b);
        return { name: a.getAttribute('download'), b64: btoa(s) };
    }""")
    assert data["name"] == "Development budget FY2027.xlsx"
    path = tmp_path / data["name"]
    path.write_bytes(base64.b64decode(data["b64"]))
    assert not errors, errors

    wb = openpyxl.load_workbook(path)
    settings = wb["Settings"]
    values = {settings.cell(r, 3).value: settings.cell(r, 9).value for r in range(17, 40) if settings.cell(r, 3).value}
    assert values["Model title"] == "Development budget FY2027"
    assert values["Entity"] == "Demo Development LP"
    assert values["Last month of actuals (period number, 0 for none)"] == 6
    assert values["Months in the model"] == 24
    assert "Revenue" in wb.sheetnames and wb["Contents"]["B1"].value == "=Model_Entity"

    if shutil.which("soffice"):
        sys.path.insert(0, str(ROOT))
        from hfgmodels.verify import recalculate
        calc = recalculate(path, wb.sheetnames)
        assert all(not e for e in calc.errors.values()), calc.errors
        assert calc.get("Contents", 1, 2) == "Demo Development LP"
        assert calc.get("Settings", 6, 10) == "Actual" and calc.get("Settings", 6, 16) == "Forecast"
