"""The TypeScript engine's package writer, checked from Python.

1. In the proof's frame, the workbook the engine writes equals the Python proof's: every cell, every
   defined name and every chart series.
2. In HFG's standard frame, for every entity, the workbook recalculates in LibreOffice with no
   formula errors, the statements match the reference calculation, the timeline block and the
   model name line work as the Look and wiring standard says, every link goes to a defined name,
   and every colour in the cell styles is a theme slot.

Fictional demo data only. Needs Node 22.18 or later with `npm ci` run in engine/.
"""

import re
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parent.parent
ENGINE = ROOT / "engine"
sys.path.insert(0, str(ROOT / "prototypes" / "assembly"))

from assemble import Library, Model, assemble, read_metadata, write_workbook  # noqa: E402
from reference import reference  # noqa: E402

BRANDS = ["HF", "HCP", "HCL", "KM", "TWK"]
PERIODS = 12


def _node_ok() -> bool:
    node = shutil.which("node")
    if not node or not (ENGINE / "node_modules").exists():
        return False
    major, minor, *_ = subprocess.run([node, "--version"], capture_output=True, text=True).stdout.strip().lstrip("v").split(".")
    return (int(major), int(minor)) >= (22, 18)


pytestmark = pytest.mark.skipif(not _node_ok(), reason="needs Node 22.18 or later and npm ci in engine/")
needs_lo = pytest.mark.skipif(not shutil.which("soffice"), reason="LibreOffice is not installed")


def build(out: Path, *args: str) -> Path:
    subprocess.run(["node", "tools/build.ts", str(out), *args], cwd=ENGINE, check=True, capture_output=True)
    return out


def cell_map(path: Path) -> dict:
    wb = openpyxl.load_workbook(path)
    return {ws.title: {c.coordinate: c.value for row in ws.iter_rows() for c in row if c.value is not None} for ws in wb}


def name_map(path: Path) -> dict:
    wb = openpyxl.load_workbook(path)
    return {n: d.attr_text for n, d in wb.defined_names.items()}


def chart_refs(path: Path) -> list:
    """Every chart's ranges, with quotes dropped from sheet names that do not need them."""
    out = []
    with zipfile.ZipFile(path) as z:
        for n in sorted(z.namelist()):
            if n.startswith("xl/charts/chart"):
                refs = re.findall(r"<(?:\w+:)?f>([^<]+)</(?:\w+:)?f>", z.read(n).decode())
                out.append(sorted(re.sub(r"^'([A-Za-z_][A-Za-z0-9_]*)'!", r"\1!", r) for r in refs))
    return sorted(out)


@pytest.fixture(scope="module")
def lib():
    return Library.load()


@pytest.fixture(scope="module")
def proof_pair(tmp_path_factory, lib):
    d = tmp_path_factory.mktemp("proof")
    ts = build(d / "engine.xlsx", "--proof")
    model = Model.from_dict(lib, read_metadata(ts)["model"])
    py = write_workbook(assemble(model), d / "python.xlsx", model)
    return ts, py


def test_proof_frame_cells_match_the_python_writer(proof_pair):
    ts, py = proof_pair
    a, b = cell_map(ts), cell_map(py)
    assert list(a) == list(b)
    for sheet in a:
        assert a[sheet] == b[sheet], sheet


def test_proof_frame_names_and_charts_match_the_python_writer(proof_pair):
    ts, py = proof_pair
    assert name_map(ts) == name_map(py)
    assert chart_refs(ts) and chart_refs(ts) == chart_refs(py)


@pytest.fixture(scope="module")
def standard(tmp_path_factory):
    d = tmp_path_factory.mktemp("standard")
    return {brand: build(d / f"demo_{brand}.xlsx", "--brand", brand) for brand in BRANDS}


def test_standard_frame_structure(standard, lib):
    for brand, path in standard.items():
        wb = openpyxl.load_workbook(path)
        assert wb.sheetnames[0] == "Contents" and "Settings" in wb.sheetnames
        names = set(wb.defined_names.keys())
        assert {"Model_Name", "Model_Title", "Tl_Start", "Tl_Last_Actual", "Go_Contents", "Go_Checks", "Chk_Errors"} <= names
        assert not any(n.startswith(("HL_", "Ts_", "DD_", "CB_", "LU_", "Err_")) for n in names), "HFG's own name prefixes only"
        for ws in wb:
            for link in ws._hyperlinks:
                assert link.location in names, f"{brand} {ws.title}!{link.ref} links to {link.location}"
                assert link.tooltip, f"{brand} {ws.title}!{link.ref} has no screen tip"
        rev = wb["Revenue"]
        assert rev.freeze_panes == "J16"
        assert rev.sheet_format.defaultRowHeight == pytest.approx(11.4)
        assert [rev.row_dimensions[r].outlineLevel for r in range(7, 16)] == [2] * 9
        assert rev["B2"].value == "=Model_Name" and rev["A1"].value == "⌂"
        meta = read_metadata(path)
        assert meta["model"]["info"]["entity"]["brand"] == brand


def test_standard_styles_use_theme_colours_only(standard):
    for brand, path in standard.items():
        with zipfile.ZipFile(path) as z:
            styles = z.read("xl/styles.xml").decode()
            theme = z.read("xl/theme/theme1.xml").decode()
        cell_part = styles.split("<dxfs")[0]
        assert "rgb=" not in cell_part, f"{brand}: a cell style has a typed colour"
        assert "HFG Heading 1" in styles and "HFG Input Number" in styles
        assert 'typeface="Segoe UI"' in theme
        assert "Modano" not in styles + theme


@needs_lo
def test_standard_frame_recalculates_and_matches_the_reference(standard, lib):
    from hfgmodels.verify import recalculate

    for brand, path in standard.items():
        meta = read_metadata(path)
        model = Model.from_dict(lib, meta["model"])
        sheets = list(meta["rows"])
        calc = recalculate(path, sheets)
        assert all(not e for e in calc.errors.values()), f"{brand}: {calc.errors}"
        ref = reference(model)
        rows = meta["rows"]["Statements"]

        def series(key):
            r = 17 + rows.index(f"demo.statements#1/{key}")
            return [calc.get("Statements", r, 10 + t) for t in range(PERIODS)]

        assert series("revenue") == pytest.approx(ref["revenue"], abs=1e-6)
        assert series("cash") == pytest.approx(ref["cash"], abs=1e-6)
        assert series("debt") == pytest.approx(ref["debt"], abs=1e-6)
        assert series("equity") == pytest.approx(ref["equity"], abs=1e-6)
        # Timeline block, worked out on each calculation sheet from Settings (which has no months across it):
        # April 2026 start, March year end, no actuals yet.
        assert calc.get("Revenue", 5, 10) == 46142             # 30 April 2026
        assert calc.get("Revenue", 10, 10) == 2027              # in the year to March 2027
        assert calc.get("Revenue", 11, 10) == 1                 # first month of that year
        assert calc.get("Revenue", 6, 10) == "Forecast"
        assert calc.get("Revenue", 5, 21) == calc.get("Statements", 5, 21) == 46477   # 31 March 2027
        assert calc.get("Settings", 5, 10) in (None, "")
        assert calc.get("Contents", 2, 2) == "Demo operating model"
        assert calc.get("Revenue", 2, 1) == "✓"


@needs_lo
def test_model_name_line_shows_alerts_when_raised(standard):
    """Raise an alert (cash below zero) and the model name line on every sheet says so; switch the phrase off and it goes."""
    from hfgmodels.verify import libreoffice, _prop

    path = standard["HF"].resolve()
    with libreoffice() as desktop:
        doc = desktop.loadComponentFromURL("file://" + str(path), "_blank", 0, (_prop("Hidden", True),))
        try:
            doc.calculateAll()
            names = doc.NamedRanges
            amount = names.getByName("Cost1_Amount").getReferredCells().getCellByPosition(0, 0)
            show = names.getByName("Opt_Show_Alerts").getReferredCells().getCellByPosition(0, 0)
            contents = doc.Sheets.getByName("Contents").getCellRangeByName("B2")
            revenue = doc.Sheets.getByName("Revenue").getCellRangeByName("B2")
            assert contents.getString() == "Demo operating model"
            amount.setValue(900)
            doc.calculateAll()
            assert re.fullmatch(r"Demo operating model \(\d+ alerts\)", contents.getString()), contents.getString()
            assert revenue.getString() == contents.getString()
            show.setFormula("=FALSE()")
            doc.calculateAll()
            assert contents.getString() == "Demo operating model"
        finally:
            doc.close(True)
