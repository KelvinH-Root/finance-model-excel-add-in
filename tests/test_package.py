"""Package editing on a small workbook made with openpyxl (no template needed)."""

import zipfile

import openpyxl
from openpyxl.workbook.defined_name import DefinedName

from hfgmodels.xlsx.package import TemplatePackage
from hfgmodels.xlsx.sheet import Sheet


def _make(tmp_path):
    wb = openpyxl.Workbook()
    wb.active.title = "First"
    wb.create_sheet("Second")
    wb.create_sheet("Third")
    wb["First"]["A1"] = "keep"
    wb["Third"]["B2"] = 5
    dn = DefinedName("LocalName", localSheetId=2, attr_text="Third!$B$2")
    wb["Third"].defined_names.add(dn)
    wb.defined_names["GlobalName"] = DefinedName("GlobalName", attr_text="First!$A$1")
    p = tmp_path / "t.xlsx"
    wb.save(p)
    return p


def test_add_sheet_keeps_parts_and_shifts_local_names(tmp_path):
    src = _make(tmp_path)
    pkg = TemplatePackage(src)
    before = set(pkg.parts)
    sh = Sheet("New")
    sh.set(1, 1, "hello")
    sh.set(2, 1, formula="Third!B2*2")
    pkg.add_sheet("New", sh.to_xml(), after="First")
    pkg.set_defined_name("NewName", "New!$A$2")
    out = pkg.save(tmp_path / "out.xlsx")
    assert before <= set(zipfile.ZipFile(out).namelist())
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ["First", "New", "Second", "Third"]
    assert wb["New"]["A1"].value == "hello"
    assert wb["New"]["A2"].value == "=Third!B2*2"
    assert wb["First"]["A1"].value == "keep"
    # the sheet-scoped name moved with its sheet (index 2 -> 3)
    assert "LocalName" in wb["Third"].defined_names
    assert wb.defined_names["NewName"].attr_text == "New!$A$2"


def test_set_cell_value_and_append_cells(tmp_path):
    src = _make(tmp_path)
    pkg = TemplatePackage(src)
    pkg.set_cell_value("Third", "B2", 42)
    pkg.append_cells("Third", [("B10", "added", None), ("C10", 3.5, None)])
    out = pkg.save(tmp_path / "out2.xlsx")
    ws = openpyxl.load_workbook(out)["Third"]
    assert ws["B2"].value == 42
    assert ws["B10"].value == "added"
    assert ws["C10"].value == 3.5


def test_sheet_xml_element_order():
    sh = Sheet("X")
    sh.set(1, 1, "a")
    sh.validations.append(("A1", '"Yes,No"'))
    sh.hyperlinks.append(("A1", "X!A1", "a"))
    xml = sh.to_xml().decode()
    assert xml.index("<dataValidations") < xml.index("<hyperlinks>") < xml.index("<pageMargins")
